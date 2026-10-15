import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from data_report_api import create_app
from data_report_api.config import TestConfig
from data_report_api.services import resource_governance as module


class ResourceGovernanceTests(unittest.TestCase):
    def setUp(self):
        self.source = MagicMock()
        self.source.database = "sibebid"
        self.source.label = "MySQL · sibebid"
        self.connection = self.source.connect.return_value
        self.cursor = self.connection.cursor.return_value
        self.last_sql = ""
        self.source_patch = patch.object(module, "DataSource", return_value=self.source)
        self.source_patch.start()
        self.addCleanup(self.source_patch.stop)

    @staticmethod
    def metric(*, segments=20, segment_missing=0, search=100000, search_success=95000):
        return (
            3, search, search_success, 2000, 1800, 25, 20, segments, segment_missing,
            15000, 200, 2, 500000, 40000, 5000, 0, 18, 1,
        )

    def configure_gds(self, *, segments=20, segment_missing=0):
        summary = self.metric(segments=segments, segment_missing=segment_missing)
        group = self.metric(
            segments=None, segment_missing=None, search=10000, search_success=9500
        )

        def execute(sql, params=()):
            self.last_sql = sql

        def fetchone():
            if "MAX(etl_updated_at)" in self.last_sql:
                return summary + (datetime(2026, 10, 15, 8, 30),)
            return None

        def fetchall():
            if "SELECT DISTINCT" in self.last_sql:
                return [
                    ("CTRIP", "CTRIP_HK", "1A", "HKGH228DF", module.GDS_MP_SCOPE),
                    ("CTRIP", "CTRIP_HK", "1A", "HKGH228DF", "1A"),
                    ("CTRIP", "CTRIP_HK", "1A", "HKGH228DF", "1A_MPIS"),
                ]
            if "GROUP BY operator_date" in self.last_sql:
                return [("2026-10-14",) + summary]
            if "GROUP BY source_api_group" in self.last_sql:
                return [("1A",) + group, ("1A_MPIS",) + group]
            return [("1A",) + summary]

        self.cursor.execute.side_effect = execute
        self.cursor.fetchone.side_effect = fetchone
        self.cursor.fetchall.side_effect = fetchall

    def configure_ota(self):
        summary = self.metric(segments=None, segment_missing=None)

        def execute(sql, params=()):
            self.last_sql = sql

        def fetchone():
            return summary

        def fetchall():
            if "SELECT DISTINCT" in self.last_sql:
                return [("携程", "携程一部", "", "", "OTA")]
            if "GROUP BY operator_date" in self.last_sql:
                return [("20261014",) + summary]
            return [("携程",) + summary]

        self.cursor.execute.side_effect = execute
        self.cursor.fetchone.side_effect = fetchone
        self.cursor.fetchall.side_effect = fetchall

    def test_gds_uses_service_table_mp_total_and_exposes_api_group_quality(self):
        self.configure_gds()
        result = module.ResourceGovernanceService(
            {
                "RESOURCE_GOVERNANCE_GDS_TABLE": "bi_gds_resource_metric_day",
                "RESOURCE_GOVERNANCE_GDS_TARGETS": "1A:4000,1G:3000",
            }
        ).overview(
            "2026-10-01", "2026-10-14", data_source="gds", group_by="gds",
            filters={"gds": "1A"},
        )

        self.assertTrue(result["available"])
        self.assertEqual(result["table"], "sibebid.bi_gds_resource_metric_day")
        self.assertEqual(result["summary"]["segmentCount"], 20)
        self.assertEqual(result["summary"]["queryToSegmentRatio"], "5000.00")
        self.assertEqual(result["summary"]["excessQueryCount"], 20000)
        self.assertEqual(result["summary"]["targetRatio"], 4000)
        self.assertTrue(result["ratioApplicable"])
        self.assertEqual(result["trend"][0]["date"], "2026-10-14")
        self.assertEqual(result["options"]["apiGroup"], ["1A", "1A_MPIS"])
        self.assertEqual(result["apiGroups"][0]["searchRequests"], 10000)
        self.assertEqual(result["apiGroups"][0]["searchFailure"], 500)
        self.assertEqual(result["apiGroups"][0]["searchFailureRate"], "5.00")
        self.assertEqual(result["apiGroups"][0]["verifyFailure"], 200)
        self.assertEqual(result["apiGroups"][0]["verifyFailureRate"], "10.00")
        self.assertEqual(result["apiGroups"][0]["orderFailure"], 5)
        self.assertEqual(result["apiGroups"][0]["orderFailureRate"], "20.00")
        executed = [call.args for call in self.cursor.execute.call_args_list]
        self.assertTrue(any("metric_scope = %s" in sql and "SUM(seg_num)" in sql for sql, *_ in executed))
        self.assertTrue(any(params[:3] == ("2026-10-01", "2026-10-14", "MP_TOTAL") for _, params in executed))

    def test_gds_api_group_filter_disables_ratio_and_uses_group_scope(self):
        self.configure_gds()
        result = module.ResourceGovernanceService({}).overview(
            "2026-10-14", "2026-10-14", data_source="gds", group_by="platform",
            filters={"apiGroup": "1A_NDC"},
        )

        self.assertFalse(result["ratioApplicable"])
        self.assertEqual(result["summary"]["status"], "not_applicable")
        self.assertIsNone(result["summary"]["segmentCount"])
        self.assertIsNone(result["summary"]["queryToSegmentRatio"])
        summary_call = next(
            call.args for call in self.cursor.execute.call_args_list
            if "MAX(etl_updated_at)" in call.args[0]
        )
        self.assertNotIn("SUM(seg_num)", summary_call[0])
        self.assertEqual(summary_call[1][2:], ("API_GROUP", "1A_NDC"))

    def test_partial_mp_segment_sync_blocks_ratio_and_excess(self):
        self.configure_gds(segment_missing=1)
        result = module.ResourceGovernanceService(
            {"RESOURCE_GOVERNANCE_GDS_TARGETS": "1A:4000"}
        ).overview(
            "2026-10-14", "2026-10-14", data_source="gds", filters={"gds": "1A"},
        )
        self.assertEqual(result["summary"]["status"], "segment_partial")
        self.assertEqual(result["summary"]["segmentMissingRowCount"], 1)
        self.assertIsNone(result["summary"]["queryToSegmentRatio"])
        self.assertIsNone(result["summary"]["excessQueryCount"])
        self.assertFalse(result["segmentDataAvailable"])

    def test_ota_keeps_request_fact_table_and_does_not_query_segments(self):
        self.configure_ota()
        result = module.ResourceGovernanceService(
            {
                "RESOURCE_GOVERNANCE_OTA_SUCCESS_TARGETS":
                    '{"携程":{"search":96,"verify":95,"order":90}}'
            }
        ).overview(
            "2026-10-14", "2026-10-14", data_source="ota", group_by="pcc",
            filters={"platform": "携程"},
        )
        self.assertEqual(result["table"], "sibebid.bi_ota_gds_request_day")
        self.assertEqual(result["dimension"]["key"], "platform")
        self.assertFalse(result["ratioApplicable"])
        self.assertEqual(result["summary"]["status"], "not_applicable")
        self.assertEqual(result["summary"]["searchFailure"], 5000)
        self.assertEqual(result["summary"]["successTargets"]["search"], "96.00")
        executed_sql = "\n".join(call.args[0] for call in self.cursor.execute.call_args_list)
        self.assertNotIn("bi_gds_resource_metric_day", executed_sql)
        self.assertNotIn("SUM(seg_num)", executed_sql)
        self.assertNotIn("create_time", executed_sql)
        self.assertTrue(any(call.args[1][:3] == ("20261014", "20261014", 0) for call in self.cursor.execute.call_args_list))

    def test_gds_rows_flag_objects_above_target_and_breach_only_filters(self):
        self.configure_gds()
        result = module.ResourceGovernanceService(
            {"RESOURCE_GOVERNANCE_GDS_TARGETS": "1A:4000"}
        ).overview(
            "2026-10-14", "2026-10-14", data_source="gds", group_by="gds",
            breach_only=True,
        )
        self.assertFalse(result["summary"]["breach"])
        self.assertTrue(result["dimension"]["rows"][0]["breach"])
        self.assertEqual(result["breachObjectCount"], 1)

    def test_ota_rows_flag_stages_below_platform_target(self):
        self.configure_ota()
        result = module.ResourceGovernanceService(
            {
                "RESOURCE_GOVERNANCE_OTA_SUCCESS_TARGETS":
                    '{"携程":{"search":96,"verify":95,"order":10}}'
            }
        ).overview(
            "2026-10-14", "2026-10-14", data_source="ota", group_by="platform",
            filters={"platform": "携程"},
        )
        reasons = result["summary"]["breachReasons"]
        self.assertIn("查询成功率95.00%低于目标96.00%", reasons)
        self.assertIn("验价成功率90.00%低于目标95.00%", reasons)
        self.assertFalse(any("生单成功率" in reason for reason in reasons))

    def test_invalid_inputs_fail_before_connecting(self):
        service = module.ResourceGovernanceService({})
        for kwargs in (
            {"start_value": "bad-date"},
            {"start_value": "2026-10-15", "end_value": "2026-10-14"},
            {"data_source": "all"},
            {"group_by": "airline"},
        ):
            with self.assertRaises(ValueError):
                service.overview(**kwargs)
        self.source.connect.assert_not_called()

    def test_separate_routes_fix_their_own_data_layer(self):
        app = create_app(TestConfig)
        with patch("data_report_api.routes.ResourceGovernanceService") as service_class:
            service_class.return_value.overview.return_value = {"available": True}
            client = app.test_client()
            self.assertEqual(client.get("/api/v1/resource-governance/gds/overview").status_code, 200)
            self.assertEqual(service_class.return_value.overview.call_args.kwargs["data_source"], "gds")
            self.assertEqual(client.get("/api/v1/resource-governance/ota/overview").status_code, 200)
            self.assertEqual(service_class.return_value.overview.call_args.kwargs["data_source"], "ota")

    def test_route_passes_breach_only_flag(self):
        app = create_app(TestConfig)
        with patch("data_report_api.routes.ResourceGovernanceService") as service_class:
            service_class.return_value.overview.return_value = {"available": True}
            client = app.test_client()
            client.get("/api/v1/resource-governance/gds/overview?breachOnly=true")
            self.assertTrue(service_class.return_value.overview.call_args.kwargs["breach_only"])

    def test_gds_report_export_contains_current_summary_and_analysis_sheets(self):
        self.configure_gds()
        content, filename = module.ResourceGovernanceService(
            {"RESOURCE_GOVERNANCE_GDS_TARGETS": "1A:4000"}
        ).export_report(
            start_value="2026-10-01", end_value="2026-10-14", data_source="gds",
            group_by="gds", filters={"gds": "1A"},
        )

        workbook = load_workbook(content, read_only=True, data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["管理摘要", "日趋势", "维度对比", "接口分组质量"],
        )
        summary = dict(workbook["管理摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["MP查询量"], 100000)
        self.assertEqual(summary["航段数"], 20)
        self.assertEqual(summary["查定比"], "5000.00:1")
        self.assertEqual(summary["治理目标"], "4000:1")
        self.assertIn("GDS资源治理分析_2026-10-01_2026-10-14", filename)

    def test_ota_report_keeps_query_volume_and_has_no_gds_group_sheet(self):
        self.configure_ota()
        content, filename = module.ResourceGovernanceService({}).export_report(
            start_value="2026-10-14", end_value="2026-10-14", data_source="ota",
            group_by="platform", filters={"platform": "携程"},
        )

        workbook = load_workbook(content, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["管理摘要", "日趋势", "维度对比"])
        headers = next(workbook["日趋势"].iter_rows(values_only=True))
        self.assertIn("查询请求量", headers)
        self.assertIn("OTA资源治理分析_2026-10-14_2026-10-14", filename)

    def test_export_route_downloads_workbook_for_fixed_layer(self):
        app = create_app(TestConfig)
        with patch("data_report_api.routes.ResourceGovernanceService") as service_class:
            service_class.return_value.export_report.return_value = (
                BytesIO(b"report"), "GDS资源治理分析.xlsx"
            )
            response = app.test_client().get(
                "/api/v1/resource-governance/gds/export?startDate=2026-10-14&endDate=2026-10-14"
            )

            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data, b"report")
            self.assertEqual(
                service_class.return_value.export_report.call_args.kwargs["data_source"],
                "gds",
            )


if __name__ == "__main__":
    unittest.main()
