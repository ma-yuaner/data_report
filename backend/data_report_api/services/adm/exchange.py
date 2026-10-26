"""Versioned ADM workbooks and atomic-per-case, idempotent intake.

Legacy imports retain original fields as source evidence; they never grant review,
payment, recovery or closure states simply because a spreadsheet contains them.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import tempfile
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from ...time_utils import business_now_naive
from ...shared.permissions import permission_codes
from . import AdmError, AdmService
from .service import CASE_FIELDS
from .attachments import validate_document


CASE_COLUMNS = [
    ("key", "主单标识"),
    ("admNo", "ADM单号"),
    ("supplierName", "下发方"),
    ("supplierType", "下发方类型"),
    ("currency", "币种"),
    ("originalAmount", "下发金额"),
    ("issuedDate", "下发日期"),
    ("deadline", "截止时间"),
    ("description", "问题描述"),
    ("remark", "备注"),
    ("supplierCode", "下发方编码"), ("airline", "航司"), ("ticketCount", "原票号数量"),
    ("originalAmountCny", "人民币金额"), ("businessCategory", "业务大类编码"),
]
LINE_COLUMNS = [
    ("id", "明细标识"),
    ("caseKey", "主单标识"),
    ("businessType", "业务类型"),
    ("platform", "平台"),
    ("shop", "店铺"),
    ("site", "细分站点"),
    ("orderNo", "OTA订单号"),
    ("ticketNo", "票号"),
    ("pnr", "PNR"),
    ("system", "系统"),
    ("pcc", "PCC"),
    ("currency", "币种"),
    ("amount", "分摊金额"),
    ("diffType", "差异类型"),
    ("diffSecondType", "差异二级类型"),
    ("differenceDescription", "差异说明"),
    ("description", "问题描述"),
    ("remark", "备注"),
    ("platformCode", "平台编码"), ("airline", "航司"), ("ticketCount", "票号数量"),
    ("amountCny", "人民币分摊金额"), ("businessCategory", "业务大类编码"),
    ("actualOwner", "实际责任人"), ("actualHandler", "实际处理人"),
    ("recoveryDeadline", "编码恢复时限"), ("resolution", "处理结果说明"),
]
LEGACY38 = [
    ("id", "主键ID"),
    ("adm_no", "ADM单号（业务唯一标识）"),
    ("ota_code", "OTA平台"),
    ("ota_order_no", "OTA订单号"),
    ("lock_flag", "锁状态"),
    ("lock_operator_id", "加锁人id"),
    ("lock_operator", "加锁人"),
    ("supplier_code", "供应商"),
    ("airline", "航司（多个逗号隔开）"),
    ("supplier_type_name", "类型"),
    ("ticket_no", "票号"),
    ("ticket_count", "票号数量（冗余统计，对应adm_details行数，方便列表展示）"),
    ("total_amount", "ADM总金额（冗余汇总，= SUM(adm_details.amount)，方便列表排序展示）"),
    ("currency", "主币种（CNY/USD/EUR等，明细行可不同币种时以明细为准）"),
    ("supply_issue_date", "供应下发日期"),
    ("adm_deadline", "ADM最晚时限（回复截至时间）"),
    ("category_name", "大类"),
    ("diff_type_name", "差异类型"),
    ("diff_detail_reason", "细分差异原因（可备注，自由文本描述具体差异情况）"),
    ("owner", "当前责任人（处理人，中间可转手，历史记录见flow_logs）"),
    ("actual_owner", "实际责任人"),
    ("adm_status_name", "处理状态"),
    ("appeal_status_name", "申诉状态"),
    ("appeal_result_name", "申诉结果"),
    ("appeal_reason", "申诉原因（status=申诉时必填）"),
    ("resolution", "结案处理结果（status=已结案时必填，如：已退款¥XXX / 差异已消化 / 关单无差异等）"),
    ("handle_time", "处理时间"),
    ("auditor_name", "审核人"),
    ("auditor_time", "审核时间"),
    ("auditor_remark", "审核备注"),
    ("actual_handler", "实际处理人"),
    ("file_url", "附件地址"),
    ("status", "状态：0-无效(删除),1-有效(正常),2-暂停"),
    ("remark", "备注"),
    ("create_time", "创建时间"),
    ("update_time", "更新时间"),
    ("create_user_name", "创建人"),
    ("update_user_name", "修改人"),
]
LEGACY_EXTRA = {
    "currency": [("total_amount_cny", "ADM总金额（CNY）"), ("pnr_no", "PNR编码"), ("pnr_recover_time", "PNR恢复时间")],
    "diff_type_name": [("diff_second_type", "差异二级类型")],
    "actual_owner": [("assist_owner", "责任人")],
    "appeal_reason": [("appeal_target", "申诉对象")],
}
LEGACY44 = [column for item in LEGACY38 for column in [item, *LEGACY_EXTRA.get(item[0], [])]]


def check_permission(user: dict, permission: str) -> None:
    if not user.get("id"):
        raise AdmError("请先登录", 401, "ADM_AUTH")
    if not user.get("is_admin") and permission not in permission_codes(user):
        raise AdmError("没有此ADM操作权限", 403, "ADM_FORBIDDEN")


def excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def text_cell(sheet, value):
    # Always mark user-authored text as text, including values starting with '='.
    cell = WriteOnlyCell(sheet, value="" if value is None else str(value))
    cell.data_type = "s"
    cell.number_format = "@"
    return cell


def append_text(sheet, values) -> None:
    sheet.append([text_cell(sheet, value) for value in values])


def template(format_name: str = "center") -> tuple[io.BytesIO, str]:
    if format_name not in {"center", "java44", "python38"}:
        raise AdmError("模板格式须为center、java44或python38", 422, "ADM_TEMPLATE")
    book = Workbook(write_only=True)
    if format_name == "center":
        for title, columns in [("ADM主单", CASE_COLUMNS), ("订单明细", LINE_COLUMNS)]:
            sheet = book.create_sheet(title)
            append_text(sheet, [label for _, label in columns])
        info = book.create_sheet("填写说明")
        for values in [
            ["模板版本", "ADM_CENTER_V1"],
            ["建档说明", "主单标识用同一文件内唯一文本；明细填写对应主单标识。一张ADM可有多个出退改明细。"],
            ["业务类型", "ISSUE/出票、REFUND/退票、CHANGE/改签、UNKNOWN/待核实"],
            ["金额", "原币精确数值，建议文本格式；未知留空，不自动补0或默认币种。"],
            ["导入范围", "本版仅新单建档，不覆盖既有人工数据，不从Excel直接授予审核/结案。"],
        ]:
            append_text(info, values)
    else:
        sheet = book.create_sheet("ADM跟单")
        append_text(sheet, [label for _, label in (LEGACY44 if format_name == "java44" else LEGACY38)])
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    return content, f"ADM_{format_name}_导入模板.xlsx"


class AdmExchangeService:
    def __init__(self, config: dict):
        self.config = config
        self.root = Path(str(config.get("ADM_IMPORT_ROOT", "/app/var/adm-imports"))).resolve()
        self.maximum = int(config.get("ADM_ATTACHMENT_MAX_BYTES", 20 * 1024 * 1024))
        self.row_limit = int(config.get("ADM_IMPORT_MAX_ROWS", 10000))

    def _rows(self, sheet, columns: list[tuple[str, str]], errors: list[dict]) -> list[dict]:
        iterator = sheet.iter_rows()
        header_cells = next(iterator, [])
        headers = [excel_text(cell.value) for cell in header_cells]
        nonempty = [value for value in headers if value]
        if len(set(nonempty)) != len(nonempty):
            raise AdmError(f"{sheet.title}存在重复列名", 422, "ADM_HEADERS")
        positions = {value: index for index, value in enumerate(headers) if value}
        required = [label for key, label in columns if key in {"key", "caseKey", "admNo", "supplierName"}]
        if any(label not in positions for label in required):
            raise AdmError(f"{sheet.title}缺少必需列：{'、'.join(required)}", 422, "ADM_HEADERS")
        result = []
        for number, cells in enumerate(iterator, 2):
            if number > self.row_limit + 1:
                raise AdmError(f"单个工作表最多{self.row_limit}行，请分批上传", 413, "ADM_IMPORT_LIMIT")
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                errors.append({"row": number, "sheet": sheet.title, "message": "不接受公式，请粘贴为值后上传"})
                continue
            row = {
                key: excel_text(cells[positions[label]].value)
                if label in positions and positions[label] < len(cells)
                else ""
                for key, label in columns
            }
            row["_row"] = number
            row["_fields"] = {
                headers[index]: excel_text(cell.value)
                for index, cell in enumerate(cells)
                if index < len(headers) and headers[index]
            }
            result.append(row)
        return result

    def _validate(self, cases: list[dict], errors: list[dict]) -> None:
        validator = AdmService(self.config)
        for case in cases:
            if case.get("issuedDate") and re.match(r"^\d{4}-\d{2}-\d{2}T00:00:00$", case["issuedDate"]):
                case["issuedDate"] = case["issuedDate"][:10]
            try:
                validator._clean(case, CASE_FIELDS)
                for line in case["lines"]:
                    validator._new_line(case, line)
            except AdmError as error:
                errors.append({"row": case["source"]["row"], "message": str(error)})
            for field in ("issuedDate", "deadline"):
                if case.get(field):
                    try:
                        if field == "issuedDate" or len(case[field]) == 10:
                            date.fromisoformat(case[field])
                        else:
                            datetime.fromisoformat(case[field].replace("Z", "+00:00"))
                    except ValueError:
                        errors.append({"row": case["source"]["row"], "message": f"{field}日期格式无效"})
            for key, label in [("admNo", "ADM号"), ("supplierName", "下发方")]:
                if not case.get(key):
                    errors.append({"row": case["source"]["row"], "message": f"{label}不能为空"})
            if case.get("currency") and not re.fullmatch(r"[A-Z]{3}", case["currency"]):
                errors.append({"row": case["source"]["row"], "message": "币种须为三位大写代码"})
            if len(case["lines"]) > 1000:
                errors.append({"row": case["source"]["row"], "message": "单张ADM最多1000条初始明细，请先拆批核实"})
            for value, label in [
                (case.get("originalAmount"), "下发金额"),
                *[(line.get("amount"), "明细金额") for line in case["lines"]],
            ]:
                if value:
                    try:
                        number = Decimal(value)
                        if (
                            not number.is_finite()
                            or number < 0
                            or number > Decimal("999999999999999.9999")
                            or number.as_tuple().exponent < -4
                        ):
                            raise InvalidOperation
                    except (InvalidOperation, ValueError):
                        errors.append(
                            {"row": case["source"]["row"], "message": f"{label}应为非负精确金额，最多4位小数"}
                        )

    def preview(self, user: dict, file) -> dict:
        check_permission(user, "adm.import")
        if str(file.filename or "").lower().endswith(".json"):
            from .history import preview
            return preview(self, user, file)
        if not str(file.filename or "").lower().endswith(".xlsx"):
            raise AdmError("请上传XLSX文件", 422, "ADM_IMPORT_TYPE")
        self.root.mkdir(parents=True, exist_ok=True)
        preview_id = uuid.uuid4().hex
        source = self.root / f"{preview_id}.xlsx"
        try:
            size = 0
            digest = hashlib.sha256()
            with source.open("xb") as target:
                while chunk := file.stream.read(1024 * 1024):
                    size += len(chunk)
                    if size > self.maximum:
                        raise AdmError("文件超过导入大小限制", 413, "ADM_IMPORT_LIMIT")
                    digest.update(chunk)
                    target.write(chunk)
            validate_document(source, ".xlsx", self.maximum)
            errors: list[dict] = []
            warnings: list[str] = []
            cases = []
            book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            try:
                if "ADM主单" in book.sheetnames and "订单明细" in book.sheetnames:
                    format_name = "center"
                    rows = self._rows(book["ADM主单"], CASE_COLUMNS, errors)
                    detail_rows = self._rows(book["订单明细"], LINE_COLUMNS, errors)
                    groups = {}
                    for row in rows:
                        key = row.pop("key")
                        if not key or key in groups:
                            errors.append({"row": row["_row"], "message": "主单标识为空或重复"})
                            continue
                        case = {field: row[field] for field, _ in CASE_COLUMNS if field != "key"}
                        case["source"] = {
                            "system": "center-workbook",
                            "row": row["_row"],
                            "recordId": key,
                            "fields": row["_fields"],
                        }
                        case["lines"] = []
                        groups[key] = case
                        cases.append(case)
                    types = {"出票": "ISSUE", "退票": "REFUND", "改签": "CHANGE", "待核实": "UNKNOWN"}
                    for row in detail_rows:
                        case = groups.get(row["caseKey"])
                        if not case:
                            errors.append({"row": row["_row"], "sheet": "订单明细", "message": "主单标识未匹配"})
                            continue
                        line = {
                            key: value for key, value in row.items() if key not in {"id", "caseKey", "_row", "_fields"}
                        }
                        line["businessType"] = types.get(line["businessType"], line["businessType"] or "UNKNOWN")
                        if line["businessType"] not in {"ISSUE", "REFUND", "CHANGE", "UNKNOWN"}:
                            errors.append({"row": row["_row"], "message": "业务类型应为出/退/改或待核实"})
                        line["currency"] = line["currency"] or case["currency"]
                        line["source"] = {"row": row["_row"], "fields": row["_fields"]}
                        case["lines"].append(line)
                else:
                    sheet = book.active
                    headings = [excel_text(cell.value) for cell in next(sheet.iter_rows(), [])]
                    if "ADM单号（业务唯一标识）" not in headings:
                        raise AdmError("未识别模板，请下载中台V1、Java44或Python38模板", 422, "ADM_HEADERS")
                    columns = LEGACY44 if "申诉对象" in headings or "PNR编码" in headings else LEGACY38
                    format_name = "java44" if columns is LEGACY44 else "python38"
                    rows = self._rows(sheet, columns, errors)
                    for row in rows:
                        case = {
                            "admNo": row["adm_no"],
                            "supplierName": row["supplier_code"],
                            "supplierType": {"1":"PLATFORM","2":"AIRLINE","3":"SUPPLIER","平台":"PLATFORM","航司":"AIRLINE","供应商":"SUPPLIER"}.get(row["supplier_type_name"], "OTHER"),
                            "supplierCode": row["_fields"].get("供应商编码", ""),
                            "airline": row["airline"], "ticketCount": row["ticket_count"],
                            "originalAmountCny": row.get("total_amount_cny", ""),
                            "businessCategory": {"票务":"0","业务":"1","客服":"2"}.get(row["category_name"], row["category_name"]),
                            "currency": row["currency"],
                            "originalAmount": row["total_amount"],
                            "issuedDate": row["supply_issue_date"],
                            "deadline": row["adm_deadline"],
                            "description": row["diff_detail_reason"],
                            "remark": row["remark"],
                            "source": {
                                "system": format_name,
                                "recordId": row["id"] or f"row:{row['_row']}",
                                "row": row["_row"],
                                "fields": row["_fields"],
                            },
                            "lines": [
                                {
                                    "businessType": "UNKNOWN",
                                    "platform": row["ota_code"],
                                    "orderNo": row["ota_order_no"],
                                    "ticketNo": row["ticket_no"],
                                    "pnr": row.get("pnr_no", ""),
                                    "currency": row["currency"],
                                    "amount": row["total_amount"],
                                    "description": row["diff_detail_reason"],
                                    "differenceDescription": row["diff_detail_reason"],
                                    "diffType": row["diff_type_name"],
                                    "diffSecondType": row.get("diff_second_type", ""),
                                    "remark": row["remark"],
                                    "airline": row["airline"], "ticketCount": row["ticket_count"],
                                    "amountCny": row.get("total_amount_cny", ""),
                                    "businessCategory": {"票务":"0","业务":"1","客服":"2"}.get(row["category_name"], row["category_name"]),
                                    "actualOwner": row["actual_owner"], "actualHandler": row["actual_handler"],
                                    "recoveryDeadline": row.get("pnr_recover_time", ""), "resolution": row["resolution"],
                                    "appealReason": row["appeal_reason"],
                                    "appealDecision": {"0":"REQUIRED","1":"NOT_REQUIRED","申诉":"REQUIRED","不申诉":"NOT_REQUIRED"}.get(row["appeal_status_name"], "UNDECIDED"),
                                }
                            ],
                        }
                        cases.append(case)
                    warnings.append(
                        "旧格式按来源行分别建草稿，不按同号自动合并；历史状态/责任人/备注/附件地址原样保存，不自动认定审核、恢复或财务已完成。"
                    )
                self._validate(cases, errors)
                if not cases:
                    errors.append({"row": 0, "message": "没有可建档的主单"})
                for case in cases:
                    case["source"]["fileSha256"] = digest.hexdigest()
                    case["idempotencyKey"] = f"xlsx:{digest.hexdigest()}:{case['source']['recordId']}"
                plan = {
                    "id": preview_id,
                    "userId": str(user["id"]),
                    "createdAt": business_now_naive().isoformat(),
                    "format": format_name,
                    "filename": str(file.filename),
                    "sha256": digest.hexdigest(),
                    "size": size,
                    "cases": cases,
                    "errors": errors,
                    "warnings": warnings,
                }
                self._save_plan(plan)
                return {
                    "previewId": preview_id,
                    "format": format_name,
                    "headers": [label for _, label in (CASE_COLUMNS if format_name == "center" else columns)],
                    "rowCount": sum(len(item["lines"]) for item in cases),
                    "cases": [
                        {
                            "admNo": item["admNo"],
                            "supplierName": item["supplierName"],
                            "currency": item["currency"],
                            "originalAmount": item["originalAmount"],
                            "lineCount": len(item["lines"]),
                        }
                        for item in cases
                    ],
                    "errors": errors[:200],
                    "errorCount": len(errors),
                    "warnings": warnings,
                    "mode": "CREATE_ONLY",
                }
            finally:
                book.close()
        except Exception:
            source.unlink(missing_ok=True)
            raise

    def _save_plan(self, plan: dict) -> None:
        destination = self.root / f"{plan['id']}.json"
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=self.root, delete=False) as output:
            json.dump(plan, output, ensure_ascii=False)
            temporary = Path(output.name)
        temporary.replace(destination)

    def commit(self, user: dict, preview_id: str) -> dict:
        check_permission(user, "adm.import")
        if not re.fullmatch(r"[a-f0-9]{32}", preview_id):
            raise AdmError("导入预览不存在", 404, "ADM_IMPORT_NOT_FOUND")
        path = self.root / f"{preview_id}.json"
        if not path.is_file():
            raise AdmError("导入预览不存在或已清理", 404, "ADM_IMPORT_NOT_FOUND")
        plan = json.loads(path.read_text(encoding="utf-8"))
        if plan["userId"] != str(user["id"]):
            raise AdmError("没有此导入预览权限", 403, "ADM_FORBIDDEN")
        if datetime.fromisoformat(plan["createdAt"]) < business_now_naive() - timedelta(hours=2):
            raise AdmError("预览已过期，请重新上传", 409, "ADM_IMPORT_EXPIRED")
        if plan["errors"]:
            raise AdmError("预览有错误，请修正文件后重传", 422, "ADM_IMPORT_INVALID")
        service = AdmService(self.config)
        outcomes = []
        for item in plan["cases"]:
            try:
                values = copy.deepcopy(item)
                if plan.get("mode") == "HISTORY_ARCHIVE":
                    check_permission(user, "adm.import")
                    if not user.get("is_admin"):
                        raise AdmError("历史归档仅限管理员", 403)
                    history = values.pop("history")
                    detail = service.create_case(user, values, history=history)
                else:
                    detail = service.import_case(user, values)
                outcomes.append(
                    {
                        "admNo": detail["admNo"],
                        "id": detail["id"],
                        "status": "CREATED_OR_REUSED",
                        "version": detail["version"],
                    }
                )
            except AdmError as error:
                outcomes.append({"admNo": item["admNo"], "status": "FAILED", "message": str(error), "code": error.code})
        return {
            "previewId": preview_id,
            "items": outcomes,
            "successCount": sum(item["status"] != "FAILED" for item in outcomes),
            "failedCount": sum(item["status"] == "FAILED" for item in outcomes),
            "notice": "历史JSON按ERP旧id全局防重，只读归档不进入活跃待办；变更冲突不会覆盖。" if plan.get("mode") == "HISTORY_ARCHIVE" else "按主单原子建档，重复提交复用来源幂等键；导入文件尚未作为下发原件，请在主单附件补充真实供应原件。",
        }


def export_workbook(cases: list[dict], kind: str = "adm_cases") -> tuple[io.BytesIO, str]:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("ADM主单" if kind == "adm_cases" else "订单明细")
    columns = (
        CASE_COLUMNS
        if kind == "adm_cases"
        else [
            *(item for item in CASE_COLUMNS if item[0] in {"admNo", "supplierName"}),
            *LINE_COLUMNS,
            ("status", "处理状态"),
            ("ownerName", "当前处理人"),
        ]
    )
    append_text(sheet, [label for _, label in columns])
    for case in cases:
        if kind == "adm_cases":
            append_text(sheet, [case.get("id") if key == "key" else case.get(key) for key, _ in columns])
        else:
            for line in case["lines"]:
                values = {**line, "caseKey": case["id"], "admNo": case["admNo"], "supplierName": case["supplierName"]}
                append_text(sheet, [values.get(key) for key, _ in columns])
    content = io.BytesIO()
    if any(case.get("history") for case in cases):
        historical = book.create_sheet("历史原始字段")
        append_text(historical, ["中台主单ID", "ERP旧ID", "原字段", "原值"])
        for case in cases:
            for key, value in case.get("history", {}).get("fields", {}).items():
                text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
                if len(text) > 32000:
                    raise AdmError("历史原始字段过长，不能无损导出Excel，请缩小范围并核对原始JSON", 422)
                append_text(historical, [case["id"], case["history"]["oldId"], key, text])
    book.save(content)
    content.seek(0)
    return content, f"ADM_{'主单' if kind == 'adm_cases' else '订单明细'}_{business_now_naive():%Y%m%d_%H%M}.xlsx"


def select_export_cases(config: dict, user: dict, filters: dict) -> tuple[list[dict], list[dict]]:
    """Freeze authorized IDs and versions; never trust selection metadata from callers."""
    check_permission(user, "adm.export")
    service = AdmService(config)
    if filters.get("caseId"):
        cases = [service.case_detail(user, str(filters["caseId"]))]
    else:
        cases = []
        page = 1
        while True:
            result = service.list_cases(user, {**filters, "page": page, "pageSize": 100})
            if result["pagination"]["total"] > 10000:
                raise AdmError("单次最多导出10000张ADM，请缩小筛选范围", 422, "ADM_EXPORT_LIMIT")
            cases.extend(service.case_detail(user, item["id"]) for item in result["items"])
            if page * 100 >= result["pagination"]["total"]:
                break
            page += 1
    if sum(len(item["lines"]) for item in cases) > 10000:
        raise AdmError("单次最多导出10000条明细，请缩小筛选范围", 422, "ADM_EXPORT_LIMIT")
    if not cases:
        raise AdmError("当前条件没有可导出数据", 422, "ADM_EXPORT_EMPTY")
    selection = [
        {
            "id": item["id"],
            "version": item["version"],
            "lineIds": [line["id"] for line in item["lines"]],
            "partial": bool(item.get("partial")),
        }
        for item in cases
    ]
    return cases, selection


def validate_export_cases(config: dict, user: dict, selection: list[dict]) -> list[dict]:
    check_permission(user, "adm.export")
    service = AdmService(config)
    cases = []
    for item in selection:
        current = service.case_detail(user, item["id"])
        if (
            current["version"] != item["version"]
            or bool(current.get("partial")) != item["partial"]
            or {line["id"] for line in current["lines"]} != set(item["lineIds"])
        ):
            raise AdmError("导出范围、权限或数据已变化，请重新创建任务", 409, "ADM_EXPORT_CHANGED")
        cases.append(current)
    return cases
