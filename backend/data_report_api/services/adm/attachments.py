"""Persistent, permission checked ADM evidence; never stores public file URLs."""

from __future__ import annotations

import csv
import hashlib
import io
import mimetypes
import re
import tempfile
import uuid
import warnings
import zipfile
from pathlib import Path
from typing import Any

from ...time_utils import business_now_naive
from . import AdmError, AdmService


FILE_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".eml": "message/rfc822",
}
CATEGORIES = {"ORIGINAL", "VERIFICATION", "RECOVERY", "APPEAL", "FINANCE", "ANALYSIS"}
STORAGE_KEY = re.compile(r"^[a-f0-9]{32}\.(?:pdf|png|jpe?g|gif|webp|xlsx|docx|csv|txt|eml)$")


def clean_filename(value: str) -> str:
    name = str(value or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not name or len(name) > 240 or any(ord(char) < 32 for char in name):
        raise AdmError("附件名称不能为空、包含控制字符或超过240字", 422, "ADM_FILE_NAME")
    return name


def validate_document(path: Path, extension: str, maximum: int) -> None:
    """Check signatures and expansion budgets before any preview parser sees a file."""
    with path.open("rb") as source:
        prefix = source.read(32)
    checks = {
        ".pdf": prefix.startswith(b"%PDF-"),
        ".png": prefix.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": prefix.startswith(b"\xff\xd8\xff"),
        ".jpeg": prefix.startswith(b"\xff\xd8\xff"),
        ".gif": prefix.startswith((b"GIF87a", b"GIF89a")),
        ".webp": prefix.startswith(b"RIFF") and prefix[8:12] == b"WEBP",
    }
    if extension in checks and not checks[extension]:
        raise AdmError("附件内容与扩展名不一致", 422, "ADM_FILE_SIGNATURE")
    if extension == ".pdf":
        from pypdf import PdfReader

        try:
            with path.open("rb") as source:
                reader = PdfReader(source, strict=True)
                if reader.is_encrypted:
                    raise AdmError("请先解密PDF后上传", 422, "ADM_FILE_ENCRYPTED")
                if not 0 < len(reader.pages) <= 2000:
                    raise AdmError("PDF为空或页数超过2000页", 422, "ADM_FILE_PAGES")
                # Force access to the page tree; a matching header is not a readable PDF.
                for page in reader.pages:
                    _ = page.mediabox
        except AdmError:
            raise
        except Exception as error:
            raise AdmError("PDF损坏，无法读取页面", 422, "ADM_FILE_CORRUPT") from error
    if extension in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        from PIL import Image

        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(path) as picture:
                    if picture.width * picture.height > 40_000_000:
                        raise AdmError("图片像素超过安全预览上限", 413, "ADM_FILE_PIXELS")
                    picture.verify()
        except AdmError:
            raise
        except Exception as error:
            raise AdmError("图片损坏，无法读取", 422, "ADM_FILE_CORRUPT") from error
    if extension in {".xlsx", ".docx"}:
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 5000 or sum(item.file_size for item in entries) > maximum * 20:
                    raise AdmError("附件解压内容超出安全预览范围", 413, "ADM_FILE_EXPANSION")
                if any(item.flag_bits & 1 for item in entries):
                    raise AdmError("不支持加密附件，请先解密后上传", 422, "ADM_FILE_ENCRYPTED")
                names = {item.filename for item in entries}
                required = "xl/workbook.xml" if extension == ".xlsx" else "word/document.xml"
                if required not in names or "[Content_Types].xml" not in names:
                    raise AdmError("文件不是有效的Office文档", 422, "ADM_FILE_SIGNATURE")
                if any("vbaproject" in name.lower() for name in names):
                    raise AdmError("不接受包含宏的附件", 422, "ADM_FILE_MACRO")
                if archive.testzip() is not None:
                    raise AdmError("附件压缩数据校验失败", 422, "ADM_FILE_CORRUPT")
        except (zipfile.BadZipFile, OSError) as error:
            raise AdmError("附件损坏，无法读取", 422, "ADM_FILE_CORRUPT") from error
    if extension in {".txt", ".csv", ".eml"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeError as error:
            raise AdmError("文本附件请使用UTF-8编码", 422, "ADM_FILE_ENCODING") from error
        if "\x00" in text:
            raise AdmError("文本附件包含非法二进制内容", 422, "ADM_FILE_CONTENT")


class AdmAttachmentService:
    def __init__(self, config: dict[str, Any], service: AdmService | None = None):
        self.config = config
        self.service = service or AdmService(config)
        self.maximum = int(config.get("ADM_ATTACHMENT_MAX_BYTES", 20 * 1024 * 1024))
        self.root = Path(str(config.get("ADM_ATTACHMENT_ROOT", "/app/var/adm-attachments"))).resolve()

    def upload(self, user: dict, case_id: str, file: Any, body: dict) -> dict:
        # Reject unauthorized objects before touching disk. Mutation rechecks permissions and CAS.
        self.service.case_detail(user, case_id)
        name = clean_filename(file.filename)
        extension = Path(name).suffix.lower()
        if extension not in FILE_TYPES:
            raise AdmError("支持PDF、图片、XLSX、DOCX、CSV、TXT及EML附件", 422, "ADM_FILE_TYPE")
        category = str(body.get("category", "VERIFICATION")).upper()
        if category not in CATEGORIES:
            raise AdmError("附件分类无效", 422, "ADM_FILE_CATEGORY")
        try:
            version = int(body.get("version", 0))
        except (TypeError, ValueError) as error:
            raise AdmError("请提供有效的主单版本", 422, "ADM_VERSION") from error
        if version < 1:
            raise AdmError("请刷新主单后重新上传", 409, "ADM_VERSION_CONFLICT")
        self.root.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        final_path: Path | None = None
        try:
            digest = hashlib.sha256()
            size = 0
            with tempfile.NamedTemporaryFile(dir=self.root, prefix=".upload-", delete=False) as target:
                temporary = Path(target.name)
                while chunk := file.stream.read(1024 * 1024):
                    size += len(chunk)
                    if size > self.maximum:
                        raise AdmError("附件超过配置的单文件大小限制", 413, "ADM_FILE_LIMIT")
                    digest.update(chunk)
                    target.write(chunk)
            if size == 0:
                raise AdmError("不能上传空附件", 422, "ADM_FILE_EMPTY")
            validate_document(temporary, extension, self.maximum)
            storage_key = f"{uuid.uuid4().hex}{extension}"
            final_path = self.root / storage_key
            temporary.replace(final_path)
            temporary = None
            metadata = {
                "id": uuid.uuid4().hex,
                "name": name,
                "category": category,
                "lineId": str(body.get("lineId") or "") or None,
                "mime": FILE_TYPES[extension],
                "size": size,
                "sha256": digest.hexdigest(),
                "storageKey": storage_key,
                "createdAt": business_now_naive().isoformat(timespec="seconds"),
            }
            result = self.service.add_attachment(user, case_id, {"version": version, "attachment": metadata})
            final_path = None  # Committed evidence is never removed by temporary-upload cleanup.
            return result
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
            if final_path is not None:
                final_path.unlink(missing_ok=True)

    def file(self, user: dict, case_id: str, attachment_id: str) -> tuple[Path, dict]:
        metadata = self.service.attachment(user, case_id, attachment_id)
        key = str(metadata.get("storageKey", ""))
        if not STORAGE_KEY.fullmatch(key):
            raise AdmError("附件存储信息异常，请联系管理员", 503, "ADM_FILE_STORAGE")
        path = self.root / key
        if path.parent != self.root or path.is_symlink() or not path.is_file():
            raise AdmError("附件原件缺失，请补传或联系管理员", 404, "ADM_FILE_MISSING")
        return path, metadata

    def table_preview(self, path: Path, *, sheet: str = "") -> dict:
        maximum_rows, maximum_columns = 200, 50
        if path.suffix == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as source:
                iterator = csv.reader(source)
                rows = []
                for row in iterator:
                    rows.append([str(value)[:5000] for value in row[:maximum_columns]])
                    if len(rows) > maximum_rows + 1:
                        break
            names = ["CSV"]
            selected = "CSV"
        else:
            from openpyxl import load_workbook

            try:
                book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                try:
                    names = book.sheetnames
                    if sheet and sheet not in names:
                        raise AdmError("指定工作表不存在", 422, "ADM_FILE_SHEET")
                    selected = sheet or names[0]
                    worksheet = book[selected]
                    rows = []
                    for row in worksheet.iter_rows(max_row=maximum_rows + 2, max_col=maximum_columns, values_only=True):
                        rows.append(["" if value is None else str(value)[:5000] for value in row])
                finally:
                    book.close()
            except AdmError:
                raise
            except Exception as error:
                raise AdmError("表格损坏，无法预览；可下载原件核查", 422, "ADM_FILE_CORRUPT") from error
        columns = rows[0] if rows else []
        return {
            "type": "table",
            "sheets": names,
            "sheet": selected,
            "columns": columns,
            "rows": rows[1 : maximum_rows + 1],
            "truncated": len(rows) > maximum_rows + 1 or len(columns) >= maximum_columns,
            "totalPreviewRows": min(maximum_rows, max(0, len(rows) - 1)),
            "notice": "只读预览，最多200行、50列；完整材料请下载原件",
        }

    def preview(self, user: dict, case_id: str, attachment_id: str, sheet: str = "") -> tuple[str, Any, dict]:
        path, metadata = self.file(user, case_id, attachment_id)
        if path.suffix in {".xlsx", ".csv"}:
            return "json", self.table_preview(path, sheet=sheet), metadata
        if metadata.get("mime") == "application/pdf" or str(metadata.get("mime", "")).startswith("image/"):
            return "file", path, metadata
        if path.suffix in {".txt", ".eml"}:
            text = path.read_text(encoding="utf-8-sig")
            return "json", {"type": "text", "text": text[:100000], "truncated": len(text) > 100000}, metadata
        return "json", {"type": "unsupported", "message": "此格式暂不支持在线预览，请下载原件"}, metadata


def file_mime(metadata: dict) -> str:
    return str(metadata.get("mime") or mimetypes.guess_type(str(metadata.get("name")))[0] or "application/octet-stream")


def verify_stored_attachment(config: dict, metadata: dict) -> bool:
    """Recheck evidence existence and integrity at closure, not only at upload time."""
    key = str(metadata.get("storageKey", ""))
    if not STORAGE_KEY.fullmatch(key):
        return False
    root = Path(str(config.get("ADM_ATTACHMENT_ROOT", "/app/var/adm-attachments"))).resolve()
    path = root / key
    try:
        if path.is_symlink() or not path.is_file() or path.stat().st_size != int(metadata.get("size", -1)):
            return False
        digest = hashlib.sha256()
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                digest.update(chunk)
        return digest.hexdigest() == metadata.get("sha256")
    except (OSError, TypeError, ValueError):
        return False
