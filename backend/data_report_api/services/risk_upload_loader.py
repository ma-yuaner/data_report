from __future__ import annotations

import datetime as dt
import re
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable

from .data_source import DataSource
from .risk_upload_definitions import RiskUploadDefinition


MAX_SQL_BYTES = 4 * 1024 * 1024


def normalize_type(value: str) -> str:
    return re.sub(r"\s+", "", value.lower())


def hive_string(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace("'", "\\'")
        .replace("\n", "\\n")
        .replace("\r", "\\r")
        .replace("\t", "\\t")
        .replace("\x00", "\\0")
    )
    return f"'{escaped}'"


def normalize_value(value: Any, target_type: str) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        text = "true" if value else "false"
    elif isinstance(value, dt.datetime):
        text = value.isoformat(sep=" ", timespec="seconds")
    elif isinstance(value, (dt.date, dt.time)):
        text = value.isoformat()
    else:
        text = str(value)

    data_type = normalize_type(target_type)
    if data_type == "string" or data_type.startswith("varchar(") or data_type.startswith("char("):
        return text
    stripped = text.strip().replace(",", "")
    if data_type in {"tinyint", "smallint", "int", "bigint"}:
        try:
            number = Decimal(stripped)
            if not number.is_finite() or number != number.to_integral_value():
                raise ValueError("非整数")
            return str(int(number))
        except (InvalidOperation, ValueError, OverflowError) as error:
            raise ValueError(f"无法将{value!r}转换为{data_type}") from error
    if data_type.startswith("decimal("):
        try:
            number = Decimal(stripped)
            precision, scale = (int(item) for item in re.findall(r"\d+", data_type))
            if not number.is_finite() or abs(number) >= Decimal(10) ** (precision - scale):
                raise ValueError("超出范围")
            return format(number, "f")
        except (InvalidOperation, ValueError) as error:
            raise ValueError(f"无法将{value!r}转换为{data_type}") from error
    if data_type in {"float", "double"}:
        try:
            number = Decimal(stripped)
            if not number.is_finite():
                raise ValueError("非有限数")
            return format(number, "f")
        except (InvalidOperation, ValueError) as error:
            raise ValueError(f"无法将{value!r}转换为{data_type}") from error
    if data_type == "boolean":
        lowered = stripped.lower()
        if lowered not in {"true", "false", "1", "0"}:
            raise ValueError(f"无法将{value!r}转换为boolean")
        return "true" if lowered in {"true", "1"} else "false"
    if data_type == "date":
        return text[:10]
    if data_type == "timestamp":
        return text
    raise ValueError(f"暂不支持Hive字段类型：{target_type}")


def hive_literal(value: Any, target_type: str) -> str:
    normalized = normalize_value(value, target_type)
    if normalized is None:
        return "NULL"
    data_type = normalize_type(target_type)
    if data_type == "string" or data_type.startswith("varchar(") or data_type.startswith("char("):
        return hive_string(normalized)
    if data_type in {"date", "timestamp"}:
        return f"CAST({hive_string(normalized)} AS {data_type.upper()})"
    return normalized


def trim_headers(values: Iterable[Any]) -> list[str]:
    headers = list(values)
    while headers and (headers[-1] is None or not str(headers[-1]).strip()):
        headers.pop()
    if not headers:
        raise ValueError("Excel首行没有有效表头")
    result: list[str] = []
    for index, value in enumerate(headers, start=1):
        if value is None or not str(value).strip():
            raise ValueError(f"Excel表头第{index}列为空，无法安全匹配字段")
        result.append(str(value).strip())
    duplicates = sorted({name for name in result if result.count(name) > 1})
    if duplicates:
        raise ValueError(f"Excel表头存在重复字段：{duplicates}")
    return result


def describe_target(cursor: Any, target_table: str) -> tuple[list[tuple[str, str]], bool]:
    cursor.execute(f"DESCRIBE {target_table}")
    columns: list[tuple[str, str]] = []
    partitioned = False
    partition_section = False
    for row in cursor.fetchall():
        name = str(row[0]).strip().lower() if row and row[0] is not None else ""
        if name.startswith("# partition"):
            partitioned = True
            partition_section = True
            continue
        if partition_section or not name or name.startswith("#"):
            continue
        columns.append((name, normalize_type(str(row[1]).strip())))
    return columns, partitioned


def repair_refund_row(values: list[Any], header_index: dict[str, int]) -> list[Any]:
    pcc_index = header_index.get("订位PCC")
    ticket_index = header_index.get("票数")
    if pcc_index is None or ticket_index is None:
        return values
    value = values[pcc_index] if pcc_index < len(values) else None
    looks_like_date = isinstance(value, (dt.date, dt.datetime)) or (
        isinstance(value, str)
        and re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", value.strip()) is not None
    )
    if not looks_like_date:
        return values
    repaired = list(values)
    repaired.insert(pcc_index, None)
    if ticket_index >= len(repaired):
        raise ValueError("旧格式退票行缺少票数和月份，无法安全修复")
    repaired.pop(ticket_index)
    return repaired


def flush_rows(cursor: Any, connection: Any, staging_table: str,
               rows: list[str]) -> int:
    if not rows:
        return 0
    cursor.execute(f"INSERT INTO TABLE {staging_table} VALUES\n" + ",\n".join(rows))
    connection.commit()
    count = len(rows)
    rows.clear()
    return count


def ticket_key_stats(cursor: Any, table: str) -> tuple[int, int, int]:
    """Return total, non-null and distinct normalized issue ticket keys."""
    cursor.execute(
        "SELECT COUNT(1), "
        "COUNT(CASE WHEN `issue_ticket_no` IS NOT NULL "
        "AND TRIM(`issue_ticket_no`)<>'' "
        "THEN 1 END), "
        "COUNT(DISTINCT CASE WHEN `issue_ticket_no` IS NOT NULL "
        "AND TRIM(`issue_ticket_no`)<>'' THEN TRIM(`issue_ticket_no`) END) "
        f"FROM {table}"
    )
    total, non_null, distinct_count = cursor.fetchone()
    return int(total), int(non_null), int(distinct_count)


def load_excel_to_hive(
    *,
    definition: RiskUploadDefinition,
    excel_path: str | Path,
    sheet_name: str,
    load_date: str | None,
    config: dict[str, Any],
    log: Callable[[str], None],
) -> tuple[int, int]:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("导入Worker缺少openpyxl依赖") from error

    path = Path(excel_path).resolve()
    if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("上传文件不存在或格式不是.xlsx/.xlsm")
    workbook = openpyxl.load_workbook(
        path, read_only=True, data_only=True, keep_links=False
    )
    connection = None
    cursor = None
    staging_table = ""
    merge_table = ""
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"工作表不存在：{sheet_name}；可选工作表={workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        header_rows = worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
        try:
            headers = trim_headers(next(header_rows))
        except StopIteration as error:
            raise ValueError("Excel工作表为空") from error
        header_index = {name: index for index, name in enumerate(headers)}
        required_headers = [source for source, _ in definition.source_to_target]
        missing = [name for name in required_headers if name not in header_index]
        if missing:
            raise ValueError(f"Excel缺少目标表所需字段：{missing}")
        ignored = [name for name in headers if name not in required_headers]
        if ignored:
            log(f"忽略{len(ignored)}个非目标字段")

        source_indexes = [header_index[name] for name in required_headers]
        expected_columns = [target for _, target in definition.source_to_target]
        source = DataSource({**config, "DATA_MODE": "hive"})
        connection = source.connect()
        cursor = connection.cursor()
        actual_schema, partitioned = describe_target(cursor, definition.target_table)
        actual_columns = [name for name, _ in actual_schema]
        if actual_columns != expected_columns:
            first = next(
                (
                    index
                    for index, (actual, expected) in enumerate(
                        zip(actual_columns, expected_columns), start=1
                    )
                    if actual != expected
                ),
                min(len(actual_columns), len(expected_columns)) + 1,
            )
            raise RuntimeError(
                f"Hive目标表字段不一致：期望{len(expected_columns)}列，"
                f"实际{len(actual_columns)}列，首个差异在第{first}列"
            )
        if partitioned and not definition.partition_date_required:
            raise RuntimeError("目标表意外成为分区表，当前上传口径按非分区整表设计，已停止覆盖")
        if partitioned and not load_date:
            raise ValueError("分区表导入必须提供dt日期")
        log(f"Excel表头与Hive目标表{len(actual_schema)}列校验通过")

        database = definition.target_table.split(".", 1)[0]
        staging_table = f"{database}.tmp_risk_upload_{uuid.uuid4().hex}"
        column_ddl = ",\n".join(
            f"  `{name}` {data_type.upper()}" for name, data_type in actual_schema
        )
        cursor.execute(
            f"CREATE TEMPORARY TABLE {staging_table} (\n{column_ddl}\n) STORED AS ORC"
        )

        batch_size = max(1, min(int(config.get("RISK_UPLOAD_INSERT_BATCH_SIZE", 2000)), 5000))
        prefix_bytes = len(f"INSERT INTO TABLE {staging_table} VALUES\n".encode("utf-8"))
        sql_bytes = prefix_bytes
        pending: list[str] = []
        total = 0
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=len(headers), values_only=True),
            start=2,
        ):
            row_values = list(values)
            if definition.key == "refund":
                row_values = repair_refund_row(row_values, header_index)
            selected = [
                row_values[index] if index < len(row_values) else None
                for index in source_indexes
            ]
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in selected):
                continue
            try:
                row_sql = "(" + ", ".join(
                    hive_literal(value, data_type)
                    for value, (_, data_type) in zip(selected, actual_schema)
                ) + ")"
            except ValueError as error:
                raise ValueError(f"Excel第{row_number}行：{error}") from error
            row_bytes = len(row_sql.encode("utf-8"))
            if prefix_bytes + row_bytes > MAX_SQL_BYTES:
                raise ValueError(f"Excel第{row_number}行内容过大，已停止写入")
            separator_bytes = 2 if pending else 0
            if pending and (
                len(pending) >= batch_size
                or sql_bytes + separator_bytes + row_bytes > MAX_SQL_BYTES
            ):
                total += flush_rows(cursor, connection, staging_table, pending)
                log(f"Hive临时表已写入{total:,}行")
                sql_bytes = prefix_bytes
                separator_bytes = 0
            pending.append(row_sql)
            sql_bytes += separator_bytes + row_bytes
        total += flush_rows(cursor, connection, staging_table, pending)
        if total <= 0:
            raise ValueError("Excel没有有效数据行，禁止覆盖Hive目标表")
        log(f"Excel读取完成，共{total:,}行")

        cursor.execute(f"SELECT COUNT(1) FROM {staging_table}")
        staging_rows = int(cursor.fetchone()[0])
        if staging_rows != total:
            raise RuntimeError(
                f"Hive临时表行数不一致：Excel {total}行，临时表{staging_rows}行"
            )
        selected_columns = ", ".join(f"`{name}`" for name in expected_columns)
        stage_total, stage_non_null, stage_distinct = ticket_key_stats(
            cursor, staging_table
        )
        if stage_non_null != stage_total:
            raise ValueError(
                f"{definition.label}增量存在{stage_total - stage_non_null:,}行空出票票号，禁止合并"
            )
        if stage_distinct != stage_total:
            raise ValueError(
                f"{definition.label}增量存在{stage_total - stage_distinct:,}个重复出票票号，禁止合并"
            )

        target_total, target_non_null, target_distinct = ticket_key_stats(
            cursor, definition.target_table
        )
        invalid_target_rows = target_total - target_non_null
        if target_distinct != target_non_null:
            raise RuntimeError(
                f"Hive{definition.label}原表存在{target_non_null - target_distinct:,}个重复出票票号，"
                "请先清理原表后再增量导入"
            )

        cursor.execute(
            "SELECT "
            "SUM(CASE WHEN `estimated_profit_cny`=0 THEN 1 ELSE 0 END), "
            "SUM(CASE WHEN `estimated_profit_cny` IS NULL "
            "OR `estimated_profit_cny`<>0 THEN 1 ELSE 0 END) "
            f"FROM {staging_table}"
        )
        delete_requests_raw, upsert_rows_raw = cursor.fetchone()
        delete_requests = int(delete_requests_raw or 0)
        upsert_rows = int(upsert_rows_raw or 0)

        cursor.execute(
            "SELECT "
            "SUM(CASE WHEN new_rows.`estimated_profit_cny`=0 THEN 1 ELSE 0 END), "
            "SUM(CASE WHEN new_rows.`estimated_profit_cny` IS NULL "
            "OR new_rows.`estimated_profit_cny`<>0 THEN 1 ELSE 0 END) "
            f"FROM {definition.target_table} old_rows "
            f"JOIN {staging_table} new_rows "
            "ON TRIM(old_rows.`issue_ticket_no`)="
            "TRIM(new_rows.`issue_ticket_no`)"
        )
        deleted_rows_raw, replaced_rows_raw = cursor.fetchone()
        deleted_rows = int(deleted_rows_raw or 0)
        replaced_rows = int(replaced_rows_raw or 0)
        added_rows = upsert_rows - replaced_rows
        unmatched_delete_rows = delete_requests - deleted_rows
        expected_rows = (
            target_total - invalid_target_rows - deleted_rows
            - replaced_rows + upsert_rows
        )
        log(
            f"{definition.label}票号校验通过：原表{target_total:,}行，"
            f"新增{added_rows:,}行，整行替换{replaced_rows:,}行，"
            f"删除{deleted_rows:,}行，清理空票号{invalid_target_rows:,}行，"
            f"未命中删除指令{unmatched_delete_rows:,}行"
        )

        merge_table = f"{database}.tmp_risk_merge_{uuid.uuid4().hex}"
        cursor.execute(
            f"CREATE TEMPORARY TABLE {merge_table} (\n{column_ddl}\n) STORED AS ORC"
        )
        old_columns = ", ".join(
            f"old_rows.`{name}`" for name in expected_columns
        )
        new_columns = ", ".join(
            f"new_rows.`{name}`" for name in expected_columns
        )
        cursor.execute(
            f"INSERT OVERWRITE TABLE {merge_table}\n"
            f"SELECT {old_columns}\n"
            f"FROM {definition.target_table} old_rows\n"
            f"LEFT JOIN {staging_table} new_keys\n"
            "ON TRIM(old_rows.`issue_ticket_no`)="
            "TRIM(new_keys.`issue_ticket_no`)\n"
            "WHERE new_keys.`issue_ticket_no` IS NULL\n"
            "AND old_rows.`issue_ticket_no` IS NOT NULL\n"
            "AND TRIM(old_rows.`issue_ticket_no`)<>''\n"
            "UNION ALL\n"
            f"SELECT {new_columns}\nFROM {staging_table} new_rows\n"
            "WHERE new_rows.`estimated_profit_cny` IS NULL "
            "OR new_rows.`estimated_profit_cny`<>0"
        )
        connection.commit()
        merged_total, merged_non_null, merged_distinct = ticket_key_stats(
            cursor, merge_table
        )
        if (
            merged_total != expected_rows
            or merged_non_null != merged_total
            or merged_distinct != merged_total
        ):
            raise RuntimeError(
                f"{definition.label}合并临时表校验失败："
                f"计划{expected_rows:,}行，实际{merged_total:,}行，"
                f"唯一票号{merged_distinct:,}个"
            )
        log(f"{definition.label}合并临时表校验通过，共{merged_total:,}行，开始重写正式表")
        cursor.execute(
            f"INSERT OVERWRITE TABLE {definition.target_table}\n"
            f"SELECT {selected_columns}\nFROM {merge_table}"
        )
        connection.commit()
        target_rows, target_non_null, target_distinct = ticket_key_stats(
            cursor, definition.target_table
        )
        if (
            target_rows != expected_rows
            or target_non_null != target_rows
            or target_distinct != target_rows
        ):
            raise RuntimeError(
                f"{definition.label}正式表校验失败："
                f"计划{expected_rows:,}行，实际{target_rows:,}行，"
                f"唯一票号{target_distinct:,}个"
            )
        log(
            f"{definition.label}增量导入完成：新增{added_rows:,}行，"
            f"替换{replaced_rows:,}行，删除{deleted_rows:,}行，"
            f"清理空票号{invalid_target_rows:,}行，"
            f"正式表共{target_rows:,}行"
        )
        return total, target_rows
    finally:
        if cursor is not None:
            for temporary_table in (merge_table, staging_table):
                if not temporary_table:
                    continue
                try:
                    cursor.execute(f"DROP TABLE IF EXISTS {temporary_table}")
                except Exception:
                    pass
            try:
                cursor.close()
            except Exception:
                pass
        if connection is not None:
            try:
                connection.close()
            except Exception:
                pass
        workbook.close()
