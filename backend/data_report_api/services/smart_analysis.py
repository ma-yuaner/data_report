"""MySQL-backed smart ticketing overview, dimensions, and ticket detail."""
from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .data_source import DataSource


DATABASE = "sibebid"
DIMENSION_TABLE = "bi_smart_issue_dimension_year"
DETAIL_TABLE = "bi_smart_issue_detail_year"
MAX_RANGE_DAYS = 366
MAX_PAGE_SIZE = 100
MAX_OPTION_ROWS = 200
MAX_DETAIL_EXPORT_ROWS = 50_000
NULL_TOKEN = "__NULL__"

TREND_DATES: dict[str, tuple[str, str]] = {
    "issue": ("issue_ticket_date", "出票日期"),
    "departure": ("first_dep_date", "首段起飞日期"),
    "order": ("order_create_date", "进单日期"),
}

DIMENSIONS: dict[str, tuple[str, str]] = {
    "airline": ("air_line", "航司"),
    "platform": ("ota_cname", "OTA平台"),
    "site": ("ota_site_cname", "OTA站点"),
    "policy": ("policy_name", "政策员"),
    "leader": ("policy_operator", "组长"),
    "member": ("policy_operator_group_name", "组员"),
    "supplier": ("issue_supplier_cname", "出票供应商"),
    "gds": ("issue_gds_code", "GDS系统"),
    "productType": ("product_type", "产品类型"),
    "depAirport": ("dep_airport", "出发机场"),
    "arrAirport": ("arr_airport", "到达机场"),
    "route": ("airroute_code", "航线"),
    "depCity": ("dep_city_cname", "出发城市"),
    "arrCity": ("arr_city_cname", "到达城市"),
    "depCountry": ("dep_country_cname", "出发国家"),
    "arrCountry": ("arr_country_cname", "到达国家"),
    "depZone": ("dep_zone_cname", "出发大区"),
    "arrZone": ("arr_zone_cname", "到达大区"),
    "depArea": ("dep_area_cname", "出发TC区"),
    "arrArea": ("arr_area_cname", "到达TC区"),
    "flightWay": ("flight_way", "航程类型"),
    "transferType": ("transfer_type", "航段类型"),
    "orderSource": ("manual_flag", "订单来源"),
    "routeCategory": ("airroute_flag", "航线类别"),
    "ageType": ("age_type", "乘客类型"),
    "ndc": ("ndc_flag", "NDC标识"),
    "routeLength": ("route_length_type", "长短航线"),
    "cityRoute": ("city_route", "城市航线"),
    "countryRoute": ("country_route", "国家航线"),
    "zoneRoute": ("zone_route", "大区航线"),
    "tcRoute": ("tc_route", "TC区航线"),
    "profitBand": ("profit_band_name", "利润区间"),
}

AGGREGATES: tuple[tuple[str, str, str], ...] = (
    ("ticket_count", "issue_count", "count"),
    ("total_profit", "issue_profit_sum", "money"),
    ("segment_count", "segment_count_sum", "count"),
    ("refund_count", "refund_issue_count", "count"),
    ("smart_participated", "smart_participated_issue_count", "count"),
    ("smart_succeeded", "smart_succeeded_issue_count", "count"),
    ("smart_used", "smart_used_issue_count", "count"),
    ("smart_profit", "smart_profit_sum", "money"),
    ("down_participated", "down_participated_issue_count", "count"),
    ("down_succeeded", "down_succeeded_issue_count", "count"),
    ("down_used", "down_used_issue_count", "count"),
    ("down_profit", "down_profit_sum", "money"),
    ("void_participated", "void_participated_issue_count", "count"),
    ("void_succeeded", "void_succeeded_issue_count", "count"),
    ("void_used", "void_used_issue_count", "count"),
    ("void_profit", "void_profit_sum", "money"),
)
AGGREGATE_NAMES = ("source_rows", *(item[0] for item in AGGREGATES), "updated_at")
SMART_PREFIXES = {"all": "smart", "down": "down", "void": "void"}

DIMENSION_EXPORT_METRICS: dict[str, tuple[str, tuple[str, ...], str]] = {
    "ticketCount": ("出票量", ("ticketCount",), "count"),
    "totalProfit": ("全部预估利润", ("totalProfit",), "money"),
    "segmentCount": ("总航段数", ("segmentCount",), "count"),
    "refundCount": ("退票中/已退票量", ("refundCount",), "count"),
    "perSegmentProfit": ("单航段利润", ("perSegmentProfit",), "money"),
}
for _prefix, _label in (("smart", "智能整体"), ("down", "智能降舱"), ("void", "智能废票")):
    DIMENSION_EXPORT_METRICS.update({
        f"{_prefix}Participated": (f"{_label}参与出票量", (_prefix, "participatedCount"), "count"),
        f"{_prefix}Succeeded": (f"{_label}成功出票量", (_prefix, "succeededCount"), "count"),
        f"{_prefix}Used": (f"{_label}使用出票量", (_prefix, "usedCount"), "count"),
        f"{_prefix}ParticipationRate": (f"{_label}参与率", (_prefix, "participationRate"), "rate"),
        f"{_prefix}SuccessRate": (f"{_label}成功率", (_prefix, "successRate"), "rate"),
        f"{_prefix}AdoptionRate": (f"{_label}成功使用率", (_prefix, "adoptionRate"), "rate"),
        f"{_prefix}ParticipationUseRate": (f"{_label}参与使用率", (_prefix, "participationUseRate"), "rate"),
        f"{_prefix}UseRate": (f"{_label}使用率", (_prefix, "useRate"), "rate"),
        f"{_prefix}Profit": (f"{_label}利润", (_prefix, "profit"), "money"),
        f"{_prefix}AverageProfit": (f"{_label}单张利润", (_prefix, "averageProfit"), "money"),
    })


@dataclass(frozen=True)
class DetailField:
    key: str
    title: str
    expression: str
    value_type: str = "text"
    category: str = "订单与出票"
    width: int = 140


def detail_field(
    key: str, title: str, field: str, value_type: str = "text",
    category: str = "订单与出票", width: int = 140, *, raw: bool = False,
) -> DetailField:
    return DetailField(key, title, field if raw else f"src.{field}", value_type, category, width)


DETAIL_FIELDS: tuple[DetailField, ...] = (
    detail_field("issueId", "出票ID", "issue_id", "text", width=120),
    detail_field("issueTicketTime", "出票时间", "issue_ticket_time", "datetime", width=175),
    detail_field("issueTicketNo", "出票票号", "issue_ticket_no", width=155),
    detail_field("airline", "航司", "air_line", width=90),
    detail_field("smartConfirm", "使用智能定位", "smart_confirm", category="智能及PNR信息", width=130),
    detail_field("issueProfit", "预估利润", "issue_profit", "money", "金额与派生字段", 120),
    detail_field("gds", "出票GDS", "issue_gds_code", width=100),
    detail_field("supplier", "出票供应商", "issue_supplier_cname", width=150),
    detail_field("issuePcc", "出票PCC", "issue_pcc_code", category="智能及PNR信息", width=110),
    detail_field("refundFlag", "是否退票", "refund_flag", width=100),
    detail_field("issueCreateTime", "出票表创建时间", "issue_create_time", "datetime", width=175),
    detail_field("firstDepTime", "首段起飞时间", "first_dep_time", "datetime", "航程信息", 175),
    detail_field("orderId", "订单ID", "order_id", "text", width=130),
    detail_field("platform", "OTA平台", "ota_cname", width=110),
    detail_field("site", "OTA站点", "ota_site_cname", width=125),
    detail_field("depAirport", "出发机场", "dep_airport", category="航程信息", width=100),
    detail_field("arrAirport", "到达机场", "arr_airport", category="航程信息", width=100),
    detail_field("route", "航线", "airroute_code", category="航程信息", width=130),
    detail_field("depCityCode", "出发城市码", "dep_city", category="航程信息", width=110),
    detail_field("depCity", "出发城市", "dep_city_cname", category="航程信息", width=110),
    detail_field("depCountryCode", "出发国家码", "dep_country", category="航程信息", width=110),
    detail_field("depCountry", "出发国家", "dep_country_cname", category="航程信息", width=110),
    detail_field("depZoneCode", "出发大区码", "dep_zone", category="航程信息", width=110),
    detail_field("depZone", "出发大区", "dep_zone_cname", category="航程信息", width=110),
    detail_field("depAreaCode", "出发TC区码", "dep_area", category="航程信息", width=110),
    detail_field("depArea", "出发TC区", "dep_area_cname", category="航程信息", width=110),
    detail_field("arrCityCode", "到达城市码", "arr_city", category="航程信息", width=110),
    detail_field("arrCity", "到达城市", "arr_city_cname", category="航程信息", width=110),
    detail_field("arrCountryCode", "到达国家码", "arr_country", category="航程信息", width=110),
    detail_field("arrCountry", "到达国家", "arr_country_cname", category="航程信息", width=110),
    detail_field("arrZoneCode", "到达大区码", "arr_zone", category="航程信息", width=110),
    detail_field("arrZone", "到达大区", "arr_zone_cname", category="航程信息", width=110),
    detail_field("arrAreaCode", "到达TC区码", "arr_area", category="航程信息", width=110),
    detail_field("arrArea", "到达TC区", "arr_area_cname", category="航程信息", width=110),
    detail_field("leader", "组长", "policy_operator", category="人员信息", width=110),
    detail_field("member", "组员", "policy_operator_group_name", category="人员信息", width=110),
    detail_field("policy", "政策员", "policy_name", category="人员信息", width=110),
    detail_field("orderStatus", "订单状态", "order_status", width=110),
    detail_field("issueStatus", "出票状态", "issue_status", width=110),
    detail_field("orderSource", "订单来源", "manual_flag", width=110),
    detail_field("orderCreateTime", "订单表创建时间", "order_create_time", "datetime", width=175),
    detail_field("otaOrderNo", "OTA订单号", "ota_order_no", width=155),
    detail_field("cabin", "原舱", "cabin", category="智能及PNR信息", width=100),
    detail_field("pnrPassengerCount", "订位乘客人数", "pnr_pass_count", "count", "智能及PNR信息", 120),
    detail_field("saleDiffPrice", "销售差价", "sale_diff_price", "money", "金额与派生字段", 120),
    detail_field("bookPnrNo", "订位PNR", "book_pnr_no", category="智能及PNR信息", width=120),
    detail_field("businessFlag", "行李符合与否", "business_flag", category="智能及PNR信息", width=120),
    detail_field("actualCabin", "智能任务后舱位", "actual_cabin", category="智能及PNR信息", width=135),
    detail_field("flightWay", "航程类型", "flight_way", category="航程信息", width=110),
    detail_field("transferType", "航段类型", "transfer_type", category="航程信息", width=110),
    detail_field("tripCount", "总航段数", "trip_count", "count", width=100),
    detail_field("rewardTripCount", "有奖航段数", "reward_trip_count", "count", width=110),
    detail_field("realCabin", "实际出票舱位", "real_cabin", category="智能及PNR信息", width=120),
    detail_field("ndcFlag", "NDC标识", "ndc_flag", width=100),
    detail_field("ageType", "乘客类型", "age_type", category="航程信息", width=100),
    detail_field("originalBookPnr", "原始-订位PNR", "iss_book_pnr_no", category="智能及PNR信息", width=130),
    detail_field("bookOfficeNo", "原始-订位PCC号", "book_office_no", width=135),
    detail_field("issuePnrNo", "出票PNR", "issue_pnr_no", category="智能及PNR信息", width=120),
    detail_field("configuredPcc", "占位配置PCC号", "pcc_code", category="智能及PNR信息", width=135),
    detail_field("pnrSwitchCount", "占位成功人数", "pnr_switch_count", "count", "智能及PNR信息", 120),
    detail_field("routeCategory", "航线类别", "airroute_flag", category="航程信息", width=110),
    detail_field("otaLastIssueTime", "OTA最晚出票时间", "ota_last_issue_time", "datetime", width=175),
    detail_field("productType", "产品类型", "product_type", width=110),
    detail_field("policyTicketPrice", "政策票面价（外币）", "policy_ticket_price", "money", "金额与派生字段", 150),
    detail_field("policyTax", "政策税费（外币）", "policy_tax", "money", "金额与派生字段", 145),
    detail_field("exchangeRate", "汇率", "policy_exchange_rate", "money", "金额与派生字段", 110),
    detail_field("saleCurrency", "销售结算币种", "sale_currency", category="金额与派生字段", width=125),
    detail_field("policyTicketPriceCny", "政策票面价（人民币）", "policy_ticket_price_cny", "money", "金额与派生字段", 160),
    detail_field("policyTaxCny", "政策税费（人民币）", "policy_tax_cny", "money", "金额与派生字段", 155),
    detail_field("saleTotalAmountCny", "销售结算总价（人民币）", "sale_total_amount_cny", "money", "金额与派生字段", 175),
    detail_field("issueTicketDate", "出票日期", "issue_ticket_date", "date", width=115),
    detail_field("firstDepDate", "首段起飞日期", "first_dep_date", "date", "航程信息", 125),
    detail_field("orderCreateDate", "进单日期", "order_create_date", "date", width=115),
    detail_field("downParticipated", "降舱是否参与", "down_participated", "boolean", "智能状态", 115),
    detail_field("downSucceeded", "降舱是否成功", "down_succeeded", "boolean", "智能状态", 115),
    detail_field("voidParticipated", "废票是否参与", "void_participated", "boolean", "智能状态", 115),
    detail_field("voidSucceeded", "废票是否成功", "void_succeeded", "boolean", "智能状态", 115),
    detail_field("cityRoute", "城市航线", "city_route", category="金额与派生字段", width=140),
    detail_field("countryRoute", "国家航线", "country_route", category="金额与派生字段", width=140),
    detail_field("zoneRoute", "大区航线", "zone_route", category="金额与派生字段", width=140),
    detail_field("tcRoute", "TC区航线", "tc_route", category="金额与派生字段", width=140),
    detail_field("routeLength", "长短航线", "route_length_type", category="金额与派生字段", width=110),
    detail_field("profitBandCode", "利润区间编码", "profit_band_code", category="金额与派生字段", width=125),
    detail_field("profitBand", "利润区间", "profit_band_name", category="金额与派生字段", width=110),
    detail_field("scheduledAt", "调度时间", "create_time", "datetime", width=175),
    detail_field("issueMonth", "出票月份", "date_format(src.issue_ticket_date,'%%Y-%%m')", category="金额与派生字段", width=100, raw=True),
    detail_field("issueYear", "出票年度", "date_format(src.issue_ticket_date,'%%Y')", category="金额与派生字段", width=90, raw=True),
    detail_field("depMonth", "首段起飞月份", "date_format(src.first_dep_date,'%%Y-%%m')", category="金额与派生字段", width=120, raw=True),
    detail_field("depYear", "首段起飞年度", "date_format(src.first_dep_date,'%%Y')", category="金额与派生字段", width=110, raw=True),
    detail_field("orderMonth", "进单月份", "date_format(src.order_create_date,'%%Y-%%m')", category="金额与派生字段", width=100, raw=True),
    detail_field("orderYear", "进单年度", "date_format(src.order_create_date,'%%Y')", category="金额与派生字段", width=90, raw=True),
)
DETAIL_BY_KEY = {item.key: item for item in DETAIL_FIELDS}
DETAIL_FIELD_POSITION = {item.key: index for index, item in enumerate(DETAIL_FIELDS)}
DETAIL_METRIC_TYPES = {"money", "count", "boolean"}
DEFAULT_DETAIL_FIELDS = (
    "issueId", "issueTicketNo", "otaOrderNo", "airline", "platform",
    "issueTicketTime", "smartConfirm", "issueProfit",
)
EVIDENCE_FIELDS = (
    "downParticipated", "downSucceeded", "voidParticipated", "voidSucceeded",
)


def ordered_detail_fields(keys: list[str] | tuple[str, ...]) -> list[DetailField]:
    unique_keys = dict.fromkeys(("issueId", *keys))
    return sorted(
        (DETAIL_BY_KEY[key] for key in unique_keys),
        key=lambda item: (
            item.key != "issueId",
            item.value_type in DETAIL_METRIC_TYPES,
            DETAIL_FIELD_POSITION[item.key],
        ),
    )


def _today() -> date:
    return datetime.now(ZoneInfo("Asia/Shanghai")).date()


def _optional_period(
    start_value: str | None, end_value: str | None, label: str,
) -> tuple[date | None, date | None]:
    if not start_value and not end_value:
        return None, None
    if not start_value or not end_value:
        raise ValueError(f"请同时选择{label}的开始和结束日期")
    try:
        first = date.fromisoformat(start_value)
        last = date.fromisoformat(end_value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label}必须使用YYYY-MM-DD格式") from error
    if first > last:
        raise ValueError(f"{label}开始日期不能晚于结束日期")
    if (last - first).days >= MAX_RANGE_DAYS:
        raise ValueError(f"{label}单次最多查询{MAX_RANGE_DAYS}天")
    return first, last


def parse_query_periods(
    start_value: str | None, end_value: str | None,
    first_dep_start: str | None, first_dep_end: str | None,
    order_start: str | None, order_end: str | None,
) -> tuple[date | None, date | None, date | None, date | None, date | None, date | None]:
    issue_first, issue_last = _optional_period(start_value, end_value, "出票日期")
    dep_first, dep_last = _optional_period(first_dep_start, first_dep_end, "首段起飞日期")
    order_first, order_last = _optional_period(order_start, order_end, "进单日期")
    if not any((issue_first, dep_first, order_first)):
        raise ValueError("出票、首段起飞、进单日期至少选择一个完整范围")
    return issue_first, issue_last, dep_first, dep_last, order_first, order_last


def parse_filters(filters: dict[str, str | None]) -> dict[str, str]:
    cleaned: dict[str, str] = {}
    for key, value in filters.items():
        if key not in DIMENSIONS or value in (None, ""):
            continue
        text = str(value).strip()
        if not text or len(text) > 200:
            raise ValueError(f"{DIMENSIONS[key][1]}筛选值不合法")
        cleaned[key] = text
    return cleaned


def build_where(
    *, start: date | None, end: date | None, first_dep_start: date | None = None,
    first_dep_end: date | None = None, order_start: date | None = None,
    order_end: date | None = None, filters: dict[str, str | None] | None = None,
    table_alias: str = "src",
) -> tuple[str, list[Any], dict[str, str]]:
    cleaned = parse_filters(filters or {})
    prefix = f"{table_alias}." if table_alias else ""
    conditions: list[str] = []
    params: list[Any] = []
    for field, value, operator in (
        ("issue_ticket_date", start, ">="), ("issue_ticket_date", end, "<="),
        ("first_dep_date", first_dep_start, ">="), ("first_dep_date", first_dep_end, "<="),
        ("order_create_date", order_start, ">="), ("order_create_date", order_end, "<="),
    ):
        if value:
            conditions.append(f"{prefix}{field} {operator} %s")
            params.append(value.isoformat())
    for key, value in cleaned.items():
        field = DIMENSIONS[key][0]
        if value == NULL_TOKEN:
            conditions.append(f"nullif(trim({prefix}{field}), '') IS NULL")
        else:
            conditions.append(f"nullif(trim({prefix}{field}), '') = %s")
            params.append(value)
    return " AND ".join(conditions), params, cleaned


def _month_start(value: date, months_back: int) -> date:
    month_index = value.year * 12 + value.month - 1 - months_back
    return date(month_index // 12, month_index % 12 + 1, 1)


def build_trend_where(
    *, start: date | None, end: date | None, trend_date: str,
    first_dep_start: date | None = None, first_dep_end: date | None = None,
    order_start: date | None = None, order_end: date | None = None,
    filters: dict[str, str | None] | None = None, table_alias: str = "src",
) -> tuple[str, list[Any], dict[str, Any]]:
    if trend_date not in TREND_DATES:
        raise ValueError("趋势统计日期不合法")
    cutoff = {
        "issue": end or _today(),
        "departure": first_dep_end or _today(),
        "order": order_end or _today(),
    }[trend_date]
    cutoff_source = {
        "issue": "出票结束日期" if end else "当前日期",
        "departure": "首段起飞结束日期" if first_dep_end else "当前日期",
        "order": "进单结束日期" if order_end else "当前日期",
    }[trend_date]
    monthly_start = _month_start(cutoff, 11)
    daily_start = cutoff - timedelta(days=29)
    trend_field, trend_label = TREND_DATES[trend_date]
    cleaned = parse_filters(filters or {})
    prefix = f"{table_alias}." if table_alias else ""
    conditions = [f"{prefix}{trend_field} >= %s", f"{prefix}{trend_field} <= %s"]
    params: list[Any] = [monthly_start.isoformat(), cutoff.isoformat()]

    # 趋势所选日期轴用固定时间窗口，另外两个顶部日期条件仍同时生效。
    if trend_date != "issue":
        for value, operator in ((start, ">="), (end, "<=")):
            if value:
                conditions.append(f"{prefix}issue_ticket_date {operator} %s")
                params.append(value.isoformat())
    if trend_date != "departure":
        for value, operator in ((first_dep_start, ">="), (first_dep_end, "<=")):
            if value:
                conditions.append(f"{prefix}first_dep_date {operator} %s")
                params.append(value.isoformat())
    if trend_date != "order":
        for value, operator in ((order_start, ">="), (order_end, "<=")):
            if value:
                conditions.append(f"{prefix}order_create_date {operator} %s")
                params.append(value.isoformat())
    for key, value in cleaned.items():
        field = DIMENSIONS[key][0]
        if value == NULL_TOKEN:
            conditions.append(f"nullif(trim({prefix}{field}), '') IS NULL")
        else:
            conditions.append(f"nullif(trim({prefix}{field}), '') = %s")
            params.append(value)
    return " AND ".join(conditions), params, {
        "key": trend_date, "field": trend_field, "label": trend_label,
        "cutoff_source": cutoff_source,
        "cutoff": cutoff, "daily_start": daily_start, "monthly_start": monthly_start,
    }


def _workbook(sheet_name: str, headers: list[str], rows: list[list[Any]]) -> BytesIO:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="E9F2FF")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="23446F")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(len(header) * 2 + 4, 12), 28)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _nested_metric(metrics: dict[str, Any], path: tuple[str, ...]) -> Any:
    value: Any = metrics
    for key in path:
        value = value.get(key) if isinstance(value, dict) else None
    return value


def aggregate_select(alias: str = "src") -> str:
    prefix = f"{alias}." if alias else ""
    parts = ["count(1) AS source_rows"]
    parts.extend(f"sum({prefix}{column}) AS {key}" for key, column, _kind in AGGREGATES)
    parts.append(f"max({prefix}create_time) AS updated_at")
    return ", ".join(parts)


def _decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    try:
        return value if isinstance(value, Decimal) else Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _money(value: Any) -> str | None:
    number = _decimal(value)
    return format(number, "f") if number is not None else None


def _ratio(numerator: int | None, denominator: int | None) -> float | None:
    if numerator is None or denominator is None:
        return None
    return round(numerator / denominator, 6) if denominator else 0.0


def _money_ratio(value: Any, denominator: int | None) -> str | None:
    number = _decimal(value)
    if number is None or denominator is None:
        return None
    if denominator == 0:
        return "0"
    return format((number / Decimal(denominator)).quantize(Decimal("0.0001")), "f")


def _raw_aggregate(row: dict[str, Any]) -> dict[str, Any]:
    output: dict[str, Any] = {"source_rows": int(row.get("source_rows") or 0)}
    for key, _column, kind in AGGREGATES:
        value = row.get(key)
        output[key] = int(value) if value is not None and kind == "count" else _decimal(value)
    output["updated_at"] = row.get("updated_at")
    return output


def combine_aggregates(rows: list[dict[str, Any]]) -> dict[str, Any]:
    if not rows:
        return {"source_rows": 0, **{key: None for key, _column, _kind in AGGREGATES}, "updated_at": None}
    combined: dict[str, Any] = {"source_rows": sum(int(row.get("source_rows") or 0) for row in rows)}
    for key, _column, kind in AGGREGATES:
        values = [row.get(key) for row in rows if row.get(key) is not None]
        if not values:
            combined[key] = None
        elif kind == "count":
            combined[key] = sum(int(value) for value in values)
        else:
            combined[key] = sum((_decimal(value) or Decimal(0) for value in values), Decimal(0))
    combined["updated_at"] = max((row.get("updated_at") for row in rows if row.get("updated_at") is not None), default=None)
    return combined


def public_metrics(raw: dict[str, Any]) -> dict[str, Any]:
    has_data = int(raw.get("source_rows") or 0) > 0
    if not has_data:
        return {
            "ticketCount": None, "totalProfit": None, "segmentCount": None,
            "refundCount": None, "perSegmentProfit": None,
            **{prefix: _public_smart_metrics(raw, prefix, False) for prefix in ("smart", "down", "void")},
        }
    ticket_count = raw.get("ticket_count")
    segment_count = raw.get("segment_count")
    return {
        "ticketCount": ticket_count,
        "totalProfit": _money(raw.get("total_profit")),
        "segmentCount": segment_count,
        "refundCount": raw.get("refund_count"),
        "perSegmentProfit": _money_ratio(raw.get("total_profit"), segment_count),
        **{prefix: _public_smart_metrics(raw, prefix, True) for prefix in ("smart", "down", "void")},
    }


def _public_smart_metrics(raw: dict[str, Any], prefix: str, available: bool) -> dict[str, Any]:
    participated = raw.get(f"{prefix}_participated") if available else None
    succeeded = raw.get(f"{prefix}_succeeded") if available else None
    used = raw.get(f"{prefix}_used") if available else None
    profit = raw.get(f"{prefix}_profit") if available else None
    tickets = raw.get("ticket_count") if available else None
    return {
        "participatedCount": participated,
        "succeededCount": succeeded,
        "usedCount": used,
        "profit": _money(profit),
        "participationRate": _ratio(participated, tickets),
        "successRate": _ratio(succeeded, participated),
        "adoptionRate": _ratio(used, succeeded),
        "participationUseRate": _ratio(used, participated),
        "useRate": _ratio(used, tickets),
        "averageProfit": _money_ratio(profit, used),
    }


def _updated_at(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return str(value or "")


def _period_payload(first: date | None, last: date | None) -> dict[str, str]:
    return {
        "startDate": first.isoformat() if first else "",
        "endDate": last.isoformat() if last else "",
    }


def _period_suffix(*ranges: tuple[date | None, date | None]) -> str:
    for first, last in ranges:
        if first and last:
            return f"{first}_{last}"
    return _today().isoformat()


def _empty_overview(
    first: date | None, last: date | None, error: str = "", *, trend_meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    cutoff = last or _today()
    trend_meta = trend_meta or {
        "key": "issue", "label": "出票日期", "cutoff": cutoff,
        "cutoff_source": "出票结束日期" if last else "当前日期",
        "daily_start": cutoff - timedelta(days=29), "monthly_start": _month_start(cutoff, 11),
    }
    return {
        "source": f"MySQL · {DATABASE}.{DIMENSION_TABLE}",
        "available": not bool(error), "hasData": False, "error": error,
        "period": _period_payload(first, last),
        "trend": {
            "dateType": trend_meta["key"], "dateLabel": trend_meta["label"],
            "cutoffSource": trend_meta["cutoff_source"],
            "cutoffDate": trend_meta["cutoff"].isoformat(),
            "dailyStartDate": trend_meta["daily_start"].isoformat(),
            "monthlyStartDate": trend_meta["monthly_start"].isoformat(),
        },
        "updatedAt": "", "metrics": public_metrics({"source_rows": 0}),
        "dailyTrend": [], "monthlyTrend": [],
        "notes": [
            "参与、成功、使用直接汇总智能日指标表；比率按当前筛选范围重新计算。",
            "智能利润是使用智能票单的预估利润，不是智能带来的增量收益。",
            "源表没有独立覆盖标记；查询无记录时展示为暂无数据，不将其认定为真实零。",
        ],
    }


def _serialize_detail(value: Any, value_type: str) -> Any:
    if value is None:
        return None
    if value_type == "count":
        return int(value)
    if value_type == "boolean":
        return bool(int(value))
    if value_type == "money":
        return _money(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


class SmartAnalysisService:
    def __init__(self, config: dict[str, Any], source: DataSource | None = None):
        fixed_config = {**config, "DATA_MODE": "mysql", "MYSQL_DATABASE": DATABASE}
        self.source = source or DataSource(fixed_config)

    def overview(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        trend_date="issue",
    ) -> dict[str, Any]:
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        total_where, total_params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=filters,
        )
        if trend_date not in TREND_DATES:
            raise ValueError("趋势统计日期不合法")
        response = _empty_overview(first, last)
        total_query = (
            f"SELECT {aggregate_select()} FROM {DATABASE}.{DIMENSION_TABLE} src "
            f"WHERE {total_where}"
        )
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(total_query, tuple(total_params))
            total_values = cursor.fetchone()
            total = _raw_aggregate(dict(zip(AGGREGATE_NAMES, total_values))) if total_values else {
                "source_rows": 0
            }
            response["hasData"] = int(total.get("source_rows") or 0) > 0
            response["metrics"] = public_metrics(total)
            response["updatedAt"] = _updated_at(total.get("updated_at"))
            return response
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart overview query failed")
            return _empty_overview(first, last, "智能总览查询失败，请检查MySQL连接、表结构和数据调度。")
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def trends(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        trend_date="issue",
    ) -> dict[str, Any]:
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        trend_where, trend_params, trend_meta = build_trend_where(
            start=first, end=last, trend_date=trend_date,
            first_dep_start=dep_first, first_dep_end=dep_last,
            order_start=order_first, order_end=order_last, filters=filters,
        )
        response = _empty_overview(first, last, trend_meta=trend_meta)
        trend_field = trend_meta["field"]
        query = (
            f"SELECT src.{trend_field}, {aggregate_select()} "
            f"FROM {DATABASE}.{DIMENSION_TABLE} src WHERE {trend_where} "
            f"GROUP BY src.{trend_field} ORDER BY src.{trend_field}"
        )
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(query, tuple(trend_params))
            trend_raw = []
            for values in cursor.fetchall():
                aggregate = _raw_aggregate(dict(zip(AGGREGATE_NAMES, values[1:])))
                trend_raw.append({"period": str(values[0]), **aggregate})
            response["available"] = True
            response["hasData"] = bool(trend_raw)
            response["dailyTrend"] = [
                {"period": row["period"], "metrics": public_metrics(row)}
                for row in trend_raw if row["period"] >= trend_meta["daily_start"].isoformat()
            ]
            monthly: dict[str, list[dict[str, Any]]] = defaultdict(list)
            for row in trend_raw:
                monthly[row["period"][:7]].append(row)
            response["monthlyTrend"] = [
                {"period": month, "metrics": public_metrics(combine_aggregates(rows))}
                for month, rows in sorted(monthly.items())
            ]
            combined = combine_aggregates(trend_raw)
            response["updatedAt"] = _updated_at(combined.get("updated_at"))
            return response
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart trend query failed")
            return _empty_overview(
                first, last, "智能趋势查询失败，请检查MySQL连接、表结构和数据调度。",
                trend_meta=trend_meta,
            )
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def dimensions(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        group_by="airline,platform", smart_type="all", sort_by="smartProfit",
        direction="desc", page_value=1, page_size_value=20,
    ) -> dict[str, Any]:
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        groups = [item.strip() for item in str(group_by or "").split(",") if item.strip()]
        if not groups or len(set(groups)) != len(groups) or any(item not in DIMENSIONS for item in groups):
            raise ValueError("分析维度不合法")
        if smart_type not in SMART_PREFIXES:
            raise ValueError("智能类型不合法")
        if direction not in {"asc", "desc"}:
            raise ValueError("排序方向不合法")
        try:
            page, page_size = int(page_value), int(page_size_value)
        except (TypeError, ValueError) as error:
            raise ValueError("分页参数不合法") from error
        if page < 1 or page > 10_000 or page_size < 1 or page_size > MAX_PAGE_SIZE:
            raise ValueError(f"页码必须大于0，每页最多{MAX_PAGE_SIZE}条")
        prefix = SMART_PREFIXES[smart_type]
        sort_columns = {
            "ticketCount": ("ticket_count", "sum(src.issue_count)"),
            "totalProfit": ("total_profit", "sum(src.issue_profit_sum)"),
            "segments": ("segment_count", "sum(src.segment_count_sum)"),
            "usedCount": (f"{prefix}_used", f"sum(src.{prefix}_used_issue_count)"),
            "smartProfit": (f"{prefix}_profit", f"sum(src.{prefix}_profit_sum)"),
        }
        if sort_by not in sort_columns:
            raise ValueError("排序指标不合法")
        where, params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=filters,
        )
        group_columns = [DIMENSIONS[key][0] for key in groups]
        group_sql = ", ".join(f"src.{field}" for field in group_columns)
        grouped_select = f"{group_sql}, {aggregate_select()}"
        grouped_query = (
            f"SELECT {grouped_select} FROM {DATABASE}.{DIMENSION_TABLE} src "
            f"WHERE {where} GROUP BY {group_sql}"
        )
        count_query = f"SELECT count(1) FROM ({grouped_query}) grouped_count"
        total_query = f"SELECT {aggregate_select()} FROM {DATABASE}.{DIMENSION_TABLE} src WHERE {where}"
        sort_column, sort_expression = sort_columns[sort_by]
        order_parts = [f"({sort_expression} IS NULL) ASC", f"{sort_column} {direction.upper()}"]
        for field in group_columns:
            order_parts.extend((f"({field} IS NULL) ASC", f"{field} ASC"))
        offset = (page - 1) * page_size
        data_query = (
            f"{grouped_query} ORDER BY {', '.join(order_parts)} "
            f"LIMIT {offset}, {page_size}"
        )
        response = {
            "source": f"MySQL · {DATABASE}.{DIMENSION_TABLE}",
            "available": False, "hasData": False, "error": "",
            "period": _period_payload(first, last),
            "dimensions": [{"key": key, "label": DIMENSIONS[key][1]} for key in groups],
            "page": page, "pageSize": page_size, "total": 0,
            "totals": public_metrics({"source_rows": 0}), "rows": [], "updatedAt": "",
        }
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(total_query, tuple(params))
            total_values = cursor.fetchone()
            total_raw = _raw_aggregate(dict(zip(AGGREGATE_NAMES, total_values)))
            cursor.execute(count_query, tuple(params))
            response["total"] = int((cursor.fetchone() or (0,))[0] or 0)
            cursor.execute(data_query, tuple(params))
            rows = []
            for values in cursor.fetchall():
                dimension_values = values[:len(groups)]
                raw = _raw_aggregate(dict(zip(AGGREGATE_NAMES, values[len(groups):])))
                rows.append({
                    "key": "\u001f".join(str(value or "") for value in dimension_values),
                    "path": [
                        {
                            "key": key, "label": DIMENSIONS[key][1],
                            "value": NULL_TOKEN if value in (None, "") else str(value),
                            "name": str(value or "未填写"),
                        }
                        for key, value in zip(groups, dimension_values)
                    ],
                    "metrics": public_metrics(raw),
                })
            response.update({
                "available": True,
                "hasData": int(total_raw.get("source_rows") or 0) > 0,
                "totals": public_metrics(total_raw), "rows": rows,
                "updatedAt": _updated_at(total_raw.get("updated_at")),
            })
            return response
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart dimension query failed")
            response["error"] = "智能维度分析查询失败，请检查MySQL连接、表结构和查询范围。"
            return response
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def filter_options(
        self, *, field: str, search: str | None = None, start_value=None,
        end_value=None, first_dep_start=None, first_dep_end=None,
        order_start=None, order_end=None, filters=None,
    ) -> dict[str, Any]:
        if field not in DIMENSIONS:
            raise ValueError("筛选字段不合法")
        query_search = str(search or "").strip()
        if len(query_search) > 100:
            raise ValueError("搜索内容不能超过100个字符")
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        other_filters = dict(filters or {})
        other_filters.pop(field, None)
        where, params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=other_filters,
        )
        column = DIMENSIONS[field][0]
        conditions = [where, f"nullif(trim(src.{column}), '') IS NOT NULL"]
        if query_search:
            conditions.append(f"instr(lower(cast(src.{column} as char)), lower(%s)) > 0")
            params.append(query_search)
        query = (
            f"SELECT src.{column}, count(1) FROM {DATABASE}.{DIMENSION_TABLE} src "
            f"WHERE {' AND '.join(conditions)} GROUP BY src.{column} "
            f"ORDER BY src.{column} LIMIT {MAX_OPTION_ROWS + 1}"
        )
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(query, tuple(params))
            rows = cursor.fetchall()
            return {
                "available": True, "error": "", "field": field,
                "options": [{"value": str(row[0]), "label": str(row[0])} for row in rows[:MAX_OPTION_ROWS]],
                "truncated": len(rows) > MAX_OPTION_ROWS,
            }
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart filter options query failed")
            return {"available": False, "error": "筛选项加载失败。", "field": field, "options": [], "truncated": False}
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def details(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        smart_type="all", stage="all", fields_value=None, page_value=1,
        page_size_value=20,
    ) -> dict[str, Any]:
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        if smart_type not in SMART_PREFIXES:
            raise ValueError("智能类型不合法")
        if stage == "profit":
            stage = "use"
        if stage not in {"all", "participation", "success", "use", "unsuccessful", "unused", "refund"}:
            raise ValueError("票单范围不合法")
        try:
            page, page_size = int(page_value), int(page_size_value)
        except (TypeError, ValueError) as error:
            raise ValueError("分页参数不合法") from error
        if page < 1 or page > 10_000 or page_size < 1 or page_size > MAX_PAGE_SIZE:
            raise ValueError(f"页码必须大于0，每页最多{MAX_PAGE_SIZE}条")
        requested = [key.strip() for key in str(fields_value or "").split(",") if key.strip()]
        if not requested:
            requested = list(DEFAULT_DETAIL_FIELDS)
        if any(key not in DETAIL_BY_KEY for key in requested):
            raise ValueError("明细字段不合法")
        selected_fields = ordered_detail_fields([*requested, *EVIDENCE_FIELDS])
        where, params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=filters,
        )
        stage_condition = self._stage_condition(smart_type, stage)
        if stage_condition:
            where += f" AND ({stage_condition})"
        summary_query = (
            f"SELECT count(1), count(distinct src.order_id), sum(src.issue_profit), max(src.create_time) "
            f"FROM {DATABASE}.{DETAIL_TABLE} src WHERE {where}"
        )
        offset = (page - 1) * page_size
        data_query = (
            f"SELECT {', '.join(item.expression for item in selected_fields)} "
            f"FROM {DATABASE}.{DETAIL_TABLE} src WHERE {where} "
            "ORDER BY (src.issue_ticket_time IS NULL) ASC, "
            f"src.issue_ticket_time DESC, src.issue_id DESC LIMIT {offset}, {page_size}"
        )
        response = {
            "source": f"MySQL · {DATABASE}.{DETAIL_TABLE}",
            "available": False, "hasData": False, "error": "",
            "period": _period_payload(first, last),
            "page": page, "pageSize": page_size, "total": 0,
            "summary": {"ticketCount": None, "orderCount": None, "profit": None},
            "columns": [self._public_detail_column(item) for item in selected_fields],
            "availableColumns": [self._public_detail_column(item) for item in DETAIL_FIELDS],
            "rows": [], "updatedAt": "",
        }
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(summary_query, tuple(params))
            summary = cursor.fetchone() or (0, 0, None, None)
            cursor.execute(data_query, tuple(params))
            rows = []
            for values in cursor.fetchall():
                row = {
                    item.key: _serialize_detail(value, item.value_type)
                    for item, value in zip(selected_fields, values)
                }
                row["recordKey"] = str(row.get("issueId") or f"{page}-{len(rows) + 1}")
                rows.append(row)
            total = int(summary[0] or 0)
            response.update({
                "available": True, "hasData": total > 0, "total": total,
                "summary": {
                    "ticketCount": total, "orderCount": int(summary[1] or 0),
                    "profit": _money(summary[2]),
                },
                "rows": rows, "updatedAt": _updated_at(summary[3]),
            })
            return response
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart detail query failed")
            response["error"] = "智能出票明细查询失败，请检查MySQL连接、表结构和查询范围。"
            return response
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def ticket(self, issue_id: str) -> dict[str, Any]:
        ticket_id = str(issue_id or "").strip()
        if not ticket_id or len(ticket_id) > 100 or any(ord(char) < 32 for char in ticket_id):
            raise ValueError("出票ID不合法")
        fields = ordered_detail_fields(tuple(DETAIL_BY_KEY))
        aliases = [f"field_{index}" for index in range(len(fields))]
        query = (
            f"SELECT {', '.join(f'{item.expression} AS {alias}' for item, alias in zip(fields, aliases))} "
            f"FROM {DATABASE}.{DETAIL_TABLE} src WHERE src.issue_id = %s "
            "ORDER BY (src.issue_ticket_time IS NULL) ASC, src.issue_ticket_time DESC LIMIT 1"
        )
        response = {
            "source": f"MySQL · {DATABASE}.{DETAIL_TABLE}", "available": False,
            "found": False, "error": "", "columns": [self._public_detail_column(item) for item in fields],
            "record": None,
        }
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(query, (ticket_id,))
            values = cursor.fetchone()
            response["available"] = True
            if values:
                response["found"] = True
                response["record"] = {
                    item.key: _serialize_detail(value, item.value_type)
                    for item, value in zip(fields, values)
                }
            return response
        except Exception:
            logging.getLogger(__name__).exception("MySQL smart ticket query failed")
            response["error"] = "出票记录加载失败，请检查MySQL连接和明细表。"
            return response
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def export_trend(self, *, kind="daily", smart_type="all", trend_date="issue", **kwargs):
        if kind not in {"daily", "monthly"}:
            raise ValueError("趋势导出类型不合法")
        if smart_type not in SMART_PREFIXES:
            raise ValueError("智能类型不合法")
        result = self.trends(trend_date=trend_date, **kwargs)
        if not result["available"]:
            raise ValueError(result["error"] or "趋势数据暂时不可用")
        key = SMART_PREFIXES[smart_type]
        rows = result["dailyTrend"] if kind == "daily" else result["monthlyTrend"]
        headers = [result["trend"]["dateLabel"], "全部出票量", "智能使用出票量", "智能利润", "智能使用率"]
        values = [[
            row["period"], row["metrics"]["ticketCount"], row["metrics"][key]["usedCount"],
            row["metrics"][key]["profit"], row["metrics"][key]["useRate"],
        ] for row in rows]
        label = "每日" if kind == "daily" else "每月"
        return _workbook(f"{label}趋势", headers, values), f"智能分析_{label}趋势_{result['trend']['cutoffDate']}.xlsx"

    def export_dimensions(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        group_by="airline,platform", metrics_value="ticketCount,smartUsed,smartProfit",
        smart_type="all", sort_by="smartProfit", direction="desc",
    ):
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        groups = [item.strip() for item in str(group_by or "").split(",") if item.strip()]
        metrics = [item.strip() for item in str(metrics_value or "").split(",") if item.strip()]
        if not groups or len(set(groups)) != len(groups) or any(item not in DIMENSIONS for item in groups):
            raise ValueError("分析维度不合法")
        if not metrics or any(item not in DIMENSION_EXPORT_METRICS for item in metrics):
            raise ValueError("导出指标不合法")
        if smart_type not in SMART_PREFIXES or direction not in {"asc", "desc"}:
            raise ValueError("排序条件不合法")
        prefix = SMART_PREFIXES[smart_type]
        sort_columns = {
            "ticketCount": ("ticket_count", "sum(src.issue_count)"),
            "totalProfit": ("total_profit", "sum(src.issue_profit_sum)"),
            "segments": ("segment_count", "sum(src.segment_count_sum)"),
            "usedCount": (f"{prefix}_used", f"sum(src.{prefix}_used_issue_count)"),
            "smartProfit": (f"{prefix}_profit", f"sum(src.{prefix}_profit_sum)"),
        }
        if sort_by not in sort_columns:
            raise ValueError("排序指标不合法")
        where, params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=filters,
        )
        group_columns = [DIMENSIONS[key][0] for key in groups]
        group_sql = ", ".join(f"src.{field}" for field in group_columns)
        sort_column, sort_expression = sort_columns[sort_by]
        order_parts = [f"({sort_expression} IS NULL) ASC", f"{sort_column} {direction.upper()}"]
        for field in group_columns:
            order_parts.extend((f"({field} IS NULL) ASC", f"{field} ASC"))
        query = (
            f"SELECT {group_sql}, {aggregate_select()} FROM {DATABASE}.{DIMENSION_TABLE} src "
            f"WHERE {where} GROUP BY {group_sql} ORDER BY {', '.join(order_parts)}"
        )
        total_query = f"SELECT {aggregate_select()} FROM {DATABASE}.{DIMENSION_TABLE} src WHERE {where}"
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(total_query, tuple(params))
            total_values = cursor.fetchone()
            total_metrics = public_metrics(_raw_aggregate(dict(zip(AGGREGATE_NAMES, total_values))))
            cursor.execute(query, tuple(params))
            export_rows: list[list[Any]] = []
            total_row = ["合计", *([""] * (len(groups) - 1))]
            total_row.extend(_nested_metric(total_metrics, DIMENSION_EXPORT_METRICS[key][1]) for key in metrics)
            export_rows.append(total_row)
            for values in cursor.fetchall():
                dimension_values = [str(value or "未填写") for value in values[:len(groups)]]
                row_metrics = public_metrics(_raw_aggregate(dict(zip(AGGREGATE_NAMES, values[len(groups):]))))
                export_rows.append(dimension_values + [
                    _nested_metric(row_metrics, DIMENSION_EXPORT_METRICS[key][1]) for key in metrics
                ])
            headers = [DIMENSIONS[key][1] for key in groups] + [DIMENSION_EXPORT_METRICS[key][0] for key in metrics]
            suffix = _period_suffix((first, last), (dep_first, dep_last), (order_first, order_last))
            return _workbook("维度分析", headers, export_rows), f"智能分析_维度分析_{suffix}.xlsx"
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    def export_details(
        self, *, start_value=None, end_value=None, first_dep_start=None,
        first_dep_end=None, order_start=None, order_end=None, filters=None,
        smart_type="all", stage="all", fields_value=None, all_fields=False,
    ):
        first, last, dep_first, dep_last, order_first, order_last = parse_query_periods(
            start_value, end_value, first_dep_start, first_dep_end, order_start, order_end,
        )
        if smart_type not in SMART_PREFIXES:
            raise ValueError("智能类型不合法")
        if stage == "profit":
            stage = "use"
        if stage not in {"all", "participation", "success", "use", "unsuccessful", "unused", "refund"}:
            raise ValueError("票单范围不合法")
        requested = [key.strip() for key in str(fields_value or "").split(",") if key.strip()]
        if all_fields:
            selected_fields = ordered_detail_fields(tuple(DETAIL_BY_KEY))
        else:
            if not requested:
                requested = list(DEFAULT_DETAIL_FIELDS)
            if any(key not in DETAIL_BY_KEY for key in requested):
                raise ValueError("明细字段不合法")
            selected_fields = ordered_detail_fields(requested)
        where, params, _cleaned = build_where(
            start=first, end=last, first_dep_start=dep_first,
            first_dep_end=dep_last, order_start=order_first,
            order_end=order_last, filters=filters,
        )
        stage_condition = self._stage_condition(smart_type, stage)
        if stage_condition:
            where += f" AND ({stage_condition})"
        count_query = f"SELECT count(1) FROM {DATABASE}.{DETAIL_TABLE} src WHERE {where}"
        query = (
            f"SELECT {', '.join(item.expression for item in selected_fields)} "
            f"FROM {DATABASE}.{DETAIL_TABLE} src WHERE {where} "
            "ORDER BY (src.issue_ticket_time IS NULL) ASC, "
            f"src.issue_ticket_time DESC, src.issue_id DESC LIMIT {MAX_DETAIL_EXPORT_ROWS}"
        )
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(count_query, tuple(params))
            total = int((cursor.fetchone() or (0,))[0] or 0)
            if total > MAX_DETAIL_EXPORT_ROWS:
                raise ValueError(f"当前匹配{total}条，超过{MAX_DETAIL_EXPORT_ROWS}条导出上限，请缩小筛选范围后再导出。")
            cursor.execute(query, tuple(params))
            rows = [[
                _serialize_detail(value, field.value_type) for field, value in zip(selected_fields, values)
            ] for values in cursor.fetchall()]
            headers = [field.title for field in selected_fields]
            suffix = _period_suffix((first, last), (dep_first, dep_last), (order_first, order_last))
            return _workbook("出票明细", headers, rows), f"智能分析_出票明细_{suffix}.xlsx"
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()

    @staticmethod
    def _public_detail_column(item: DetailField) -> dict[str, Any]:
        return {
            "key": item.key, "title": item.title, "valueType": item.value_type,
            "category": item.category, "width": item.width,
        }

    @staticmethod
    def _stage_condition(smart_type: str, stage: str) -> str:
        if stage == "all":
            return ""
        if stage == "refund":
            return "src.refund_flag IN ('退票中','已退票')"
        if smart_type == "down":
            participated, succeeded = "src.down_participated = 1", "src.down_succeeded = 1"
            used = "src.smart_confirm = '智能降舱'"
            not_succeeded = "coalesce(src.down_succeeded, 0) <> 1"
            not_used = "(src.smart_confirm IS NULL OR src.smart_confirm <> '智能降舱')"
        elif smart_type == "void":
            participated, succeeded = "src.void_participated = 1", "src.void_succeeded = 1"
            used = "src.smart_confirm = '智能废票'"
            not_succeeded = "coalesce(src.void_succeeded, 0) <> 1"
            not_used = "(src.smart_confirm IS NULL OR src.smart_confirm <> '智能废票')"
        else:
            participated = "(src.down_participated = 1 OR src.void_participated = 1)"
            succeeded = "(src.down_succeeded = 1 OR src.void_succeeded = 1)"
            used = "src.smart_confirm IN ('智能降舱','智能废票')"
            not_succeeded = "(coalesce(src.down_succeeded, 0) <> 1 AND coalesce(src.void_succeeded, 0) <> 1)"
            not_used = "(src.smart_confirm IS NULL OR src.smart_confirm NOT IN ('智能降舱','智能废票'))"
        return {
            "participation": participated,
            "success": succeeded,
            "use": used,
            "unsuccessful": f"({participated}) AND ({not_succeeded})",
            "unused": f"({succeeded}) AND ({not_used})",
        }[stage]
