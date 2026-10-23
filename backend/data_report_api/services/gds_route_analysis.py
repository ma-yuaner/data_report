from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

from ..time_utils import business_now
from .data_source import DataSource


class GdsRouteAnalysisUnavailable(RuntimeError):
    """The Hive route aggregate is unavailable or does not match its contract."""


REQUIRED_COLUMNS = {
    "operator_date", "data_source", "api_method", "source_api_group",
    "gds_code", "pcc_code", "airroute_code", "dep_airport", "arr_airport",
    "dep_city", "arr_city", "ota_code", "ota_site_code", "poi_code",
    "poi_list", "api_request_count_all", "api_request_success",
    "api_request_error", "api_execute_time", "api_result_success",
    "api_cache_count", "api_timeout_count", "api_price_count",
    "api_price_change", "direct_flight_count", "connecting_flight_count",
    "operator_mintes_max", "dt",
}

DIMENSION_COLUMNS = (
    "airroute_code", "dep_city", "arr_city", "dep_airport", "arr_airport",
    "ota_code", "ota_site_code", "gds_code", "pcc_code",
)


def _identifier(value: Any, default: str) -> str:
    text = str(value or default).strip()
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", text):
        raise RuntimeError("GDS航线Hive表名配置不合法")
    return text


def _number(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, Decimal):
        return int(value)
    return int(value)


def _ratio(numerator: int, denominator: int) -> str | None:
    if denominator <= 0:
        return None
    return f"{numerator * 100 / denominator:.2f}"


def _average(total: int, count: int) -> str | None:
    if count <= 0:
        return None
    return f"{total / count:.2f}"


def _channel_expression(channel: str) -> str:
    normalized = "upper(regexp_replace(coalesce(source_api_group, ''), '[^A-Za-z0-9]', ''))"
    if channel == "MPIS":
        return f"{normalized} rlike 'MPIS$'"
    if channel == "NDC":
        return f"{normalized} rlike 'NDC$'"
    return f"not ({normalized} rlike 'MPIS$') and not ({normalized} rlike 'NDC$')"


def _resource_channel_expression(channel: str) -> str:
    """Build the user-facing MP/MPIS filter without changing metric buckets.

    The governance definition treats NDC as a child of the broad MP channel while
    the metric payload still exposes ordinary MP and NDC separately for analysis.
    """
    if channel == "MPIS":
        return _channel_expression("MPIS")
    if channel == "MP":
        return f"not ({_channel_expression('MPIS')})"
    if channel == "NDC":
        return _channel_expression("NDC")
    raise ValueError("资源通道仅支持MP、MPIS或NDC")


API_METHOD_EXPRESSION = "lower(trim(coalesce(api_method, '')))"
SEARCH_METHOD_CONDITION = f"{API_METHOD_EXPRESSION} = 'search'"
VERIFY_METHOD_CONDITION = f"{API_METHOD_EXPRESSION} = 'verify'"
ORDER_METHOD_CONDITION = f"{API_METHOD_EXPRESSION} in ('ordercreate', 'order')"


METRIC_SQL = f"""
    coalesce(sum(api_request_count_all), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_request_success else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_request_error else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_result_success else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_price_count else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_timeout_count else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_cache_count else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then api_execute_time else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_request_success else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_request_error else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_timeout_count else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_execute_time else 0 end), 0),
    coalesce(sum(case when {VERIFY_METHOD_CONDITION} then api_price_change else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_request_success else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_request_error else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_timeout_count else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_execute_time else 0 end), 0),
    coalesce(sum(case when {ORDER_METHOD_CONDITION} then api_price_change else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} and {_channel_expression('MP')} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} and {_channel_expression('MPIS')} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} and {_channel_expression('NDC')} then api_request_count_all else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then direct_flight_count else 0 end), 0),
    coalesce(sum(case when {SEARCH_METHOD_CONDITION} then connecting_flight_count else 0 end), 0),
    max(operator_mintes_max)
""".strip()


def _metric(row: Iterable[Any]) -> dict[str, Any]:
    values = list(row)
    numbers = [_number(value) for value in values[:26]]
    (
        request_count, search_requests, search_success, search_errors,
        search_result_success, search_price_count, search_timeouts, search_cache,
        search_execute_time, verify_requests, verify_success, verify_errors,
        verify_timeouts, verify_execute_time, verify_price_changes, order_requests,
        order_success, order_errors, order_timeouts, order_execute_time,
        order_price_changes, mp_requests, mpis_requests, ndc_requests,
        direct_flights, connecting_flights,
    ) = numbers
    return {
        "requestCount": request_count,
        "searchRequests": search_requests,
        "searchSuccess": search_success,
        "searchFailure": search_errors,
        "searchResultSuccess": search_result_success,
        "searchPriceCount": search_price_count,
        "searchTimeouts": search_timeouts,
        "searchCacheCount": search_cache,
        "searchSuccessRate": _ratio(search_success, search_requests),
        "searchResultSuccessRate": _ratio(search_result_success, search_requests),
        "searchTimeoutRate": _ratio(search_timeouts, search_requests),
        "searchCacheRate": _ratio(search_cache, search_requests),
        "avgSearchExecuteTime": _average(search_execute_time, search_success),
        "verifyRequests": verify_requests,
        "verifySuccess": verify_success,
        "verifyFailure": verify_errors,
        "verifySuccessRate": _ratio(verify_success, verify_requests),
        "verifyTimeoutRate": _ratio(verify_timeouts, verify_requests),
        "verifyPriceChangeRate": _ratio(verify_price_changes, verify_requests),
        "avgVerifyExecuteTime": _average(verify_execute_time, verify_success),
        "orderRequests": order_requests,
        "orderSuccess": order_success,
        "orderFailure": order_errors,
        "orderSuccessRate": _ratio(order_success, order_requests),
        "orderTimeoutRate": _ratio(order_timeouts, order_requests),
        "orderPriceChangeRate": _ratio(order_price_changes, order_requests),
        "avgOrderExecuteTime": _average(order_execute_time, order_success),
        "mpSearchRequests": mp_requests,
        "mpisSearchRequests": mpis_requests,
        "ndcSearchRequests": ndc_requests,
        "mpShare": _ratio(mp_requests, search_requests),
        "mpisShare": _ratio(mpis_requests, search_requests),
        "ndcShare": _ratio(ndc_requests, search_requests),
        "directFlightCount": direct_flights,
        "connectingFlightCount": connecting_flights,
        "sourceUpdatedAt": str(values[26]) if values[26] is not None else None,
    }


class GdsRouteAnalysisService:
    def __init__(self, config: dict[str, Any], source: DataSource | None = None):
        self.config = config
        self.source = source or DataSource({**config, "DATA_MODE": "hive"})
        self.table_name = _identifier(config.get("GDS_ROUTE_HIVE_TABLE"), "dws_gds_airroute_day")
        self.table = f"{self.source.database}.{self.table_name}"
        self.max_rows = max(50, min(int(config.get("GDS_ROUTE_MAX_ROWS", 500)), 2000))
        self.max_days = max(1, min(int(config.get("GDS_ROUTE_MAX_DAYS", 31)), 92))
        self.export_limit = max(1, min(int(config.get("GDS_ROUTE_EXPORT_MAX_ROWS", 20000)), 50000))

    def overview(
        self,
        start_value: str | None,
        end_value: str | None,
        filters: dict[str, str] | None = None,
        *, _export: bool = False,
    ) -> dict[str, Any]:
        start_date, end_date = self._period(start_value, end_value)
        normalized_filters = {key: str(value or "").strip() for key, value in (filters or {}).items()}
        if _export and not self._query_ready(normalized_filters):
            raise ValueError("请先选择GDS、PCC、平台、站点、通道、接口组或航线范围再导出")
        limit = self.export_limit if _export else self.max_rows
        connection = self.source.connect()
        cursor = connection.cursor()
        try:
            self._validate_schema(cursor)
            options = {} if _export else self._options(cursor, start_date, end_date)
            if not self._query_ready(normalized_filters):
                return self._empty_response(start_date, end_date, normalized_filters, options)
            where_sql, params = self._where(start_date, end_date, normalized_filters)
            cursor.execute(f"SELECT {METRIC_SQL} FROM {self.table} WHERE {where_sql}", tuple(params))
            summary_row = cursor.fetchone()
            summary = _metric(summary_row or [0] * 27)

            dimensions = ", ".join(DIMENSION_COLUMNS)
            cursor.execute(
                f"SELECT {dimensions}, {METRIC_SQL} FROM {self.table} WHERE {where_sql} "
                f"GROUP BY {dimensions} ORDER BY sum(case when {SEARCH_METHOD_CONDITION} "
                f"then api_request_count_all else 0 end) DESC, {dimensions} LIMIT {limit + 1}",
                tuple(params),
            )
            raw_rows = list(cursor.fetchall())
            truncated = len(raw_rows) > limit
            if _export and truncated:
                raise ValueError(f"航线名单超过导出上限{limit}项，请缩小日期或筛选范围；未导出部分名单")
            rows = [self._route_row(row) for row in raw_rows[:limit]]
            return {
                "source": self.source.label,
                "table": self.table,
                "available": True,
                "queryReady": True,
                "generatedAt": business_now().isoformat(timespec="seconds"),
                "period": {"startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
                "filters": normalized_filters,
                "summary": summary,
                "rows": rows,
                "rowCount": len(rows),
                "resultLimit": limit,
                "truncated": truncated,
                "options": options,
                "notes": [
                    "航线页直接查询Hive日汇总，不经过MySQL服务表。",
                    "日期按Hive分区dt（yyyyMMdd）过滤，operator_date仅作为业务字段保留。",
                    "查询、验价、生单分别按api_method=search、verify、orderCreate统计。",
                    "查询日志不写实际航司，首版不使用poi_code拆分查询量。",
                    "本表没有出票航段分母，航线页不计算查定比，避免产生错误结论。",
                ],
            }
        except ValueError:
            raise
        except Exception as error:
            raise GdsRouteAnalysisUnavailable(f"GDS航线Hive查询失败：{error}") from error
        finally:
            try:
                cursor.close()
            finally:
                connection.close()

    def export_report(self, start_value, end_value, filters=None):
        result = self.overview(start_value, end_value, filters, _export=True)
        columns = [
            ("GDS", "gdsCode"), ("PCC", "pccCode"), ("平台", "platform"), ("站点", "site"),
            ("有方向航线", "route"), ("出发机场", "depAirport"), ("到达机场", "arrAirport"),
            ("出发城市", "depCity"), ("到达城市", "arrCity"),
            ("查询请求量", "searchRequests"), ("查询成功量", "searchSuccess"), ("查询失败量", "searchFailure"),
            ("查询成功率(%)", "searchSuccessRate"), ("查询结果成功率(%)", "searchResultSuccessRate"),
            ("验价请求量", "verifyRequests"), ("验价成功量", "verifySuccess"), ("验价失败量", "verifyFailure"),
            ("验价成功率(%)", "verifySuccessRate"), ("生单请求量", "orderRequests"),
            ("生单成功量", "orderSuccess"), ("生单失败量", "orderFailure"), ("生单成功率(%)", "orderSuccessRate"),
            ("普通MP查询量", "mpSearchRequests"), ("MPIS查询量", "mpisSearchRequests"),
            ("NDC查询量", "ndcSearchRequests"), ("源数据时间", "sourceUpdatedAt"),
        ]
        workbook = Workbook(write_only=True)

        def append(sheet, values):
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Codes and untrusted labels must never become Excel formulas.
                cells.append(cell)
            sheet.append(cells)

        sheet = workbook.create_sheet("航线名单")
        sheet.freeze_panes = "E2"
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22
        append(sheet, [label for label, _ in columns])
        for row in result["rows"]:
            append(sheet, [row.get(key) for _, key in columns])
        sheet.auto_filter.ref = f"A1:Z{len(result['rows']) + 1}"
        scope = workbook.create_sheet("统计范围")
        scope.column_dimensions["A"].width = 24
        scope.column_dimensions["B"].width = 100
        append(scope, ["项目", "内容"])
        for row in [
            ["开始日期", result["period"]["startDate"]], ["结束日期", result["period"]["endDate"]],
            ["生成时间", result["generatedAt"]], ["数据源", result["table"]],
            ["最近源数据时间", result["summary"].get("sourceUpdatedAt")], ["名单条数", len(result["rows"])],
            ["导出上限", self.export_limit], ["范围说明", "按筛选重新查询全部匹配组合，不限当前分页；超限拒绝导出"],
            ["完整性提示", "最近源数据时间不代表全批次完整；日汇总同步覆盖仍需核实"],
            ["口径说明", "按方向航线及四维聚合；各阶段请求不是同一批订单的转化漏斗；不含航段、查定比或关闭建议"],
            *[[key, value] for key, value in result["filters"].items() if value],
        ]:
            append(scope, row)
        content = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
        try:
            workbook.save(content)
            content.seek(0)
        except Exception:
            content.close()
            raise
        finally:
            workbook.close()
        period = result["period"]
        return content, f"GDS航线名单_{period['startDate']}_{period['endDate']}.xlsx"

    def _period(self, start_value: str | None, end_value: str | None) -> tuple[date, date]:
        yesterday = business_now().date() - timedelta(days=1)
        end_date = self._date(end_value, yesterday)
        if end_date > yesterday:
            end_date = yesterday
        start_date = self._date(start_value, end_date)
        if start_date > end_date:
            raise ValueError("统计开始日期不能晚于结束日期")
        if (end_date - start_date).days + 1 > self.max_days:
            raise ValueError(f"Hive航线明细单次最多查询{self.max_days}天")
        return start_date, end_date

    @staticmethod
    def _date(value: str | None, default: date) -> date:
        if not str(value or "").strip():
            return default
        try:
            return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
        except ValueError as error:
            raise ValueError("日期格式必须为YYYY-MM-DD") from error

    def _validate_schema(self, cursor) -> None:
        cursor.execute(f"DESCRIBE {self.table}")
        columns = {
            str(row[0]).strip().lower()
            for row in cursor.fetchall()
            if row and str(row[0] or "").strip() and not str(row[0]).startswith("#")
        }
        missing = sorted(REQUIRED_COLUMNS - columns)
        if missing:
            raise GdsRouteAnalysisUnavailable(
                f"{self.table}缺少航线分析字段：{','.join(missing)}"
            )

    def _where(
        self, start_date: date, end_date: date, filters: dict[str, str]
    ) -> tuple[str, list[Any]]:
        clauses = [
            "dt >= %s", "dt <= %s",
            "(trim(data_source) = '1' or upper(trim(data_source)) = 'GDS')",
            "coalesce(trim(airroute_code), '') <> ''",
        ]
        params: list[Any] = [start_date.strftime("%Y%m%d"), end_date.strftime("%Y%m%d")]
        mappings = {
            "platform": "ota_code", "site": "ota_site_code", "gds": "gds_code",
            "pcc": "pcc_code", "apiGroup": "source_api_group", "route": "airroute_code",
        }
        for key, field in mappings.items():
            value = filters.get(key, "")
            if value:
                clauses.append(f"trim({field}) = %s")
                params.append(value)
        resource_channel = filters.get("resourceChannel", "").upper()
        if resource_channel:
            clauses.append(_resource_channel_expression(resource_channel))
        keyword = filters.get("keyword", "").upper()
        if keyword:
            clauses.append(
                "upper(concat_ws('|', coalesce(airroute_code,''), coalesce(dep_airport,''), "
                "coalesce(arr_airport,''), coalesce(dep_city,''), coalesce(arr_city,''), "
                "coalesce(poi_code,''), coalesce(poi_list,''))) like %s"
            )
            params.append(f"%{keyword}%")
        return " AND ".join(clauses), params

    def _options(self, cursor, start_date: date, end_date: date) -> dict[str, list[str]]:
        cursor.execute(
            f"SELECT DISTINCT gds_code, pcc_code, ota_code, ota_site_code, source_api_group "
            f"FROM {self.table} WHERE dt >= %s AND dt <= %s "
            "AND (trim(data_source) = '1' or upper(trim(data_source)) = 'GDS') LIMIT 2000",
            (start_date.strftime("%Y%m%d"), end_date.strftime("%Y%m%d")),
        )
        rows = cursor.fetchall()
        keys = ("gds", "pcc", "platform", "site", "apiGroup")
        result = {key: set() for key in keys}
        for row in rows:
            for index, key in enumerate(keys):
                value = str(row[index] or "").strip()
                if value:
                    result[key].add(value)
        return {key: sorted(values) for key, values in result.items()}

    @staticmethod
    def _query_ready(filters: dict[str, str]) -> bool:
        return any(filters.get(key) for key in (
            "gds", "pcc", "platform", "site", "resourceChannel", "apiGroup", "route", "keyword",
        ))

    def _empty_response(
        self,
        start_date: date,
        end_date: date,
        filters: dict[str, str],
        options: dict[str, list[str]],
    ) -> dict[str, Any]:
        return {
            "source": self.source.label,
            "table": self.table,
            "available": True,
            "queryReady": False,
            "generatedAt": business_now().isoformat(timespec="seconds"),
            "period": {"startDate": start_date.isoformat(), "endDate": end_date.isoformat()},
            "filters": filters,
            "summary": _metric([0] * 26 + [None]),
            "rows": [],
            "rowCount": 0,
            "resultLimit": self.max_rows,
            "truncated": False,
            "options": options,
            "notes": [
                "已加载所选日期的筛选项；请选择GDS、PCC、平台、站点、通道、接口组或航线后查询。",
                "航线大表默认不执行无条件明细扫描。",
            ],
        }

    @staticmethod
    def _route_row(row: Iterable[Any]) -> dict[str, Any]:
        values = list(row)
        dimensions = [str(value or "").strip() for value in values[: len(DIMENSION_COLUMNS)]]
        payload = dict(zip(DIMENSION_COLUMNS, dimensions))
        payload.update(_metric(values[len(DIMENSION_COLUMNS):]))
        payload.update({
            "id": "|".join(dimensions),
            "route": payload.pop("airroute_code"),
            "depCity": payload.pop("dep_city"),
            "arrCity": payload.pop("arr_city"),
            "depAirport": payload.pop("dep_airport"),
            "arrAirport": payload.pop("arr_airport"),
            "platform": payload.pop("ota_code"),
            "site": payload.pop("ota_site_code"),
            "gdsCode": payload.pop("gds_code"),
            "pccCode": payload.pop("pcc_code"),
        })
        return payload
