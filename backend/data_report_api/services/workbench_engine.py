"""Bounded, declarative data preparation. No user SQL or Python is evaluated."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import tempfile
import zipfile
from contextvars import ContextVar
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .workbench_formula import compile_formula, decimal_scale, install_division


_STEP_CONTEXT = ContextVar("workbench_step", default=None)


class WorkbenchEngineError(ValueError):
    def __init__(self, message, code="INVALID_STEP", report=None):
        context = _STEP_CONTEXT.get()
        if context:
            index, current_report = context
            if not message.startswith("步骤"):
                message = f"步骤{index + 1}：{message}"
            if report is None:
                report = {**current_report, "failedStep": index + 1}
        super().__init__(message)
        self.code = code
        self.report = report or {}


DEFAULT_LIMITS = {
    "max_rows": 1_000_000,
    "max_columns": 200,
    "max_output_rows": 2_000_000,
    "max_result_columns": 1000,
    "max_join_rows": 2_000_000,
    "max_join_expansion": 5,
    "memory_limit": "512MB",
    "threads": 2,
    "max_file_bytes": 200 * 1024 * 1024,
    "max_uncompressed_bytes": 1024 * 1024 * 1024,
}
OPERATORS = ["trim", "replace", "cast", "filter", "derive", "group", "join", "append", "select"]
_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,119}$")


def _q(value):
    if not isinstance(value, str) or not _IDENT.fullmatch(value):
        raise WorkbenchEngineError("字段或步骤标识无效")
    return '"' + value + '"'


def _limits(limits):
    return {**DEFAULT_LIMITS, **(limits or {})}


def _connect(limits):
    import duckdb

    return duckdb.connect(
        config={
            "memory_limit": str(limits["memory_limit"]),
            "threads": int(limits["threads"]),
            "temp_directory": str(limits.get("temp_directory") or tempfile.gettempdir()),
            "max_temp_directory_size": str(limits.get("max_temp_directory_size", "2GB")),
            "autoload_known_extensions": False,
            "autoinstall_known_extensions": False,
        }
    )


def _scalar(value):
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def import_dataset(file_path, output_path, sheet_name=None, limits=None):
    """Stream untrusted XLSX/CSV through a normalized CSV into immutable Parquet."""
    limits = _limits(limits)
    source = Path(file_path)
    if source.stat().st_size > limits["max_file_bytes"]:
        raise WorkbenchEngineError("文件超过大小限制", "FILE_TOO_LARGE")
    workbook = None
    handle = None
    suffix = source.suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            with zipfile.ZipFile(source) as archive:
                if sum(info.file_size for info in archive.infolist()) > limits["max_uncompressed_bytes"]:
                    raise WorkbenchEngineError("Excel解压大小超限", "FILE_TOO_LARGE")
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise WorkbenchEngineError("工作表不存在", "INVALID_SOURCE")
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            if sheet.max_column and sheet.max_column > limits["max_columns"]:
                raise WorkbenchEngineError("字段数量超限", "COLUMN_LIMIT")
            rows = iter(sheet.iter_rows(values_only=True))
        elif suffix == ".csv":
            handle = source.open("r", encoding="utf-8-sig", newline="")
            rows = iter(csv.reader(handle))
        else:
            raise WorkbenchEngineError("仅支持UTF-8 CSV和XLSX", "INVALID_SOURCE")
        headers = next(rows, None)
        if not headers or len(headers) > limits["max_columns"]:
            raise WorkbenchEngineError("表头为空或字段超限", "INVALID_SOURCE")
        fields = [
            {
                "id": f"c{i + 1:04d}",
                "name": str(v or f"未命名列{i + 1}"),
                "sourceName": str(v or ""),
                "position": i + 1,
                "type": "text",
            }
            for i, v in enumerate(headers)
        ]
        fields.append({"id": "_source_row", "name": "来源行号", "type": "integer"})
        count = 0
        with tempfile.TemporaryDirectory(prefix="workbench-import-", dir=limits.get("temp_directory")) as tmp:
            normalized = Path(tmp) / "normalized.csv"
            with normalized.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow([f["id"] for f in fields])
                for line, row in enumerate(rows, start=2):
                    if len(row) > len(headers):
                        raise WorkbenchEngineError(f"第{line}行字段超过表头数量", "INVALID_SOURCE")
                    count += 1
                    if count > limits["max_rows"]:
                        raise WorkbenchEngineError("数据行数超限，未发布部分结果", "ROW_LIMIT")
                    writer.writerow([_scalar(row[i]) if i < len(row) else None for i in range(len(headers))] + [line])
            with _connect(limits) as con:
                con.execute(
                    "CREATE TABLE imported AS SELECT * FROM read_csv(?, header=true, all_varchar=true, null_padding=false)",
                    [str(normalized)],
                )
                con.execute("COPY imported TO ? (FORMAT PARQUET)", [str(output_path)])
        return {
            "fields": fields,
            "rowCount": count,
            "quality": {
                "scope": "full",
                "duplicateHeaders": len(headers) - len(set(str(v) for v in headers)),
                "sourceRowsRetained": True,
                "formulaValues": "cached_only" if workbook else None,
                "note": "源字段按文本保留；空单元格为NULL，0保留；类型转换需显式步骤",
            },
        }
    finally:
        if workbook:
            workbook.close()
        if handle:
            handle.close()


def _decimal(column, scale=6):
    text = f"TRIM(CAST({column} AS VARCHAR))"
    pattern = r"^[+-]?([0-9]+|[0-9]{1,3}(,[0-9]{3})+)(\.[0-9]+)?$"
    precision = 28 if scale == 6 else 38
    return f"CASE WHEN REGEXP_FULL_MATCH({text}, '{pattern}') THEN TRY_CAST(REPLACE({text}, ',', '') AS DECIMAL({precision},{scale})) ELSE NULL END"


def _predicate(column, operator, value, params):
    if operator in ("is_null", "not_null"):
        return f"{column} IS {'NOT ' if operator == 'not_null' else ''}NULL"
    if operator == "contains":
        params.append(str(value or ""))
        return f"CONTAINS(CAST({column} AS VARCHAR), ?)"
    comparisons = {"eq": "=", "ne": "<>", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}
    if operator not in comparisons or value is None:
        raise WorkbenchEngineError("比较条件无效")
    params.append(str(value))
    return f"{column} {comparisons[operator]} ?"


def _execute(source, steps, resolve_source, output_path, limits=None, check_cancel=None):
    limits = _limits(limits)
    if not isinstance(steps, list) or len(steps) > 50:
        raise WorkbenchEngineError("最多允许50个加工步骤")
    fields = [dict(f) for f in source["fields"]]
    report = {"scope": "full", "steps": []}
    seen = set()
    division_scales = set()
    with _connect(limits) as con:
        # Integer quotient/remainder: exact six-place ROUND_HALF_UP, never DOUBLE.
        con.execute(
            "CREATE MACRO wb_ratio(x,y) AS CASE WHEN y=0 THEN NULL ELSE CAST((x // y + CASE WHEN ABS(x % y)*2 >= ABS(y) THEN SIGN(x)*SIGN(y) ELSE 0 END) AS DECIMAL(38,0))*0.000001 END"
        )
        con.execute(
            "CREATE MACRO wb_divide(a,b) AS wb_ratio(CAST(a*1000000 AS HUGEINT)*1000000,CAST(b*1000000 AS HUGEINT))"
        )
        con.execute("CREATE TABLE t0 AS SELECT * FROM read_parquet(?)", [str(source["path"])])
        right_tables = {}
        for i, step in enumerate(steps):
            if step.get("op") in ("join", "append"):
                right = resolve_source(step.get("rightSourceId"))
                name = f"right_{i}"
                con.execute(f"CREATE TABLE {name} AS SELECT * FROM read_parquet(?)", [str(right["path"])])
                right_tables[i] = (name, right["fields"])
        # All paths are loaded by trusted service; subsequent work cannot read files or URLs.
        con.execute("SET allowed_paths=?", [[str(Path(output_path).resolve())]])
        con.execute("SET enable_external_access=false")
        current = "t0"
        for i, step in enumerate(steps):
            _STEP_CONTEXT.set((i, report))
            if check_cancel:
                check_cancel()
            sid = step.get("id")
            _q(sid)
            if sid in seen:
                raise WorkbenchEngineError("步骤标识重复")
            seen.add(sid)
            op = step.get("op")
            known = {f["id"]: f for f in fields}

            def col(key, prefix=""):
                if key not in known:
                    raise WorkbenchEngineError(f"步骤{i + 1}引用的字段不存在")
                return prefix + _q(key)

            before = con.execute(f"SELECT COUNT(*) FROM {current}").fetchone()[0]
            detail = {"id": sid, "op": op, "inputRows": before}
            params = []
            projections = {f["id"]: _q(f["id"]) for f in fields}
            query = None
            if op == "trim":
                for key in step.get("columns", []):
                    projections[key] = f"TRIM(CAST({col(key)} AS VARCHAR))"
                    known[key]["type"] = "text"
            elif op == "replace":
                key = step.get("column")
                projections[key] = (
                    f"CASE WHEN CAST({col(key)} AS VARCHAR)=? THEN ? ELSE CAST({col(key)} AS VARCHAR) END"
                )
                params = [str(step.get("from", "")), str(step.get("to", ""))]
                known[key]["type"] = "text"
            elif op == "cast":
                key = step.get("column")
                c = col(key)
                kind = step.get("type")
                if kind == "decimal":
                    try:
                        scale = decimal_scale(step.get("scale", 6))
                    except ValueError as exc:
                        raise WorkbenchEngineError(str(exc)) from None
                    expression = _decimal(c, scale)
                    detail["decimalScale"] = scale
                    known[key]["scale"] = scale
                    detail["precisionChanges"] = con.execute(
                        f"SELECT COUNT(*) FROM {current} WHERE ({expression}) IS NOT NULL AND REGEXP_MATCHES(TRIM(CAST({c} AS VARCHAR)), '\\.[0-9]{{{scale}}}[0-9]*[1-9][0-9]*$')"
                    ).fetchone()[0]
                    source_row = '"_source_row"' if "_source_row" in known else "NULL"
                    changed = f"({expression}) IS NOT NULL AND REGEXP_MATCHES(TRIM(CAST({c} AS VARCHAR)), '\\.[0-9]{{{scale}}}[0-9]*[1-9][0-9]*$')"
                    detail["precisionSamples"] = [
                        {"sourceRow": _scalar(r[0]), "before": _scalar(r[1]), "after": _scalar(r[2])}
                        for r in con.execute(f"SELECT {source_row},{c},{expression} FROM {current} WHERE {changed} LIMIT 10").fetchall()
                    ]
                elif kind == "date":
                    expression = f"TRY_CAST(REPLACE(TRIM(CAST({c} AS VARCHAR)), '/', '-') AS DATE)"
                elif kind == "text":
                    expression = f"CAST({c} AS VARCHAR)"
                else:
                    raise WorkbenchEngineError("不支持的类型")
                projections[key] = expression
                detail["invalidValues"] = con.execute(
                    f"SELECT COUNT(*) FROM {current} WHERE {c} IS NOT NULL AND TRIM(CAST({c} AS VARCHAR))<>'' AND ({expression}) IS NULL"
                ).fetchone()[0]
                sample_id = '"_source_row"' if "_source_row" in known else "NULL"
                detail["invalidSamples"] = [
                    {"sourceRow": _scalar(r[0]), "value": _scalar(r[1])}
                    for r in con.execute(
                        f"SELECT {sample_id},{c} FROM {current} WHERE {c} IS NOT NULL AND TRIM(CAST({c} AS VARCHAR))<>'' AND ({expression}) IS NULL LIMIT 20"
                    ).fetchall()
                ]
                known[key]["type"] = kind
            elif op == "filter":
                predicate = _predicate(col(step.get("column")), step.get("operator"), step.get("value"), params)
                query = f"SELECT * FROM {current} WHERE {predicate}"
            elif op == "derive":
                key = "d_" + sid
                _q(key)
                if key in known:
                    raise WorkbenchEngineError("派生字段标识冲突")
                kind = step.get("kind")
                if kind == "formula":
                    try:
                        scale = decimal_scale(step.get("scale", 6))
                        if scale not in division_scales:
                            install_division(con, scale)
                            division_scales.add(scale)
                        expression, result_type, used, zero_checks = compile_formula(step.get("expression"), known, _q, scale)
                    except ValueError as exc:
                        raise WorkbenchEngineError(str(exc)) from None
                    projections[key] = expression
                    non_additive = bool(re.search(r"\bgroup_(sum|count)\s*\(", step.get("expression", ""))) or any(known[f].get("nonAdditive") for f in used)
                    fields.append({"id": key, "name": str(step.get("name") or key)[:200], "type": result_type, "scale": scale, "nonAdditive": non_additive})
                    detail.update(decimalScale=scale, rounding="ROUND_HALF_UP", nullPolicy="propagate_unless_coalesce")
                    if non_additive:
                        detail["warnings"] = ["分组特征在每条明细上重复，不提供逐行相加总额"]
                    if zero_checks:
                        zero_sql = " OR ".join(zero_checks)
                        detail["divisionByZero"] = con.execute(f"SELECT COUNT(*) FROM (SELECT ({zero_sql}) AS wb_zero FROM {current}) checked WHERE wb_zero").fetchone()[0]
                    names = [known[f]["name"] for f in used]
                    if any("原币" in n for n in names) and any("CNY" in n.upper() or "本币" in n for n in names):
                        detail.setdefault("warnings", []).append("计算混合原币与本币字段，请核对币种和单位；结果不代表已确认财务指标")
                    query = f"SELECT {','.join(f'{v} AS {_q(k)}' for k, v in projections.items())} FROM {current}"
                else:
                    left = col(step.get("left"))
                right = col(step["right"]) if step.get("right") else "?"
                if kind == "formula":
                    pass
                elif kind == "condition":
                    predicate = _predicate(left, step.get("operator"), step.get("value"), params)
                    expression = f"CASE WHEN {predicate} THEN ? ELSE ? END"
                    params.extend([step.get("then"), step.get("else")])
                elif kind == "concat":
                    expression = f"CAST({left} AS VARCHAR) || CAST({right} AS VARCHAR)"
                    if right == "?":
                        params.append(step.get("value"))
                elif kind in ("add", "subtract", "multiply", "divide"):
                    scale = decimal_scale(step.get("scale", max(known[step["left"]].get("scale", 6), known.get(step.get("right"), {}).get("scale", 6))))
                    if known[step["left"]]["type"] != "decimal" or (
                        step.get("right") and known[step["right"]]["type"] != "decimal"
                    ):
                        raise WorkbenchEngineError("数值计算前请先将参与字段显式转换为数值")
                    if not step.get("right"):
                        try:
                            literal = Decimal(str(step.get("value")))
                            if not literal.is_finite() or abs(literal) >= Decimal("1e22") or literal != literal.quantize(Decimal(1).scaleb(-scale)):
                                raise InvalidOperation
                        except (InvalidOperation, ValueError):
                            raise WorkbenchEngineError("数值常量无效或超出范围") from None
                    for operand in [left] + ([right] if right != "?" else []):
                        invalid = con.execute(
                            f"SELECT COUNT(*) FROM {current} WHERE {operand} IS NOT NULL AND {_decimal(operand, scale)} IS NULL"
                        ).fetchone()[0]
                        if invalid:
                            raise WorkbenchEngineError("参与计算的金额超出所选精度支持范围")
                    a, b = _decimal(left, scale), _decimal(right, scale)
                    if right == "?":
                        params.append(format(literal, "f"))
                    if kind == "divide":
                        if scale not in division_scales:
                            install_division(con, scale)
                            division_scales.add(scale)
                        expression = f"wb_formula_divide_{scale}({a},{b})"
                        detail["divisionByZero"] = con.execute(
                            f"SELECT COUNT(*) FROM {current} WHERE {b}=0", params
                        ).fetchone()[0]
                    else:
                        symbol = {"add": "+", "subtract": "-", "multiply": "*"}[kind]
                        expression = f"CAST(({a} {symbol} {b}) AS DECIMAL(38,{scale}))"
                    detail["decimalScale"] = scale
                else:
                    raise WorkbenchEngineError("派生算子无效")
                if kind != "formula":
                    projections[key] = expression
                    fields.append(
                        {
                            "id": key,
                            "name": str(step.get("name") or key)[:200],
                            "type": "text" if kind in ("condition", "concat") else "decimal",
                            **({"scale": scale} if kind in ("add", "subtract", "multiply", "divide") else {}),
                            "nonAdditive": bool(known.get(step.get("left"), {}).get("nonAdditive") or known.get(step.get("right"), {}).get("nonAdditive")),
                        }
                    )
            elif op == "select":
                keys = step.get("columns", [])
                if not keys or len(keys) != len(set(keys)):
                    raise WorkbenchEngineError("请选择不重复的字段")
                if "_source_row" in known and "_source_row" not in keys:
                    keys = [*keys, "_source_row"]
                query = f"SELECT {','.join(col(k) for k in keys)} FROM {current}"
                fields = [known[k] for k in keys]
            elif op == "group":
                keys = step.get("keys", [])
                if len(keys) != len(set(keys)):
                    raise WorkbenchEngineError("分组字段不能重复")
                selections = [col(k) for k in keys]
                next_fields = [known[k] for k in keys]
                for j, agg in enumerate(step.get("aggregations", [])):
                    c = col(agg.get("column"))
                    fn = agg.get("fn")
                    if fn not in ("sum", "count", "min", "max"):
                        raise WorkbenchEngineError("汇总函数无效")
                    if fn == "sum" and known[agg["column"]]["type"] != "decimal":
                        raise WorkbenchEngineError("金额求和前请显式转换为数值")
                    key = f"g_{sid}_{j}"
                    selections.append(f"{fn}({c}) AS {_q(key)}")
                    next_fields.append(
                        {
                            "id": key,
                            "name": str(agg.get("name") or key)[:200],
                            "type": "integer" if fn == "count" else known[agg["column"]]["type"],
                            **({"scale": known[agg["column"]].get("scale", 6)} if known[agg["column"]]["type"] == "decimal" else {}),
                        }
                    )
                if not selections:
                    raise WorkbenchEngineError("汇总字段为空")
                query = f"SELECT {','.join(selections)} FROM {current}" + (
                    f" GROUP BY {','.join(col(k) for k in keys)}" if keys else ""
                )
                fields = next_fields
            elif op == "join":
                rt, rf = right_tables[i]
                rk = {f["id"]: f for f in rf}
                lk, rkeys = step.get("leftKeys", []), step.get("rightKeys", [])
                how = step.get("how")
                if not lk or len(lk) != len(rkeys) or how not in ("left", "inner", "full", "anti"):
                    raise WorkbenchEngineError("关联参数无效")
                for k in rkeys:
                    if k not in rk:
                        raise WorkbenchEngineError("右表关联字段不存在")
                lcols, rcols = [col(k) for k in lk], [_q(k) for k in rkeys]
                pairs = " AND ".join(f"l.{a}=r.{b}" for a, b in zip(lcols, rcols))
                lg = ",".join(lcols)
                rg = ",".join(rcols)
                cte = f"WITH l AS (SELECT {lg},COUNT(*) n FROM {current} GROUP BY {lg}), r AS (SELECT {rg},COUNT(*) n FROM {rt} GROUP BY {rg}) "
                stats = con.execute(
                    cte
                    + f"SELECT COALESCE(SUM(CASE WHEN l.n IS NOT NULL AND r.n IS NOT NULL THEN CAST(l.n AS HUGEINT)*r.n ELSE 0 END),0), COALESCE(SUM(CASE WHEN r.n IS NULL THEN l.n ELSE 0 END),0),COALESCE(SUM(CASE WHEN l.n IS NULL THEN r.n ELSE 0 END),0),COALESCE(MAX(CASE WHEN l.n>1 AND r.n>1 THEN 1 ELSE 0 END),0) FROM l FULL JOIN r ON {pairs}"
                ).fetchone()
                matched, unmatched_left, unmatched_right, many = map(int, stats)
                estimated = {
                    "inner": matched,
                    "left": matched + unmatched_left,
                    "full": matched + unmatched_left + unmatched_right,
                    "anti": unmatched_left,
                }[how]
                expansion = estimated / max(before, 1)
                detail["join"] = {
                    "estimatedRows": estimated,
                    "expansion": expansion,
                    "manyToMany": bool(many),
                    "unmatchedLeft": unmatched_left,
                    "unmatchedRight": unmatched_right,
                    "nullsMatch": False,
                    "expansionRatio": expansion,
                }
                samples = {}
                for side, table, keys in (("left", current, lcols), ("right", rt, rcols)):
                    group_sql = f"SELECT {','.join(keys)},COUNT(*) n FROM {table} GROUP BY {','.join(keys)} HAVING COUNT(*)>1"
                    detail["join"][side + "DuplicateKeys"] = con.execute(f"SELECT COUNT(*) FROM ({group_sql}) duplicates").fetchone()[0]
                    samples[side] = [{"keys": [_scalar(v) for v in row[:-1]], "count": row[-1]} for row in con.execute(group_sql + " ORDER BY n DESC LIMIT 5").fetchall()]
                detail["join"]["duplicateSamples"] = samples
                if estimated > limits["max_join_rows"] or expansion > limits["max_join_expansion"]:
                    report["steps"].append(detail)
                    raise WorkbenchEngineError("关联结果规模超限；请增加关联字段或先分组汇总", "JOIN_LIMIT", report)
                selections = [f"l.{_q(f['id'])}" for f in fields]
                if how != "anti":
                    for f in rf:
                        key = f"r_{sid}_{f['id']}"
                        selections.append(f"r.{_q(f['id'])} AS {_q(key)}")
                        fields.append({**f, "id": key, "name": "右表·" + f["name"]})
                query = f"SELECT {','.join(selections)} FROM {current} l {how.upper()} JOIN {rt} r ON {pairs}"
            elif op == "append":
                rt, rf = right_tables[i]
                rk = {f["id"]: f for f in rf}
                mapping = step.get("mapping", {})
                if set(mapping) != set(known):
                    raise WorkbenchEngineError("追加必须显式映射所有左表字段")
                values = []
                for key, field in known.items():
                    right_key = mapping[key]
                    if right_key not in rk or rk[right_key]["type"] != field["type"]:
                        raise WorkbenchEngineError("追加字段不存在或类型不兼容")
                    values.append(f"{_q(right_key)} AS {_q(key)}")
                query = f"SELECT * FROM {current} UNION ALL SELECT {','.join(values)} FROM {rt}"
            else:
                raise WorkbenchEngineError("不支持的加工算子")
            if len(fields) > limits["max_result_columns"]:
                raise WorkbenchEngineError("结果字段数量超限，请先选择需要的字段", "COLUMN_LIMIT")
            if len({f["id"] for f in fields}) != len(fields):
                raise WorkbenchEngineError("结果字段标识冲突，请修改步骤标识")
            if query is None:
                query = f"SELECT {','.join(f'{v} AS {_q(k)}' for k, v in projections.items())} FROM {current}"
            target = f"t{i + 1}"
            try:
                con.execute(f"CREATE TABLE {target} AS {query}", params)
            except Exception as exc:
                raise WorkbenchEngineError(
                    f"步骤{i + 1}执行失败，请检查字段类型、数值范围与条件", "EXECUTION_FAILED", report
                ) from exc
            after = con.execute(f"SELECT COUNT(*) FROM {target}").fetchone()[0]
            detail["outputRows"] = after
            report["steps"].append(detail)
            if after > limits["max_output_rows"]:
                raise WorkbenchEngineError("结果行数超限", "ROW_LIMIT", report)
            con.execute(f"DROP TABLE {current}")
            current = target
        if check_cancel:
            check_cancel()
        count = con.execute(f"SELECT COUNT(*) FROM {current}").fetchone()[0]
        report["rowCount"] = count
        report["totals"] = {
            f["id"]: _scalar(con.execute(f"SELECT SUM({_q(f['id'])}) FROM {current}").fetchone()[0])
            for f in fields
            if f["type"] == "decimal" and not f.get("nonAdditive")
        }
        con.execute(f"COPY {current} TO ? (FORMAT PARQUET)", [str(Path(output_path).resolve())])
        return {"fields": fields, "rowCount": count, "report": report}


def execute(source, steps, resolve_source, output_path, limits=None, check_cancel=None):
    """Public boundary returns safe errors without driver SQL, paths or stack traces."""
    context_token = _STEP_CONTEXT.set(None)
    try:
        return _execute(source, steps, resolve_source, output_path, limits, check_cancel)
    except WorkbenchEngineError:
        raise
    except (KeyError, TypeError, IndexError, AttributeError) as exc:
        raise WorkbenchEngineError("步骤参数不完整或字段结构无效", "INVALID_STEP") from exc
    except Exception as exc:
        # Cancellation exceptions belong to the orchestration layer.
        if check_cancel:
            check_cancel()
        raise WorkbenchEngineError("计算失败，请检查字段类型、资源限制和步骤参数", "EXECUTION_FAILED") from exc
    finally:
        _STEP_CONTEXT.reset(context_token)


def preview(path, fields, offset=0, limit=100, scope="source", report=None):
    offset, limit = max(0, int(offset)), min(200, max(1, int(limit)))
    with _connect(_limits(None)) as con:
        total = con.execute("SELECT COUNT(*) FROM read_parquet(?)", [str(path)]).fetchone()[0]
        values = con.execute("SELECT * FROM read_parquet(?) LIMIT ? OFFSET ?", [str(path), limit, offset]).fetchall()
    return {
        "fields": fields,
        "rows": [{f["id"]: _scalar(v) for f, v in zip(fields, row)} for row in values],
        "total": total,
        "offset": offset,
        "limit": limit,
        "scope": scope,
        "report": report,
    }


def export_csv(path, fields):
    """Yield bounded text chunks; neutralize spreadsheet formula injection."""

    def safe(value):
        value = _scalar(value)
        if value is not None and value.lstrip().startswith(("=", "+", "-", "@")):
            # Preserve actual negative numeric cells; protect untrusted text only.
            try:
                Decimal(value)
            except InvalidOperation:
                value = "'" + value
        return value

    with _connect(_limits(None)) as con:
        cursor = con.execute("SELECT * FROM read_parquet(?)", [str(path)])
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow([safe(f["name"]) for f in fields])
        yield "\ufeff" + stream.getvalue()
        while rows := cursor.fetchmany(1000):
            stream.seek(0)
            stream.truncate(0)
            writer.writerows([[safe(v) for v in row] for row in rows])
            yield stream.getvalue()
