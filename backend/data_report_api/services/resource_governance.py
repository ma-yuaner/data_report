from __future__ import annotations

import logging
import json
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from tempfile import SpooledTemporaryFile
from typing import Any, BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..time_utils import business_now
from .data_source import DataSource


REQUEST_TABLE_NAME = "bi_ota_gds_request_day"
DEFAULT_GDS_METRIC_TABLE = "bi_gds_resource_metric_day"
GDS_MP_SCOPE = "MP_TOTAL"
GDS_API_GROUP_SCOPE = "API_GROUP"
DATA_SOURCES = {
    "ota": {"value": 0, "label": "OTA"},
    "gds": {"value": 1, "label": "GDS"},
}
DIMENSIONS = {
    "platform": {"field": "ota_code", "label": "平台"},
    "site": {"field": "ota_site_code", "label": "站点"},
    "gds": {"field": "gds_code", "label": "GDS"},
    "pcc": {"field": "pcc_code", "label": "PCC"},
    "apiGroup": {"field": "source_api_group", "label": "接口分组"},
}
FILTER_FIELDS = {key: value["field"] for key, value in DIMENSIONS.items()}
METRIC_SELECT = """
COUNT(1),
SUM(search_request_count_all), SUM(search_request_success),
SUM(verify_request_count_all), SUM(verify_request_success),
SUM(order_request_count_all), SUM(order_request_success),
{segment_expression}, {segment_missing_expression},
SUM(search_cache_count), SUM(verify_cache_count), SUM(order_cache_count),
SUM(search_execute_time), SUM(verify_execute_time), SUM(order_execute_time),
SUM(search_price_change_count), SUM(verify_price_change_count), SUM(order_price_change_count)
""".strip()

REPORT_COLUMNS = (
    ("查询请求量", "searchRequests"),
    ("查询成功量", "searchSuccess"),
    ("查询失败量", "searchFailure"),
    ("查询成功率", "searchSuccessRate"),
    ("查询失败率", "searchFailureRate"),
    ("验价请求量", "verifyRequests"),
    ("验价成功量", "verifySuccess"),
    ("验价失败量", "verifyFailure"),
    ("验价成功率", "verifySuccessRate"),
    ("验价失败率", "verifyFailureRate"),
    ("生单请求量", "orderRequests"),
    ("生单成功量", "orderSuccess"),
    ("生单失败量", "orderFailure"),
    ("生单成功率", "orderSuccessRate"),
    ("生单失败率", "orderFailureRate"),
)
REPORT_MEMORY_LIMIT_BYTES = 8 * 1024 * 1024


def _report_metric_value(metric: dict[str, Any], key: str) -> Any:
    value = metric.get(key)
    if key.endswith("Rate"):
        return None if value is None else f"{value}%"
    return value


def _append_report_sheet(
    workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    header_fill = PatternFill("solid", fgColor="17375F")
    for index, value in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=index, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, 1):
        max_length = max(
            [len(str(header))]
            + [len(str(row[index - 1])) for row in rows if row[index - 1] is not None]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 34)


def _governance_workbook(result: dict[str, Any]) -> BinaryIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = result["summary"]
    filters = result["filters"]
    layer = result["dataSource"]["label"]
    filter_label = "；".join(
        f"{DIMENSIONS[key]['label']}={value}" for key, value in filters.items() if value
    ) or "全部"
    summary_rows = [
        ["数据层", layer],
        ["统计期间", f"{result['period']['startDate']} 至 {result['period']['endDate']}"],
        ["筛选条件", filter_label],
        ["数据来源", result["table"]],
        ["数据更新时间", summary.get("sourceUpdatedAt") or "—"],
        ["生成时间", result["generatedAt"]],
    ]
    if result.get("ratioApplicable"):
        ratio = "∞" if summary.get("ratioInfinite") else summary.get("queryToSegmentRatio")
        summary_rows.extend([
            ["MP查询量", summary.get("searchRequests")],
            ["航段数", summary.get("segmentCount")],
            ["查定比", "—" if ratio is None else f"{ratio}:1"],
            ["治理目标", "未配置" if summary.get("targetRatio") is None else f"{summary['targetRatio']}:1"],
            ["超额查询量", summary.get("excessQueryCount")],
        ])
    summary_rows.extend(
        [[label, _report_metric_value(summary, key)] for label, key in REPORT_COLUMNS]
    )
    summary_rows.extend([
        ["异常判定", "异常" if summary.get("breach") else "正常/未配置目标"],
        ["异常原因", "；".join(summary.get("breachReasons") or []) or "—"],
    ])
    _append_report_sheet(workbook, "管理摘要", ["项目", "数值"], summary_rows)

    trend_headers = ["日期"]
    if result.get("ratioApplicable"):
        trend_headers.extend(["MP查询量", "航段数", "查定比", "治理目标", "超额查询量"])
    trend_metric_columns = REPORT_COLUMNS[1:] if result.get("ratioApplicable") else REPORT_COLUMNS
    trend_headers.extend(label for label, _ in trend_metric_columns)
    trend_rows: list[list[Any]] = []
    for metric in result.get("trend", []):
        row: list[Any] = [metric.get("date")]
        if result.get("ratioApplicable"):
            ratio = "∞" if metric.get("ratioInfinite") else metric.get("queryToSegmentRatio")
            row.extend([
                metric.get("searchRequests"), metric.get("segmentCount"), ratio,
                metric.get("targetRatio"), metric.get("excessQueryCount"),
            ])
        row.extend(_report_metric_value(metric, key) for _, key in trend_metric_columns)
        trend_rows.append(row)
    _append_report_sheet(workbook, "日趋势", trend_headers, trend_rows)

    dimension_headers = [result["dimension"]["label"]]
    if result.get("ratioApplicable"):
        dimension_headers.extend(["MP查询量", "航段数", "查定比", "治理目标", "超额查询量"])
    dimension_metric_columns = REPORT_COLUMNS[1:] if result.get("ratioApplicable") else REPORT_COLUMNS
    dimension_headers.extend(label for label, _ in dimension_metric_columns)
    dimension_headers.extend(["是否异常", "异常原因"])
    dimension_rows: list[list[Any]] = []
    for metric in result["dimension"].get("rows", []):
        row = [metric.get("name") or metric.get("value")]
        if result.get("ratioApplicable"):
            ratio = "∞" if metric.get("ratioInfinite") else metric.get("queryToSegmentRatio")
            row.extend([
                metric.get("searchRequests"), metric.get("segmentCount"), ratio,
                metric.get("targetRatio"), metric.get("excessQueryCount"),
            ])
        row.extend(_report_metric_value(metric, key) for _, key in dimension_metric_columns)
        row.extend([
            "是" if metric.get("breach") else "否",
            "；".join(metric.get("breachReasons") or []),
        ])
        dimension_rows.append(row)
    _append_report_sheet(workbook, "维度对比", dimension_headers, dimension_rows)

    if result["dataSource"]["key"] == "gds":
        group_headers = ["接口分组", *(label for label, _ in REPORT_COLUMNS)]
        group_rows = [
            [metric.get("name") or metric.get("value"), *(
                _report_metric_value(metric, key) for _, key in REPORT_COLUMNS
            )]
            for metric in result.get("apiGroups", [])
        ]
        _append_report_sheet(workbook, "接口分组质量", group_headers, group_rows)

    output = SpooledTemporaryFile(max_size=REPORT_MEMORY_LIMIT_BYTES, mode="w+b")
    workbook.save(output)
    output.seek(0)
    return output


def _parse_date(value: str | None, fallback: date) -> date:
    if not value:
        return fallback
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise ValueError("日期格式必须为YYYY-MM-DD") from error


def _clean_filter(value: Any) -> str:
    cleaned = str(value or "").strip()
    if len(cleaned) > 100:
        raise ValueError("筛选值长度不能超过100个字符")
    return cleaned


def _safe_table_name(value: Any, fallback: str) -> str:
    table = str(value or fallback).strip()
    if not re.fullmatch(r"[A-Za-z0-9_]+", table):
        raise ValueError("资源治理表名配置不合法")
    return table


def _rate(numerator: int, denominator: int) -> str | None:
    if denominator <= 0:
        return None
    value = Decimal(numerator) * Decimal(100) / Decimal(denominator)
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _ratio(numerator: int, denominator: int) -> str | None:
    if denominator <= 0:
        return None
    value = Decimal(numerator) / Decimal(denominator)
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _date_label(value: Any) -> str:
    raw = str(value or "")
    if len(raw) >= 8 and raw[:8].isdigit():
        return f"{raw[:4]}-{raw[4:6]}-{raw[6:8]}"
    return raw


def _parse_gds_targets(raw: Any) -> dict[str, int]:
    targets: dict[str, int] = {}
    for item in str(raw or "").split(","):
        code, separator, value = item.strip().partition(":")
        if not separator or not code.strip():
            continue
        try:
            ratio = int(value.strip())
        except ValueError:
            continue
        if ratio > 0:
            targets[code.strip().upper()] = ratio
    return targets


def _parse_ota_targets(raw: Any) -> dict[str, dict[str, str]]:
    if isinstance(raw, dict):
        values = raw
    else:
        try:
            values = json.loads(str(raw or "{}"))
        except (TypeError, ValueError, json.JSONDecodeError):
            return {}
    if not isinstance(values, dict):
        return {}
    targets: dict[str, dict[str, str]] = {}
    for platform, stages in values.items():
        if not isinstance(stages, dict) or not str(platform).strip():
            continue
        parsed: dict[str, str] = {}
        for stage in ("search", "verify", "order"):
            try:
                target = Decimal(str(stages.get(stage)))
            except Exception:
                continue
            if Decimal(0) <= target <= Decimal(100):
                parsed[stage] = str(target.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        if parsed:
            targets[str(platform).strip()] = parsed
    return targets


def _metric(
    row: tuple[Any, ...],
    segment_column_available: bool,
    target_ratio: int | None,
    ratio_applicable: bool = True,
    success_targets: dict[str, str] | None = None,
) -> dict[str, Any]:
    row_count = int(row[0] or 0)
    search_requests = int(row[1] or 0)
    search_success = int(row[2] or 0)
    verify_requests = int(row[3] or 0)
    verify_success = int(row[4] or 0)
    order_requests = int(row[5] or 0)
    order_success = int(row[6] or 0)
    search_failure = max(search_requests - search_success, 0)
    verify_failure = max(verify_requests - verify_success, 0)
    order_failure = max(order_requests - order_success, 0)
    segment_count = int(row[7]) if ratio_applicable and row[7] is not None else None
    segment_missing_rows = int(row[8] or 0) if ratio_applicable and row[8] is not None else None
    search_cache = int(row[9] or 0)
    verify_cache = int(row[10] or 0)
    order_cache = int(row[11] or 0)
    search_execute_time = int(row[12] or 0)
    verify_execute_time = int(row[13] or 0)
    order_execute_time = int(row[14] or 0)
    search_price_change = int(row[15] or 0)
    verify_price_change = int(row[16] or 0)
    order_price_change = int(row[17] or 0)

    if row_count == 0:
        status = "no_records"
    elif not ratio_applicable:
        status = "not_applicable"
    elif not segment_column_available:
        status = "segment_field_missing"
    elif segment_count is None:
        status = "segment_data_missing"
    elif segment_missing_rows:
        status = "segment_partial"
    elif segment_count == 0 and search_requests > 0:
        status = "denominator_zero"
    else:
        status = "ready"

    segment_complete = status in {"ready", "denominator_zero"}
    excess = None if not segment_complete or target_ratio is None else max(search_requests - segment_count * target_ratio, 0)
    breach_reasons: list[str] = []
    targets = success_targets or {}
    if segment_complete and target_ratio is not None and (search_requests > 0 or segment_count is not None):
        if segment_count == 0 and search_requests > 0:
            breach_reasons.append(f"查定比∞，超过治理目标{target_ratio}:1")
        elif segment_count is not None and search_requests > segment_count * target_ratio:
            breach_reasons.append(f"查定比超过治理目标{target_ratio}:1")
    if row_count > 0:
        for stage, label in (("search", "查询"), ("verify", "验价"), ("order", "生单")):
            target = targets.get(stage)
            if target is None:
                continue
            actual = _rate(
                search_success if stage == "search" else verify_success if stage == "verify" else order_success,
                search_requests if stage == "search" else verify_requests if stage == "verify" else order_requests,
            )
            if actual is not None and Decimal(actual) < Decimal(str(target)):
                breach_reasons.append(f"{label}成功率{actual}%低于目标{target}%")
    return {
        "rowCount": row_count,
        "searchRequests": search_requests,
        "searchSuccess": search_success,
        "searchFailure": search_failure,
        "verifyRequests": verify_requests,
        "verifySuccess": verify_success,
        "verifyFailure": verify_failure,
        "orderRequests": order_requests,
        "orderSuccess": order_success,
        "orderFailure": order_failure,
        "segmentCount": segment_count,
        "segmentMissingRowCount": segment_missing_rows,
        "queryToSegmentRatio": _ratio(search_requests, segment_count or 0) if segment_complete else None,
        "ratioInfinite": bool(segment_complete and search_requests > 0 and segment_count == 0),
        "searchSuccessRate": _rate(search_success, search_requests),
        "searchFailureRate": _rate(search_failure, search_requests),
        "verifySuccessRate": _rate(verify_success, verify_requests),
        "verifyFailureRate": _rate(verify_failure, verify_requests),
        "orderSuccessRate": _rate(order_success, order_requests),
        "orderFailureRate": _rate(order_failure, order_requests),
        "searchCacheRate": _rate(search_cache, search_requests),
        "verifyCacheRate": _rate(verify_cache, verify_requests),
        "orderCacheRate": _rate(order_cache, order_requests),
        "avgSearchExecuteTime": _ratio(search_execute_time, search_requests),
        "avgVerifyExecuteTime": _ratio(verify_execute_time, verify_requests),
        "avgOrderExecuteTime": _ratio(order_execute_time, order_requests),
        "searchPriceChangeRate": _rate(search_price_change, search_success),
        "verifyPriceChangeRate": _rate(verify_price_change, verify_success),
        "orderPriceChangeRate": _rate(order_price_change, order_success),
        "excessQueryCount": excess,
        "targetRatio": target_ratio,
        "successTargets": success_targets or {},
        "status": status,
        "breach": bool(breach_reasons),
        "breachReasons": breach_reasons,
    }


class ResourceGovernanceService:
    def __init__(self, config: dict[str, Any]):
        mysql_config = dict(config)
        mysql_config["DATA_MODE"] = "mysql"
        self.source = DataSource(mysql_config)
        self.request_table = REQUEST_TABLE_NAME
        self.gds_metric_table = _safe_table_name(
            config.get("RESOURCE_GOVERNANCE_GDS_TABLE"), DEFAULT_GDS_METRIC_TABLE
        )
        self.gds_targets = _parse_gds_targets(config.get("RESOURCE_GOVERNANCE_GDS_TARGETS", "1A:4000"))
        self.ota_targets = _parse_ota_targets(config.get("RESOURCE_GOVERNANCE_OTA_SUCCESS_TARGETS", "{}"))

    @staticmethod
    def _where(
        first: date,
        last: date,
        cleaned_filters: dict[str, str],
        *,
        source_value: int | None = None,
        metric_scope: str | None = None,
        include_api_group: bool = True,
        iso_dates: bool = False,
    ) -> tuple[str, list[Any]]:
        clauses = ["operator_date >= %s", "operator_date <= %s"]
        parameters: list[Any] = [
            first.isoformat() if iso_dates else first.strftime("%Y%m%d"),
            last.isoformat() if iso_dates else last.strftime("%Y%m%d"),
        ]
        if source_value is not None:
            clauses.append("data_source = %s")
            parameters.append(source_value)
        if metric_scope:
            clauses.append("metric_scope = %s")
            parameters.append(metric_scope)
        for key, value in cleaned_filters.items():
            if value and (key != "apiGroup" or include_api_group):
                clauses.append(f"{FILTER_FIELDS[key]} = %s")
                parameters.append(value)
        return " AND ".join(clauses), parameters

    def _populate_gds(
        self,
        cursor: Any,
        result: dict[str, Any],
        first: date,
        last: date,
        dimension_key: str,
        cleaned_filters: dict[str, str],
        selected_target: int | None,
        breach_only: bool,
    ) -> None:
        table = f"{self.source.database}.{self.gds_metric_table}"
        result["table"] = table
        result["apiGroups"] = []

        # MP_TOTAL是唯一允许携带航段分母的口径；选择接口组后只分析该组接口质量。
        summary_scope = GDS_API_GROUP_SCOPE if cleaned_filters["apiGroup"] else GDS_MP_SCOPE
        summary_ratio_applicable = summary_scope == GDS_MP_SCOPE
        summary_select = METRIC_SELECT.format(
            segment_expression="SUM(seg_num)" if summary_ratio_applicable else "NULL",
            segment_missing_expression=(
                "SUM(CASE WHEN segment_missing_flag = 1 OR seg_num IS NULL THEN 1 ELSE 0 END)"
                if summary_ratio_applicable
                else "NULL"
            ),
        )
        summary_where, summary_params = self._where(
            first,
            last,
            cleaned_filters,
            metric_scope=summary_scope,
            include_api_group=not summary_ratio_applicable,
            iso_dates=True,
        )

        option_where, option_params = self._where(
            first, last, {key: "" for key in FILTER_FIELDS}, iso_dates=True
        )
        cursor.execute(
            "SELECT DISTINCT ota_code, ota_site_code, gds_code, pcc_code, source_api_group "
            f"FROM {table} WHERE {option_where} LIMIT 5000",
            tuple(option_params),
        )
        option_rows = cursor.fetchall()
        for index, key in enumerate(("platform", "site", "gds", "pcc", "apiGroup")):
            values = {
                str(row[index]).strip()
                for row in option_rows
                if row[index] is not None and str(row[index]).strip()
            }
            if key == "apiGroup":
                values.discard(GDS_MP_SCOPE)
            result["options"][key] = sorted(values)

        cursor.execute(
            f"SELECT {summary_select}, MAX(etl_updated_at) FROM {table} WHERE {summary_where}",
            tuple(summary_params),
        )
        summary_row = cursor.fetchone() or ((0,) + (None,) * 17 + (None,))
        summary = _metric(
            tuple(summary_row[:18]),
            True,
            selected_target if summary_ratio_applicable else None,
            summary_ratio_applicable,
        )
        summary["sourceUpdatedAt"] = (
            summary_row[18].isoformat(sep=" ", timespec="seconds")
            if isinstance(summary_row[18], datetime)
            else (str(summary_row[18]) if summary_row[18] else None)
        )
        result["summary"] = summary
        result["ratioApplicable"] = summary_ratio_applicable

        cursor.execute(
            f"SELECT operator_date, {summary_select} FROM {table} WHERE {summary_where} "
            "GROUP BY operator_date ORDER BY operator_date",
            tuple(summary_params),
        )
        result["trend"] = [
            {
                "date": _date_label(row[0]),
                **_metric(
                    tuple(row[1:]),
                    True,
                    selected_target if summary_ratio_applicable else None,
                    summary_ratio_applicable,
                ),
            }
            for row in cursor.fetchall()
        ]

        dimension_scope = (
            GDS_API_GROUP_SCOPE
            if dimension_key == "apiGroup" or cleaned_filters["apiGroup"]
            else GDS_MP_SCOPE
        )
        dimension_ratio_applicable = dimension_scope == GDS_MP_SCOPE
        dimension_select = METRIC_SELECT.format(
            segment_expression="SUM(seg_num)" if dimension_ratio_applicable else "NULL",
            segment_missing_expression=(
                "SUM(CASE WHEN segment_missing_flag = 1 OR seg_num IS NULL THEN 1 ELSE 0 END)"
                if dimension_ratio_applicable
                else "NULL"
            ),
        )
        dimension_where, dimension_params = self._where(
            first,
            last,
            cleaned_filters,
            metric_scope=dimension_scope,
            include_api_group=dimension_scope == GDS_API_GROUP_SCOPE,
            iso_dates=True,
        )
        dimension = DIMENSIONS[dimension_key]
        dimension_field = dimension["field"]
        cursor.execute(
            f"SELECT {dimension_field}, {dimension_select} FROM {table} WHERE {dimension_where} "
            f"AND {dimension_field} IS NOT NULL AND {dimension_field} <> '' GROUP BY {dimension_field} "
            "ORDER BY SUM(search_request_count_all) DESC LIMIT 500",
            tuple(dimension_params),
        )
        dimension_rows = []
        for row in cursor.fetchall():
            value = str(row[0])
            row_target = (
                self.gds_targets.get(value.upper())
                if dimension_ratio_applicable and dimension_key == "gds"
                else selected_target if dimension_ratio_applicable else None
            )
            dimension_rows.append(
                {
                    "value": value,
                    "name": value,
                    **_metric(tuple(row[1:]), True, row_target, dimension_ratio_applicable),
                }
            )
        breach_rows = [row for row in dimension_rows if row.get("breach")]
        result["breachObjectCount"] = len(breach_rows)
        result["dimension"]["rows"] = breach_rows if breach_only else dimension_rows

        group_where, group_params = self._where(
            first,
            last,
            cleaned_filters,
            metric_scope=GDS_API_GROUP_SCOPE,
            include_api_group=True,
            iso_dates=True,
        )
        group_select = METRIC_SELECT.format(
            segment_expression="NULL", segment_missing_expression="NULL"
        )
        cursor.execute(
            f"SELECT source_api_group, {group_select} FROM {table} WHERE {group_where} "
            "AND source_api_group IS NOT NULL AND source_api_group <> '' "
            "GROUP BY source_api_group ORDER BY SUM(search_request_count_all) DESC LIMIT 500",
            tuple(group_params),
        )
        result["apiGroups"] = [
            {
                "value": str(row[0]),
                "name": str(row[0]),
                **_metric(tuple(row[1:]), True, None, False),
            }
            for row in cursor.fetchall()
        ]
        result["segmentColumnAvailable"] = True
        result["segmentDataAvailable"] = (
            summary_ratio_applicable and summary["status"] in {"ready", "denominator_zero"}
        )
        result["available"] = True

    def _populate_ota(
        self,
        cursor: Any,
        result: dict[str, Any],
        first: date,
        last: date,
        dimension_key: str,
        cleaned_filters: dict[str, str],
        selected_success_targets: dict[str, str],
        breach_only: bool,
    ) -> None:
        table = f"{self.source.database}.{self.request_table}"
        result["table"] = table
        result["apiGroups"] = []
        metric_select = METRIC_SELECT.format(
            segment_expression="NULL", segment_missing_expression="NULL"
        )
        base_where, base_params = self._where(
            first,
            last,
            cleaned_filters,
            source_value=DATA_SOURCES["ota"]["value"],
        )
        option_where, option_params = self._where(
            first,
            last,
            {key: "" for key in FILTER_FIELDS},
            source_value=DATA_SOURCES["ota"]["value"],
        )
        cursor.execute(
            "SELECT DISTINCT ota_code, ota_site_code, gds_code, pcc_code, source_api_group "
            f"FROM {table} WHERE {option_where} LIMIT 5000",
            tuple(option_params),
        )
        option_rows = cursor.fetchall()
        for index, key in enumerate(("platform", "site", "gds", "pcc", "apiGroup")):
            result["options"][key] = sorted(
                {
                    str(row[index]).strip()
                    for row in option_rows
                    if row[index] is not None and str(row[index]).strip()
                }
            )

        cursor.execute(
            f"SELECT {metric_select} FROM {table} WHERE {base_where}", tuple(base_params)
        )
        summary_row = cursor.fetchone() or ((0,) + (None,) * 17)
        summary = _metric(
            tuple(summary_row[:18]), False, None, False, selected_success_targets
        )
        summary["sourceUpdatedAt"] = None
        result["summary"] = summary
        cursor.execute(
            f"SELECT operator_date, {metric_select} FROM {table} WHERE {base_where} "
            "GROUP BY operator_date ORDER BY operator_date",
            tuple(base_params),
        )
        result["trend"] = [
            {
                "date": _date_label(row[0]),
                **_metric(tuple(row[1:]), False, None, False, selected_success_targets),
            }
            for row in cursor.fetchall()
        ]

        dimension = DIMENSIONS[dimension_key]
        dimension_field = dimension["field"]
        cursor.execute(
            f"SELECT {dimension_field}, {metric_select} FROM {table} WHERE {base_where} "
            f"AND {dimension_field} IS NOT NULL AND {dimension_field} <> '' GROUP BY {dimension_field} "
            "ORDER BY SUM(search_request_count_all) DESC LIMIT 500",
            tuple(base_params),
        )
        dimension_rows = []
        for row in cursor.fetchall():
            value = str(row[0])
            row_targets = (
                self.ota_targets.get(value, {})
                if dimension_key == "platform"
                else selected_success_targets
            )
            dimension_rows.append(
                {
                    "value": value,
                    "name": value,
                    **_metric(tuple(row[1:]), False, None, False, row_targets),
                }
            )
        breach_rows = [row for row in dimension_rows if row.get("breach")]
        result["breachObjectCount"] = len(breach_rows)
        result["dimension"]["rows"] = breach_rows if breach_only else dimension_rows
        result["available"] = True

    def overview(
        self,
        start_value: str | None = None,
        end_value: str | None = None,
        data_source: str = "gds",
        group_by: str | None = None,
        filters: dict[str, Any] | None = None,
        breach_only: bool = False,
    ) -> dict[str, Any]:
        complete_day = business_now().date() - timedelta(days=1)
        first = _parse_date(start_value, complete_day.replace(day=1))
        last = _parse_date(end_value, complete_day)
        if first > last:
            raise ValueError("开始日期不能晚于结束日期")
        if last > complete_day:
            raise ValueError(
                f"日级资源数据最多查询至{complete_day.isoformat()}，今日数据待小时表接入"
            )
        if (last - first).days > 366:
            raise ValueError("单次最多查询367天")
        source_key = str(data_source or "gds").strip().lower()
        if source_key not in DATA_SOURCES:
            raise ValueError("数据层只支持OTA或GDS")
        default_group = "gds" if source_key == "gds" else "platform"
        dimension_key = str(group_by or default_group).strip()
        if dimension_key not in DIMENSIONS:
            raise ValueError("不支持的分析维度")
        if source_key == "ota" and dimension_key in {"gds", "pcc"}:
            dimension_key = "platform"
        cleaned_filters = {
            key: _clean_filter((filters or {}).get(key)) for key in FILTER_FIELDS
        }
        ratio_applicable = source_key == "gds" and not cleaned_filters["apiGroup"]
        selected_target = self.gds_targets.get(cleaned_filters["gds"].upper()) if source_key == "gds" and cleaned_filters["gds"] else None
        selected_success_targets = self.ota_targets.get(cleaned_filters["platform"], {}) if source_key == "ota" and cleaned_filters["platform"] else {}

        result = {
            "source": self.source.label,
            "table": f"{self.source.database}.{self.gds_metric_table if source_key == 'gds' else self.request_table}",
            "available": False,
            "error": "",
            "generatedAt": datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds"),
            "period": {"startDate": first.isoformat(), "endDate": last.isoformat()},
            "dataSource": {"key": source_key, **DATA_SOURCES[source_key]},
            "targetRatio": selected_target,
            "gdsTargets": self.gds_targets,
            "otaSuccessTargets": self.ota_targets,
            "ratioApplicable": ratio_applicable,
            "segmentColumnAvailable": False,
            "segmentDataAvailable": False,
            "breachOnly": bool(breach_only),
            "breachObjectCount": 0,
            "filters": cleaned_filters,
            "options": {key: [] for key in FILTER_FIELDS},
            "summary": _metric((0,) + (None,) * 17, False, selected_target, ratio_applicable, selected_success_targets),
            "trend": [],
            "apiGroups": [],
            "dimension": {"key": dimension_key, "label": DIMENSIONS[dimension_key]["label"], "rows": []},
            "notes": (
                [
                    "正式GDS查定比按非MPIS查询量÷航段数重算，分母固定使用seg_num。",
                    "MP_TOTAL承载正式查定比；API_GROUP承载1A、1A_NDC、1A_MPIS等分组接口质量，两者不会重复累计航段。",
                    "选择具体接口分组后只分析该组查询、验价和生单成功/失败，查定比显示为不适用。",
                ]
                if source_key == "gds"
                else [
                    "OTA渠道层只分析查询、验价和生单接口表现，不计算航段查定比。",
                    "OTA与GDS是两个独立分析层，页面不跨层相加。",
                ]
            ),
        }
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            if source_key == "gds":
                self._populate_gds(
                    cursor, result, first, last, dimension_key, cleaned_filters,
                    selected_target, bool(breach_only),
                )
            else:
                self._populate_ota(
                    cursor, result, first, last, dimension_key, cleaned_filters,
                    selected_success_targets, bool(breach_only),
                )
        except Exception as error:
            logging.getLogger(__name__).exception("GDS/OTA resource governance query failed")
            result["error"] = "GDS/OTA资源汇总查询失败，请检查MySQL连接、表结构与同步批次。"
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()
        return result

    def export_report(self, **kwargs: Any) -> tuple[BinaryIO, str]:
        result = self.overview(**kwargs)
        if not result["available"]:
            raise ValueError(result["error"] or "当前范围暂无可导出的资源治理数据")
        layer = result["dataSource"]["label"]
        period = f"{result['period']['startDate']}_{result['period']['endDate']}"
        return _governance_workbook(result), f"{layer}资源治理分析_{period}.xlsx"
