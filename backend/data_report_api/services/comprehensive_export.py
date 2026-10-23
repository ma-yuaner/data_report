"""Persistent, cancellable background exports for comprehensive wide-table details."""
from __future__ import annotations

import hashlib
import json
import logging
import shutil
import sqlite3
import threading
import uuid
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..shared.permissions import permission_codes
from ..time_utils import business_now
from .comprehensive_analysis import BUSINESSES, DIMENSIONS
from .comprehensive_detail import (
    BUSINESS_LABELS,
    DETAIL_SPECS,
    build_detail_queries,
    parse_detail_filters,
)
from .data_source import DataSource


LOGGER = logging.getLogger(__name__)
EXPORT_PERMISSION = "analysis.comprehensive.export"
ACTIVE_STATUSES = ("QUEUED", "RUNNING", "CANCEL_REQUESTED")
MAX_EXPORT_ROWS = 500_000
MAX_ACTIVE_TASKS_PER_USER = 3
BATCH_SIZE = 2_000
GDS_ROUTE_KIND = "gds_routes"
RESOURCE_GDS_KIND = "resource_governance_gds"
RESOURCE_OTA_KIND = "resource_governance_ota"
SMART_TREND_DAILY_KIND = "smart_trend_daily"
SMART_TREND_MONTHLY_KIND = "smart_trend_monthly"
SMART_DIMENSIONS_KIND = "smart_dimensions"
SMART_DETAILS_KIND = "smart_details"
BACKGROUND_EXPORT_KINDS = {
    GDS_ROUTE_KIND,
    RESOURCE_GDS_KIND,
    RESOURCE_OTA_KIND,
    SMART_TREND_DAILY_KIND,
    SMART_TREND_MONTHLY_KIND,
    SMART_DIMENSIONS_KIND,
    SMART_DETAILS_KIND,
}
TASK_LABELS = {
    GDS_ROUTE_KIND: "GDS航线名单",
    RESOURCE_GDS_KIND: "GDS资源治理汇报",
    RESOURCE_OTA_KIND: "OTA渠道质量汇报",
    SMART_TREND_DAILY_KIND: "智能分析每日趋势",
    SMART_TREND_MONTHLY_KIND: "智能分析每月趋势",
    SMART_DIMENSIONS_KIND: "智能分析维度分析",
    SMART_DETAILS_KIND: "智能分析出票明细",
}


class ComprehensiveExportUnavailable(RuntimeError):
    pass


class ComprehensiveExportConflict(RuntimeError):
    pass


class ComprehensiveExportCancelled(RuntimeError):
    pass


def _now_text() -> str:
    return business_now().isoformat(timespec="seconds")


def _is_admin(user: dict[str, Any]) -> bool:
    return bool(user.get("is_admin") or user.get("isAdmin"))


def _require_export(user: dict[str, Any]) -> None:
    if _is_admin(user) or EXPORT_PERMISSION in permission_codes(user):
        return
    raise PermissionError("当前账号没有宽表明细导出权限")


def _require_kind(user: dict[str, Any], kind: str) -> None:
    menu_codes = set(user.get("menu_codes") or user.get("menuCodes") or [])
    if kind == "comprehensive" or kind in BUSINESSES:
        _require_export(user)
        if not (_is_admin(user) or "analysis.comprehensive" in menu_codes):
            raise PermissionError("当前账号没有综合分析菜单权限")
        return
    required_menu = {
        GDS_ROUTE_KIND: "gds_route.diagnosis",
        RESOURCE_GDS_KIND: "governance.gds",
        RESOURCE_OTA_KIND: "governance.ota",
        SMART_TREND_DAILY_KIND: "smart.overview",
        SMART_TREND_MONTHLY_KIND: "smart.overview",
        SMART_DIMENSIONS_KIND: "smart.dimensions",
        SMART_DETAILS_KIND: "smart.details",
    }.get(kind)
    if required_menu is None:
        raise ValueError("不支持的导出任务类型")
    if not (_is_admin(user) or required_menu in menu_codes):
        raise PermissionError("当前账号没有该导出任务的菜单权限")


def _user_id(user: dict[str, Any]) -> int:
    return int(user.get("id") or 0)


def _display_name(user: dict[str, Any]) -> str:
    return str(
        user.get("display_name") or user.get("displayName")
        or user.get("username") or ""
    ).strip()[:100]


class ComprehensiveExportStore:
    def __init__(self, root: str | Path, retention_hours: int = 24):
        self.root = Path(root).expanduser().resolve()
        self.files = self.root / "files"
        self.root.mkdir(parents=True, exist_ok=True)
        self.files.mkdir(parents=True, exist_ok=True)
        self.database = self.root / "jobs.sqlite3"
        self.retention_hours = max(1, retention_hours)
        self._initialize()

    def connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA busy_timeout=30000")
        return connection

    def _initialize(self) -> None:
        with self.connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS comprehensive_export_jobs (
                    id TEXT PRIMARY KEY,
                    user_id INTEGER NOT NULL,
                    requested_by TEXT NOT NULL,
                    fingerprint TEXT NOT NULL,
                    business_type TEXT NOT NULL,
                    start_date TEXT NOT NULL,
                    end_date TEXT NOT NULL,
                    filters_json TEXT NOT NULL,
                    field_version TEXT NOT NULL,
                    status TEXT NOT NULL,
                    total_rows INTEGER NOT NULL DEFAULT 0,
                    processed_rows INTEGER NOT NULL DEFAULT 0,
                    queued_at TEXT NOT NULL,
                    started_at TEXT,
                    finished_at TEXT,
                    updated_at TEXT NOT NULL,
                    cancel_requested_at TEXT,
                    file_path TEXT,
                    download_name TEXT,
                    message TEXT NOT NULL DEFAULT ''
                )
                """
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_comprehensive_export_user_time "
                "ON comprehensive_export_jobs(user_id, queued_at DESC)"
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_comprehensive_export_active "
                "ON comprehensive_export_jobs(status, user_id, queued_at)"
            )

    @staticmethod
    def public(row: sqlite3.Row | dict[str, Any], *, reused: bool = False) -> dict[str, Any]:
        item = dict(row)
        task_label = TASK_LABELS.get(item["business_type"])
        business_label = task_label or BUSINESS_LABELS.get(item["business_type"], item["business_type"])
        return {
            "id": item["id"],
            "businessType": item["business_type"],
            "businessLabel": business_label,
            "taskLabel": task_label or business_label + "宽表明细",
            "startDate": item["start_date"],
            "endDate": item["end_date"],
            "status": item["status"],
            "totalRows": int(item.get("total_rows") or 0),
            "processedRows": int(item.get("processed_rows") or 0),
            "queuedAt": item["queued_at"],
            "startedAt": item.get("started_at"),
            "finishedAt": item.get("finished_at"),
            "updatedAt": item["updated_at"],
            "message": item.get("message") or "",
            "downloadName": item.get("download_name"),
            "canCancel": item["status"] in ("QUEUED", "RUNNING"),
            "canDownload": item["status"] == "SUCCEEDED" and bool(item.get("file_path")),
            "reused": reused,
        }

    def create(
        self, *, user_id: int, requested_by: str, fingerprint: str,
        business: str, start_date: str, end_date: str, filters: dict[str, Any],
        total_rows: int,
    ) -> dict[str, Any]:
        connection = self.connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            duplicate = connection.execute(
                "SELECT * FROM comprehensive_export_jobs "
                "WHERE user_id=? AND fingerprint=? AND status IN ('QUEUED','RUNNING','CANCEL_REQUESTED') "
                "ORDER BY queued_at DESC LIMIT 1",
                (user_id, fingerprint),
            ).fetchone()
            if duplicate is not None:
                connection.commit()
                return self.public(duplicate, reused=True)
            active_count = connection.execute(
                "SELECT COUNT(1) FROM comprehensive_export_jobs "
                "WHERE user_id=? AND status IN ('QUEUED','RUNNING','CANCEL_REQUESTED')",
                (user_id,),
            ).fetchone()[0]
            if int(active_count or 0) >= MAX_ACTIVE_TASKS_PER_USER:
                raise ComprehensiveExportConflict("已有3个未完成导出任务，请等待或取消后再试")
            task_id = uuid.uuid4().hex
            now = _now_text()
            connection.execute(
                """
                INSERT INTO comprehensive_export_jobs (
                    id,user_id,requested_by,fingerprint,business_type,start_date,end_date,
                    filters_json,field_version,status,total_rows,processed_rows,queued_at,
                    updated_at,message
                ) VALUES (?,?,?,?,?,?,?,?,?,'QUEUED',?,0,?,?,?)
                """,
                (
                    task_id, user_id, requested_by, fingerprint, business, start_date, end_date,
                    json.dumps(filters, ensure_ascii=False, sort_keys=True), "v1", total_rows,
                    now, now, "等待导出",
                ),
            )
            row = connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE id=?", (task_id,)
            ).fetchone()
            connection.commit()
            return self.public(row)
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def list(self, user_id: int, limit: int = 30) -> list[dict[str, Any]]:
        self.expire_files()
        with self.connect() as connection:
            rows = connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE user_id=? "
                "ORDER BY queued_at DESC LIMIT ?",
                (user_id, max(1, min(int(limit), 100))),
            ).fetchall()
        return [self.public(row) for row in rows]

    def get(self, user_id: int, task_id: str) -> sqlite3.Row | None:
        with self.connect() as connection:
            return connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE id=? AND user_id=?",
                (task_id, user_id),
            ).fetchone()

    def cancel(self, user_id: int, task_id: str) -> dict[str, Any]:
        connection = self.connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE id=? AND user_id=?",
                (task_id, user_id),
            ).fetchone()
            if row is None:
                raise ValueError("导出任务不存在")
            status = str(row["status"])
            now = _now_text()
            if status == "QUEUED":
                connection.execute(
                    "UPDATE comprehensive_export_jobs SET status='CANCELLED',finished_at=?,updated_at=?,"
                    "cancel_requested_at=?,message='已取消' WHERE id=?",
                    (now, now, now, task_id),
                )
            elif status == "RUNNING":
                connection.execute(
                    "UPDATE comprehensive_export_jobs SET status='CANCEL_REQUESTED',updated_at=?,"
                    "cancel_requested_at=?,message='正在取消' WHERE id=?",
                    (now, now, task_id),
                )
            elif status == "CANCEL_REQUESTED":
                pass
            else:
                raise ComprehensiveExportConflict("当前任务已结束，不能取消")
            result = connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE id=?", (task_id,)
            ).fetchone()
            connection.commit()
            return self.public(result)
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def claim_next(self) -> dict[str, Any] | None:
        connection = self.connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                """
                SELECT q.* FROM comprehensive_export_jobs q
                WHERE q.status='QUEUED' AND NOT EXISTS (
                    SELECT 1 FROM comprehensive_export_jobs r
                    WHERE r.user_id=q.user_id AND r.status IN ('RUNNING','CANCEL_REQUESTED')
                )
                ORDER BY q.queued_at LIMIT 1
                """
            ).fetchone()
            if row is None:
                connection.commit()
                return None
            now = _now_text()
            changed = connection.execute(
                "UPDATE comprehensive_export_jobs SET status='RUNNING',started_at=?,updated_at=?,"
                "message='正在生成文件' WHERE id=? AND status='QUEUED'",
                (now, now, row["id"]),
            ).rowcount
            if changed != 1:
                connection.rollback()
                return None
            connection.commit()
            result = dict(row)
            result.update(status="RUNNING", started_at=now, updated_at=now, message="正在生成文件")
            return result
        finally:
            connection.close()

    def cancellation_requested(self, task_id: str) -> bool:
        with self.connect() as connection:
            row = connection.execute(
                "SELECT status FROM comprehensive_export_jobs WHERE id=?", (task_id,)
            ).fetchone()
        return row is None or row["status"] == "CANCEL_REQUESTED"

    def progress(self, task_id: str, processed: int) -> None:
        now = _now_text()
        with self.connect() as connection:
            connection.execute(
                "UPDATE comprehensive_export_jobs SET processed_rows=?,updated_at=?,message=? "
                "WHERE id=? AND status='RUNNING'",
                (processed, now, f"已处理 {processed:,} 条", task_id),
            )

    def finish(self, task_id: str, path: Path, download_name: str, processed: int) -> None:
        now = _now_text()
        with self.connect() as connection:
            changed = connection.execute(
                "UPDATE comprehensive_export_jobs SET status='SUCCEEDED',processed_rows=?,"
                "finished_at=?,updated_at=?,file_path=?,download_name=?,message='导出完成' WHERE id=? AND status='RUNNING'",
                (processed, now, now, str(path), download_name, task_id),
            ).rowcount
            if changed != 1:
                raise ComprehensiveExportCancelled()

    def set_total(self, task_id: str, total: int) -> None:
        with self.connect() as connection:
            connection.execute("UPDATE comprehensive_export_jobs SET total_rows=?,updated_at=? WHERE id=? AND status='RUNNING'",
                               (total, _now_text(), task_id))

    def fail(self, task_id: str, message: str, *, cancelled: bool = False) -> None:
        now = _now_text()
        status = "CANCELLED" if cancelled else "FAILED"
        safe_message = "已取消" if cancelled else str(message).strip()[:500] or "导出失败"
        with self.connect() as connection:
            connection.execute(
                "UPDATE comprehensive_export_jobs SET status=?,finished_at=?,updated_at=?,"
                "file_path=NULL,download_name=NULL,message=? WHERE id=?",
                (status, now, now, safe_message, task_id),
            )

    def mark_downloaded(self, user_id: int, task_id: str) -> None:
        with self.connect() as connection:
            row = connection.execute(
                "SELECT status,file_path FROM comprehensive_export_jobs WHERE id=? AND user_id=?",
                (task_id, user_id),
            ).fetchone()
            if row is None or row["status"] != "SUCCEEDED":
                return
            path = Path(row["file_path"]) if row["file_path"] else None
            if path:
                path.unlink(missing_ok=True)
            connection.execute(
                "UPDATE comprehensive_export_jobs SET status='DOWNLOADED',file_path=NULL,"
                "updated_at=?,message='已下载，服务器文件已清理' WHERE id=? AND user_id=?",
                (_now_text(), task_id, user_id),
            )

    def expire_files(self) -> None:
        threshold = (business_now() - timedelta(hours=self.retention_hours)).isoformat(timespec="seconds")
        with self.connect() as connection:
            rows = connection.execute(
                "SELECT id,file_path FROM comprehensive_export_jobs "
                "WHERE status='SUCCEEDED' AND finished_at<?", (threshold,),
            ).fetchall()
            for row in rows:
                path = Path(row["file_path"]) if row["file_path"] else None
                if path:
                    path.unlink(missing_ok=True)
                connection.execute(
                    "UPDATE comprehensive_export_jobs SET status='EXPIRED',file_path=NULL,"
                    "download_name=NULL,message='文件已过期' WHERE id=?", (row["id"],),
                )


class ComprehensiveExportService:
    def __init__(self, config: dict[str, Any]):
        mysql_config = dict(config)
        mysql_config["DATA_MODE"] = "mysql"
        self.config = {str(key): value for key, value in mysql_config.items() if str(key).isupper()}
        self.source = DataSource(mysql_config)
        self.store = ComprehensiveExportStore(
            str(config.get("COMPREHENSIVE_EXPORT_ROOT", "/app/var/comprehensive-exports")),
            int(config.get("COMPREHENSIVE_EXPORT_RETENTION_HOURS", 24)),
        )

    @staticmethod
    def _fingerprint(user_id: int, business: str, start_date: str, end_date: str,
                     filters: dict[str, Any]) -> str:
        payload = json.dumps(
            {"user": user_id, "business": business, "start": start_date, "end": end_date,
             "filters": filters, "fields": "v1"},
            ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        )
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    def _validated_scope(self, values: dict[str, Any]) -> tuple[str, date, date, dict[str, list[str | None]]]:
        business = str(values.get("businessType") or "")
        if business not in BUSINESSES:
            raise ValueError("不支持的业务类型")
        from ..shared.periods import parse_analysis_period
        first, last = parse_analysis_period(values.get("startDate"), values.get("endDate"))
        if (last - first).days > 30:
            raise ValueError("宽表导出单次最多31天")
        filters = parse_detail_filters({key: values.get(key) for key in DIMENSIONS})
        return business, first, last, filters

    def _count(self, business: str, first: date, last: date,
               filters: dict[str, list[str | None]]) -> int:
        count_query, _, count_params, _ = build_detail_queries(
            database=self.source.database, business=business,
            start_date=first.isoformat(), next_date=(last + timedelta(days=1)).isoformat(),
            filters=filters, page=1, page_size=1,
        )
        connection = cursor = None
        try:
            connection = self.source.connect()
            cursor = connection.cursor()
            cursor.execute(count_query, count_params)
            return int(cursor.fetchone()[0] or 0)
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()
    def create(self, user: dict[str, Any], values: dict[str, Any]) -> dict[str, Any]:
        _require_kind(user, "comprehensive")
        business, first, last, filters = self._validated_scope(values)
        user_id = _user_id(user)
        fingerprint = self._fingerprint(user_id, business, first.isoformat(), last.isoformat(), filters)
        with self.store.connect() as connection:
            duplicate = connection.execute(
                "SELECT * FROM comprehensive_export_jobs WHERE user_id=? AND fingerprint=? "
                "AND status IN ('QUEUED','RUNNING','CANCEL_REQUESTED') ORDER BY queued_at DESC LIMIT 1",
                (user_id, fingerprint),
            ).fetchone()
        if duplicate is not None:
            return self.store.public(duplicate, reused=True)
        try:
            total = self._count(business, first, last, filters)
        except Exception as error:
            LOGGER.exception("Comprehensive export count failed")
            raise ComprehensiveExportUnavailable("导出数量查询失败，请稍后重试") from error
        if total <= 0:
            raise ValueError("当前条件没有可导出的宽表明细")
        if total > MAX_EXPORT_ROWS:
            raise ValueError(
                f"当前匹配 {total:,} 条，超过 {MAX_EXPORT_ROWS:,} 条导出上限。请缩小日期或筛选范围后再导出。"
            )
        task = self.store.create(
            user_id=user_id, requested_by=_display_name(user), fingerprint=fingerprint,
            business=business, start_date=first.isoformat(), end_date=last.isoformat(),
            filters=filters, total_rows=total,
        )
        self._start_worker()
        return task

    def list(self, user: dict[str, Any], limit: int = 30) -> list[dict[str, Any]]:
        tasks = []
        for task in self.store.list(_user_id(user), limit):
            try:
                _require_kind(user, task["businessType"])
            except PermissionError:
                continue
            tasks.append(task)
        if any(task["status"] == "QUEUED" for task in tasks):
            self._start_worker()
        return tasks

    def cancel(self, user: dict[str, Any], task_id: str) -> dict[str, Any]:
        row = self.store.get(_user_id(user), task_id)
        if row is None:
            raise ValueError("导出任务不存在")
        _require_kind(user, row["business_type"])
        result = self.store.cancel(_user_id(user), task_id)
        self._start_worker()
        return result

    def download(self, user: dict[str, Any], task_id: str) -> tuple[Path, str]:
        self.store.expire_files()
        row = self.store.get(_user_id(user), task_id)
        if row is None:
            raise ValueError("导出任务不存在")
        _require_kind(user, row["business_type"])
        if row["status"] != "SUCCEEDED" or not row["file_path"]:
            raise ComprehensiveExportConflict("导出文件尚不可下载")
        path = Path(row["file_path"])
        if not path.is_file():
            raise ComprehensiveExportUnavailable("导出文件不存在或已清理")
        return path, str(row["download_name"] or path.name)

    def mark_downloaded(self, user: dict[str, Any], task_id: str) -> None:
        row = self.store.get(_user_id(user), task_id)
        if row is None:
            return
        _require_kind(user, str(row["business_type"]))
        self.store.mark_downloaded(_user_id(user), task_id)

    def _start_worker(self) -> None:
        threading.Thread(
            target=self._worker_loop, name="comprehensive-export-worker", daemon=True
        ).start()

    def _worker_loop(self) -> None:
        while True:
            task = self.store.claim_next()
            if task is None:
                return
            self._execute(task)

    def _execute(self, task: dict[str, Any]) -> None:
        task_id = str(task["id"])
        output_path = self.store.files / f"{task_id}.xlsx"
        try:
            if task["business_type"] == GDS_ROUTE_KIND:
                self._write_gds_routes(task, output_path)
            elif task["business_type"] in BACKGROUND_EXPORT_KINDS:
                self._write_background(task, output_path)
            else:
                self._write_xlsx(task, output_path)
        except ComprehensiveExportCancelled:
            output_path.unlink(missing_ok=True)
            self.store.fail(task_id, "已取消", cancelled=True)
        except ValueError as error:
            output_path.unlink(missing_ok=True)
            self.store.fail(task_id, str(error))
        except Exception as error:
            output_path.unlink(missing_ok=True)
            LOGGER.exception("Comprehensive export failed; task_id=%s", task_id)
            self.store.fail(task_id, "导出失败，请稍后重试")

    def create_gds_routes(self, user: dict[str, Any], values: dict[str, Any]) -> dict[str, Any]:
        from .gds_route_analysis import GdsRouteAnalysisService
        _require_kind(user, GDS_ROUTE_KIND)
        service = GdsRouteAnalysisService(self.config)
        first, last, filters = service.export_scope(values.get("startDate"), values.get("endDate"), values)
        fingerprint = self._fingerprint(_user_id(user), GDS_ROUTE_KIND, first.isoformat(), last.isoformat(), filters)
        task = self.store.create(user_id=_user_id(user), requested_by=_display_name(user), fingerprint=fingerprint,
                                 business=GDS_ROUTE_KIND, start_date=first.isoformat(), end_date=last.isoformat(),
                                 filters=filters, total_rows=0)
        self._start_worker()
        return task

    def create_background(self, user: dict[str, Any], values: dict[str, Any]) -> dict[str, Any]:
        kind = str(values.get("kind") or "")
        if kind not in BACKGROUND_EXPORT_KINDS - {GDS_ROUTE_KIND}:
            raise ValueError("不支持的导出任务类型")
        _require_kind(user, kind)
        payload = {str(key): value for key, value in values.items() if key != "kind"}
        start_date = str(payload.get("startDate") or payload.get("orderStart") or payload.get("firstDepStart") or "")
        end_date = str(payload.get("endDate") or payload.get("orderEnd") or payload.get("firstDepEnd") or "")
        fingerprint = self._fingerprint(_user_id(user), kind, start_date, end_date, payload)
        task = self.store.create(
            user_id=_user_id(user), requested_by=_display_name(user), fingerprint=fingerprint,
            business=kind, start_date=start_date, end_date=end_date, filters=payload, total_rows=0,
        )
        self._start_worker()
        return task

    def _write_gds_routes(self, task: dict[str, Any], output_path: Path) -> None:
        from .gds_route_analysis import GdsRouteAnalysisService
        import shutil
        task_id = str(task["id"])

        def check():
            if self.store.cancellation_requested(task_id):
                raise ComprehensiveExportCancelled()

        def progress(done, total):
            self.store.set_total(task_id, total)
            self.store.progress(task_id, done)

        check()
        self.store.progress(task_id, 0)
        content, name = GdsRouteAnalysisService(self.config).export_report(
            task["start_date"], task["end_date"], json.loads(task["filters_json"]),
            check_cancel=check, progress=progress)
        try:
            check()
            content.seek(0)
            with output_path.open("wb") as target:
                shutil.copyfileobj(content, target)
            check()
            row = self.store.get(int(task["user_id"]), task_id)
            self.store.finish(task_id, output_path, name, int(row["total_rows"]))
        finally:
            content.close()

    def _write_background(self, task: dict[str, Any], output_path: Path) -> None:
        from .resource_governance import ResourceGovernanceService
        from .smart_analysis import DIMENSIONS as SMART_DIMENSIONS, SmartAnalysisService

        task_id = str(task["id"])
        kind = str(task["business_type"])
        values = json.loads(task["filters_json"])
        if self.store.cancellation_requested(task_id):
            raise ComprehensiveExportCancelled()
        if kind in {RESOURCE_GDS_KIND, RESOURCE_OTA_KIND}:
            data_source = "gds" if kind == RESOURCE_GDS_KIND else "ota"
            content, name = ResourceGovernanceService(self.config).export_report(
                start_value=values.get("startDate"), end_value=values.get("endDate"),
                data_source=data_source, group_by=values.get("groupBy"),
                filters={key: values.get(key, "") for key in (
                    "platform", "site", "gds", "resourceChannel", "pcc", "apiGroup"
                )},
                breach_only=bool(values.get("breachOnly")),
            )
        else:
            service = SmartAnalysisService(self.config)
            period_args = {
                "start_value": values.get("startDate"), "end_value": values.get("endDate"),
                "first_dep_start": values.get("firstDepStart"), "first_dep_end": values.get("firstDepEnd"),
                "order_start": values.get("orderStart"), "order_end": values.get("orderEnd"),
            }
            filters = {key: values.get(key) for key in (*SMART_DIMENSIONS, "otaOrderNo")}
            if kind in {SMART_TREND_DAILY_KIND, SMART_TREND_MONTHLY_KIND}:
                content, name = service.export_trend(
                    **period_args, filters=filters,
                    kind="daily" if kind == SMART_TREND_DAILY_KIND else "monthly",
                    smart_type=values.get("smartType", "all"),
                    trend_date=values.get("trendDate", "issue"),
                )
            elif kind == SMART_DIMENSIONS_KIND:
                content, name = service.export_dimensions(
                    **period_args, filters=filters,
                    group_by=",".join(values.get("groupBy") or []),
                    metrics_value=",".join(values.get("metrics") or []),
                    smart_type=values.get("smartType", "all"),
                    sort_by=values.get("sortBy", "smartProfit"),
                    direction=values.get("direction", "desc"),
                )
            elif kind == SMART_DETAILS_KIND:
                content, name = service.export_details(
                    **period_args, filters=filters,
                    smart_type=values.get("smartType", "all"), stage=values.get("stage", "all"),
                    fields_value=",".join(values.get("fields") or []),
                    all_fields=values.get("scope", "visible") == "all",
                )
            else:
                raise ValueError("不支持的导出任务类型")
        if self.store.cancellation_requested(task_id):
            raise ComprehensiveExportCancelled()
        try:
            content.seek(0)
            with output_path.open("wb") as target:
                shutil.copyfileobj(content, target)
            if self.store.cancellation_requested(task_id):
                raise ComprehensiveExportCancelled()
            self.store.set_total(task_id, 1)
            self.store.finish(task_id, output_path, name, 1)
        finally:
            content.close()

    def _write_xlsx(self, task: dict[str, Any], output_path: Path) -> None:
        import pymysql

        business = str(task["business_type"])
        filters = json.loads(str(task["filters_json"]))
        total = int(task["total_rows"])
        spec = DETAIL_SPECS[business]
        _, data_query, _, data_params = build_detail_queries(
            database=self.source.database, business=business,
            start_date=str(task["start_date"]),
            next_date=(date.fromisoformat(str(task["end_date"])) + timedelta(days=1)).isoformat(),
            filters=filters, page=1, page_size=MAX_EXPORT_ROWS,
        )
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(f"{BUSINESS_LABELS[business]}宽表明细")
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(spec.columns))}{total + 1}"
        )
        for index, item in enumerate(spec.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(10, min(35, item.width / 7))
        headers = []
        for item in spec.columns:
            cell = WriteOnlyCell(sheet, value=item.title)
            cell.font = Font(bold=True)
            headers.append(cell)
        sheet.append(headers)
        connection = cursor = None
        processed = 0
        try:
            connection = self.source.connect()
            cursor = connection.cursor(pymysql.cursors.SSCursor)
            cursor.execute(data_query, data_params)
            while True:
                if self.store.cancellation_requested(str(task["id"])):
                    raise ComprehensiveExportCancelled()
                rows = cursor.fetchmany(BATCH_SIZE)
                if not rows:
                    break
                for raw_row in rows:
                    cells = []
                    for item, value in zip(spec.columns, raw_row):
                        cell = WriteOnlyCell(sheet, value=value)
                        if item.value_type == "text":
                            cell.number_format = "@"
                        elif item.value_type == "date":
                            cell.number_format = "yyyy-mm-dd"
                        elif item.value_type == "datetime":
                            cell.number_format = "yyyy-mm-dd hh:mm:ss"
                        elif item.value_type == "count":
                            cell.number_format = "#,##0"
                        elif item.value_type == "money":
                            cell.number_format = "#,##0.00;-#,##0.00"
                        cells.append(cell)
                    sheet.append(cells)
                processed += len(rows)
                self.store.progress(str(task["id"]), processed)
            if processed != total:
                raise RuntimeError(f"导出期间数据量发生变化：预期{total}，实际{processed}")
            if self.store.cancellation_requested(str(task["id"])):
                raise ComprehensiveExportCancelled()
            workbook.save(output_path)
            download_name = (
                f"综合分析_{BUSINESS_LABELS[business]}宽表明细_"
                f"{task['start_date']}_{task['end_date']}.xlsx"
            )
            self.store.finish(str(task["id"]), output_path, download_name, processed)
        finally:
            if cursor is not None:
                cursor.close()
            if connection is not None:
                connection.close()
